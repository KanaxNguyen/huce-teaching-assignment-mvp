"""Copy source workbooks into ignored local storage and create an auditable manifest."""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import xlrd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_ROOT = PROJECT_ROOT.parent
TARGET_DIR = PROJECT_ROOT / "data" / "local" / "source"
MANIFEST_PATH = PROJECT_ROOT / "data" / "manifests" / "source-files.json"

NORMALIZED_NAMES = {
    "Nguyện vọng TKB 2026-2027.xlsx": ("teacher-preferences.xlsx", "preference"),
    "phan-cong-giang-day-lich-hoc-to-bo-mon-26-07-2026_TKB.xls": (
        "schedule-2026-07-26.xls",
        "schedule",
    ),
    "phan-cong-giang-day-lich-hoc-to-bo-mon-28-07-2026.xls": (
        "schedule-2026-07-28.xls",
        "schedule",
    ),
    "Phan_Cong_Giang_Day_HK1_2026_2027.xlsx": (
        "previous-assignments.xlsx",
        "previous_assignment",
    ),
    "Thoi_Khoa_Bieu_To_Bo_Mon_Nguyen_Ven_Lop.xlsx": (
        "expected-output-template.xlsx",
        "expected_output",
    ),
}


def workbook_info(path: Path) -> tuple[list[str], list[str]]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        names = book.sheet_names()
        sample = []
        for sheet in book.sheets():
            for row in range(min(sheet.nrows, 20)):
                sample.extend(
                    str(v) for v in sheet.row_values(row)[:14] if v not in ("", None)
                )
    else:
        book = openpyxl.load_workbook(path, read_only=False, data_only=True)
        names = book.sheetnames
        sample = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows(
                min_row=1, max_row=min(sheet.max_row, 20), max_col=14, values_only=True
            ):
                sample.extend(str(v) for v in row if v not in ("", None))
    keys = ["STT", "Mã học phần", "Tên môn học", "Mã lớp học", "Tên", "Thứ 2"]
    joined = "\n".join(sample).casefold()
    return names, [key for key in keys if key.casefold() in joined]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> int:
    TARGET_DIR.mkdir(parents=True, exist_ok=True)
    entries = []
    now = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).isoformat()

    for source in sorted((*INPUT_ROOT.glob("*.xls"), *INPUT_ROOT.glob("*.xlsx"))):
        normalized, detected = NORMALIZED_NAMES.get(
            source.name, (source.name, "unknown")
        )
        destination = TARGET_DIR / normalized
        if not destination.exists() or sha256(destination) != sha256(source):
            shutil.copy2(source, destination)
        sheets, headers = workbook_info(destination)
        entries.append(
            {
                "original_name": source.name,
                "original_path": str(source),
                "copied_path": str(destination.relative_to(PROJECT_ROOT)).replace(
                    "\\", "/"
                ),
                "detected_type": detected,
                "extension": source.suffix.lower(),
                "size_bytes": destination.stat().st_size,
                "sha256": sha256(destination),
                "sheet_names": sheets,
                "header_candidates": headers,
                "modified_at": datetime.fromtimestamp(
                    source.stat().st_mtime, ZoneInfo("Asia/Ho_Chi_Minh")
                ).isoformat(),
                "imported_at": now,
            }
        )

    manifest = {
        "generated_at": now,
        "primary_schedule": "data/local/source/schedule-2026-07-28.xls",
        "files": entries,
    }
    MANIFEST_PATH.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Wrote {MANIFEST_PATH} with {len(entries)} entries.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
