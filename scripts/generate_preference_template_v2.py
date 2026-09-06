"""Generate the canonical lecturer-preference template V2."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "data" / "fixtures" / "Template_Nguyen_vong_Giang_v2.xlsx"
CONSTRAINTS = [
    "UNAVAILABLE", "AVOID_PERIOD", "PREFERRED_PERIOD", "PREFER_CONSECUTIVE_PERIODS",
    "MIN_CLASSES", "MAX_CLASSES", "MAX_SESSIONS_PER_DAY", "MAX_DAYS_PER_WEEK",
    "MIN_FREE_MORNING_PER_WEEK", "REQUIRED_ASSIGNMENT", "FORBIDDEN_ASSIGNMENT", "RAW_NOTE",
    "SEMINAR_COMMITMENT", "SEMINAR_NOTE",
]
DAYS = ["T2", "T3", "T4", "T5", "T6", "T7", "CN", "ALL_WEEKDAYS", "ALL_DAYS"]
HARDNESS = ["SOFT", "HARD"]
STATUSES = ["DRAFT", "CONFIRMED", "NEEDS_REVIEW", "REJECTED"]
PREFERENCE_HEADERS = [
    "Mã GV*", "Họ tên GV", "Ngữ cảnh*", "Loại ràng buộc*", "Thứ/Phạm vi*", "Tiết bắt đầu",
    "Tiết kết thúc", "Ngày bắt đầu", "Ngày kết thúc", "Độ cứng*", "Trọng số",
    "Giá trị số", "Ghi chú / Nguyên văn", "Trạng thái",
]
SEMINAR_HEADERS = [
    "Mã seminar*", "Tên seminar*", "Mã GV tham gia*", "Thứ cho phép*",
    "Khung tiết cho phép*", "Độ cứng*", "Trọng số", "Ghi chú", "Trạng thái",
]


def _style_table(sheet, widths):
    navy = "17324D"; pale = "EAF2F8"; border = Side(style="thin", color="CBD5E1")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=border)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(widths)).coordinate}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in range(2, 202):
        for cell in sheet[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E2E8F0"))
        if row % 2 == 0:
            for cell in sheet[row]:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    sheet.conditional_formatting.add(
        f"N2:N201", FormulaRule(formula=['N2="NEEDS_REVIEW"'], fill=PatternFill("solid", fgColor="FDE68A"))
    )


def _list_validation(sheet, formula, cells):
    validation = DataValidation(type="list", formula1=formula, allow_blank=False)
    validation.error = "Chọn một giá trị trong danh mục."
    validation.errorTitle = "Giá trị không hợp lệ"
    validation.prompt = "Dùng danh mục chuẩn để hệ thống đọc chính xác."
    validation.promptTitle = "Danh mục chuẩn"
    validation.showErrorMessage = True; validation.showInputMessage = True
    sheet.add_data_validation(validation); validation.add(cells)


def generate(path: Path = DEFAULT_OUTPUT) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    guide = book.active; guide.title = "Huong_dan"
    preferences = book.create_sheet("Nguyen_vong_GV")
    seminars = book.create_sheet("Seminar_Shared")
    catalogs = book.create_sheet("Danh_muc")

    guide.append(["TEMPLATE NGUYỆN VỌNG GIẢNG VIÊN V2"])
    guide.append(["Nguyên tắc", "Mỗi dòng là một quy tắc nguyên tử; không gộp nhiều yêu cầu vào một dòng."])
    guide.append(["Nhận dạng", "Mã GV là định danh ưu tiên. Thiếu hoặc không khớp sẽ chuyển NEEDS_REVIEW, không áp dụng toàn cục."])
    guide.append(["Khoảng tiết", "Bao gồm cả hai đầu: 4–6 nghĩa là tiết 4, 5 và 6."])
    guide.append(["Hard/Soft", "HARD không phụ thuộc trọng số. SOFT weight 1.0 vẫn là SOFT."])
    guide.append(["RAW_NOTE", "Chỉ lưu nguyên văn để duyệt; không bao giờ đưa trực tiếp vào solver."])
    guide.append(["Quy trình", "Upload → tạo draft → trưởng bộ môn duyệt/chỉnh/từ chối → áp dụng → solver."])
    guide.append(["Ngày", "T2…T7, CN; ALL_WEEKDAYS = T2–T6; ALL_DAYS = T2–CN."])
    guide.append(["Seminar", "Một dòng là một sự kiện chung; mã người tham gia cách nhau bằng dấu chấm phẩy."])
    guide.append(["Ngữ cảnh", "TEACHING = lịch dạy; SEMINAR = thông tin seminar thuộc giảng viên; MIXED = phải tách thành rule nguyên tử trước khi áp dụng."])
    guide.column_dimensions["A"].width = 20; guide.column_dimensions["B"].width = 110
    guide.sheet_properties.pageSetUpPr.fitToPage = True; guide.page_setup.orientation = "landscape"
    guide.page_setup.fitToWidth = 1; guide.page_setup.fitToHeight = 1
    guide["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    guide["A1"].fill = PatternFill("solid", fgColor="17324D")
    guide.merge_cells("A1:B1"); guide.row_dimensions[1].height = 34
    for row in range(2, guide.max_row + 1):
        guide.cell(row, 1).font = Font(bold=True, color="17324D")
        guide.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    preferences.append(PREFERENCE_HEADERS)
    examples = [
        ["GV001", "Nguyễn Văn A", "TEACHING", "AVOID_PERIOD", "ALL_WEEKDAYS", 1, 3, None, None, "SOFT", 0.8, None, "Xin tránh tiết 1-3 các buổi", "DRAFT"],
        ["GV002", "Trần Thị B", "TEACHING", "PREFERRED_PERIOD", "ALL_WEEKDAYS", 1, 6, None, None, "SOFT", 1.0, None, "Ưu tiên dạy buổi sáng", "DRAFT"],
        ["GV003", "Lê Văn C", "TEACHING", "UNAVAILABLE", "T6", 1, 6, None, None, "HARD", 1.0, None, "Xin nghỉ sáng thứ 6", "DRAFT"],
        ["GV004", "Phạm Thị D", "TEACHING", "PREFER_CONSECUTIVE_PERIODS", "ALL_WEEKDAYS", None, None, None, None, "SOFT", 0.8, 6, "Xin dạy liền 6 tiết", "DRAFT"],
        ["GV005", "Đỗ Văn E", "TEACHING", "PREFERRED_PERIOD", "ALL_WEEKDAYS", 7, 12, None, None, "SOFT", 0.7, None, "Xin dạy buổi chiều", "DRAFT"],
        ["GV006", "Vũ Thị F", "TEACHING", "RAW_NOTE", "ALL_WEEKDAYS", None, None, None, None, "SOFT", 0.5, None, "Xin dạy ít", "NEEDS_REVIEW"],
        ["GV007", "Bùi Văn G", "TEACHING", "UNAVAILABLE", "T4", 10, 12, None, None, "HARD", 1.0, None, "Không thể dạy thứ 4 tiết 10-12", "DRAFT"],
        ["GV008", "Đặng Thị H", "SEMINAR", "SEMINAR_COMMITMENT", "T5", 4, 6, None, None, "SOFT", 0.8, None, "Seminar thứ 5 tiết 4-6", "DRAFT"],
    ]
    for row in examples: preferences.append(row)
    _style_table(preferences, [14, 24, 17, 31, 19, 13, 13, 15, 15, 14, 11, 12, 42, 18])
    _list_validation(preferences, "=Danh_muc!$E$2:$E$4", "C2:C201")
    _list_validation(preferences, "=Danh_muc!$A$2:$A$15", "D2:D201")
    _list_validation(preferences, "=Danh_muc!$B$2:$B$10", "E2:E201")
    _list_validation(preferences, "=Danh_muc!$C$2:$C$3", "J2:J201")
    _list_validation(preferences, "=Danh_muc!$D$2:$D$5", "N2:N201")
    period_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="15", allow_blank=True)
    preferences.add_data_validation(period_validation); period_validation.add("F2:G201")
    weight_validation = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    preferences.add_data_validation(weight_validation); weight_validation.add("K2:K201")

    seminars.append(SEMINAR_HEADERS)
    seminars.append(["SEM001", "Seminar giáo trình", "GV001;GV002;GV003", "T4;T5", "4-6;10-12", "SOFT", 0.8, "Một event chung cho mọi người tham gia", "DRAFT"])
    _style_table(seminars, [17, 28, 35, 22, 25, 14, 11, 42, 18])
    _list_validation(seminars, "=Danh_muc!$C$2:$C$3", "F2:F201")
    _list_validation(seminars, "=Danh_muc!$D$2:$D$5", "I2:I201")

    contexts = ["TEACHING", "SEMINAR", "MIXED"]
    catalogs.append(["Loại ràng buộc", "Thứ/Phạm vi", "Độ cứng", "Trạng thái", "Ngữ cảnh"])
    for index in range(max(len(CONSTRAINTS), len(DAYS), len(HARDNESS), len(STATUSES), len(contexts))):
        catalogs.append([
            CONSTRAINTS[index] if index < len(CONSTRAINTS) else None,
            DAYS[index] if index < len(DAYS) else None,
            HARDNESS[index] if index < len(HARDNESS) else None,
            STATUSES[index] if index < len(STATUSES) else None,
            contexts[index] if index < len(contexts) else None,
        ])
    _style_table(catalogs, [34, 22, 16, 20, 18])
    catalogs.sheet_state = "hidden"
    book.save(path)
    return path


if __name__ == "__main__":
    print(generate())
