from collections import defaultdict
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
import openpyxl
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.models.entities import (
    ClassSection, Course, HistoricalEvidence, Lecturer, LecturerAlias,
    LecturerCourseCapability, LecturerSemesterProfile, NormalizedPreferenceDraft,
    Semester,
)
from app.services.lecturer_master import confirm_alias, dependencies, merge_lecturers, normalize_identity, relink_source_identity

router = APIRouter(prefix="/api/v1")


class LecturerInput(BaseModel):
    code: str | None = Field(default=None, max_length=30)
    name: str = Field(min_length=1, max_length=200)
    department: str | None = None
    email: str | None = None
    note: str | None = None
    status: Literal["ACTIVE", "INACTIVE"] = "ACTIVE"
    aliases: list[str] = Field(default_factory=list)
    participation_status: Literal["ACTIVE", "NOT_PARTICIPATING", "ON_LEAVE"] | None = None


class ParticipationInput(BaseModel):
    participation_status: Literal["ACTIVE", "NOT_PARTICIPATING", "ON_LEAVE"]
    target_workload: float | None = Field(default=None, ge=0)
    min_workload: float | None = Field(default=None, ge=0)
    max_workload: float | None = Field(default=None, ge=0)
    note: str | None = None


class CapabilityInput(BaseModel):
    course_id: int
    allowed: bool
    confirmed: bool = True


class IdentityCorrection(BaseModel):
    lecturer_id: int | None = None
    confirm_alias: bool = False
    replace_alias: bool = False
    ignore: bool = False


class MergeInput(BaseModel):
    target_id: int
    confirmed: bool = False


def require_lecturer(db, lecturer_id):
    item = db.get(Lecturer, lecturer_id)
    if not item:
        raise HTTPException(404, "Không tìm thấy giảng viên.")
    return item


def commit(db):
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "Mã giảng viên hoặc liên kết đã tồn tại.") from error


@router.post("/lecturers")
def create_lecturer(payload: LecturerInput, semester_id: int | None = Query(None), db: Session = Depends(get_db)):
    if not payload.name.strip():
        raise HTTPException(422, "Họ tên không được để trống.")
    cleaned_aliases = [a.strip() for a in payload.aliases if a.strip()] if payload.aliases else []
    item = Lecturer(
        code=payload.code.strip() or None if payload.code else None,
        canonical_name=payload.name.strip(),
        department=payload.department,
        email=payload.email,
        note=payload.note,
        status=payload.status,
        aliases=cleaned_aliases,
        confirmed=True,
    )
    db.add(item)
    db.flush()
    for alias_str in cleaned_aliases:
        confirm_alias(db, item, alias_str)
    if semester_id:
        part_status = payload.participation_status or "ACTIVE"
        db.add(LecturerSemesterProfile(
            lecturer_id=item.id,
            semester_id=semester_id,
            participation_status=part_status,
        ))
    commit(db)
    return {"id": item.id, "name": item.canonical_name, "code": item.code, "status": item.status}


@router.patch("/lecturers/{lecturer_id}")
def edit_lecturer(lecturer_id: int, payload: LecturerInput, semester_id: int | None = Query(None), db: Session = Depends(get_db)):
    item = require_lecturer(db, lecturer_id)
    if not payload.name.strip():
        raise HTTPException(422, "Họ tên không được để trống.")
    item.canonical_name = payload.name.strip()
    item.code = payload.code.strip() or None if payload.code else None
    for field in ("department", "email", "note", "status"):
        setattr(item, field, getattr(payload, field))
    if payload.aliases is not None:
        cleaned_aliases = [a.strip() for a in payload.aliases if a.strip()]
        for alias_str in cleaned_aliases:
            confirm_alias(db, item, alias_str)
    if semester_id and payload.participation_status:
        profile = db.scalar(select(LecturerSemesterProfile).where(
            LecturerSemesterProfile.lecturer_id == item.id,
            LecturerSemesterProfile.semester_id == semester_id,
        ))
        if profile:
            profile.participation_status = payload.participation_status
        else:
            db.add(LecturerSemesterProfile(
                lecturer_id=item.id,
                semester_id=semester_id,
                participation_status=payload.participation_status,
            ))
    commit(db)
    return {"id": item.id, "name": item.canonical_name, "status": item.status}


@router.get("/lecturers/{lecturer_id}/dependencies")
def get_dependencies(lecturer_id: int, db: Session = Depends(get_db)):
    require_lecturer(db, lecturer_id)
    counts = dependencies(db, lecturer_id)
    return {"counts": counts, "can_delete": not any(counts.values())}


@router.delete("/lecturers/{lecturer_id}")
def delete_lecturer(lecturer_id: int, confirmed: bool = Query(False), db: Session = Depends(get_db)):
    item = require_lecturer(db, lecturer_id)
    if not confirmed:
        raise HTTPException(422, "Cần xác nhận xóa giảng viên.")
    counts = dependencies(db, lecturer_id)
    if any(counts.values()):
        raise HTTPException(409, {"code": "LECTURER_REFERENCED", "message": "Giảng viên có dữ liệu liên quan. Hãy ngừng sử dụng.", "dependencies": counts})
    db.delete(item)
    commit(db)
    return {"deleted": True}


@router.get("/lecturers/{lecturer_id}/profile")
def lecturer_profile(lecturer_id: int, semester_id: int = Query(...), db: Session = Depends(get_db)):
    item = require_lecturer(db, lecturer_id)
    if not db.get(Semester, semester_id):
        raise HTTPException(404, "Không tìm thấy kỳ học.")
    profile = db.scalar(select(LecturerSemesterProfile).where(LecturerSemesterProfile.lecturer_id == lecturer_id, LecturerSemesterProfile.semester_id == semester_id))
    return {
        "id": item.id, "code": item.code, "name": item.canonical_name, "status": item.status,
        "department": item.department, "email": item.email, "note": item.note,
        "aliases": item.aliases or [],
        "capabilities": [{"course_id": c.course_id, "allowed": c.allowed, "confirmed": c.confirmed} for c in db.scalars(select(LecturerCourseCapability).where(LecturerCourseCapability.lecturer_id == lecturer_id))],
        "participation": {field: getattr(profile, field) for field in ParticipationInput.model_fields} if profile else {"participation_status": "ACTIVE"},
        "dependencies": dependencies(db, lecturer_id),
    }


@router.put("/lecturers/{lecturer_id}/participation")
def participation(lecturer_id: int, payload: ParticipationInput, semester_id: int = Query(...), db: Session = Depends(get_db)):
    require_lecturer(db, lecturer_id)
    if not db.get(Semester, semester_id):
        raise HTTPException(404, "Không tìm thấy kỳ học.")
    if payload.min_workload is not None and payload.max_workload is not None and payload.min_workload > payload.max_workload:
        raise HTTPException(422, "Giới hạn tối thiểu vượt tối đa.")
    item = db.scalar(select(LecturerSemesterProfile).where(LecturerSemesterProfile.lecturer_id == lecturer_id, LecturerSemesterProfile.semester_id == semester_id))
    if item is None:
        item = LecturerSemesterProfile(lecturer_id=lecturer_id, semester_id=semester_id)
        db.add(item)
    for field, value in payload.model_dump().items():
        setattr(item, field, value)
    commit(db)
    return {"saved": True}


@router.put("/lecturers/{lecturer_id}/capabilities")
def capability(lecturer_id: int, payload: CapabilityInput, db: Session = Depends(get_db)):
    require_lecturer(db, lecturer_id)
    if not db.get(Course, payload.course_id):
        raise HTTPException(404, "Không tìm thấy môn học.")
    item = db.scalar(select(LecturerCourseCapability).where(LecturerCourseCapability.lecturer_id == lecturer_id, LecturerCourseCapability.course_id == payload.course_id))
    if item is None:
        item = LecturerCourseCapability(lecturer_id=lecturer_id, course_id=payload.course_id)
        db.add(item)
    item.allowed, item.confirmed, item.source = payload.allowed, payload.confirmed, "MANUAL"
    commit(db)
    return {"saved": True}


@router.post("/lecturers/{lecturer_id}/merge")
def merge(lecturer_id: int, payload: MergeInput, db: Session = Depends(get_db)):
    source, target = require_lecturer(db, lecturer_id), require_lecturer(db, payload.target_id)
    if not payload.confirmed:
        return {"preview": True, "source_id": source.id, "target_id": target.id, "affected": dependencies(db, source.id)}
    try:
        result = merge_lecturers(db, source, target)
        commit(db)
    except ValueError as error:
        db.rollback()
        raise HTTPException(409, str(error)) from error
    return {"merged": True, **result}


@router.patch("/identity/{kind}/{record_id}")
def correct_identity(kind: Literal["draft", "history"], record_id: int, payload: IdentityCorrection, semester_id: int = Query(...), db: Session = Depends(get_db)):
    model = NormalizedPreferenceDraft if kind == "draft" else HistoricalEvidence
    item = db.get(model, record_id)
    if not item or item.semester_id != semester_id:
        raise HTTPException(404, "Không tìm thấy bản ghi trong kỳ học.")
    if kind == "draft" and (item.applied_constraint_id or item.applied_seminar_id):
        raise HTTPException(409, "Bản nháp đã áp dụng; cần chỉnh ràng buộc hiện tại riêng.")
    lecturer = require_lecturer(db, payload.lecturer_id) if payload.lecturer_id else None
    if kind == 'draft' and lecturer and not payload.ignore:
        try:
            affected=relink_source_identity(db,item,lecturer,semester_id)
            if payload.confirm_alias:
                confirm_alias(db,lecturer,item.lecturer_alias or item.lecturer_code or '',payload.replace_alias)
            commit(db)
        except ValueError as error:
            db.rollback()
            raise HTTPException(409,str(error)) from error
        return {'id':item.id,'lecturer_id':item.lecturer_id,'status':item.status,'affected_draft_ids':affected}
    try:
        if payload.confirm_alias and lecturer:
            confirm_alias(db, lecturer, item.lecturer_alias if kind == "draft" else item.source_text, payload.replace_alias)
    except ValueError as error:
        db.rollback()
        raise HTTPException(409, str(error)) from error
    item.lecturer_id = lecturer.id if lecturer else None
    if kind == "draft":
        item.context_confirmed = True
        item.target = {**(item.target or {}), "identity_confirmed": True}
        item.status = "REJECTED" if payload.ignore else "NEEDS_REVIEW"
        item.needs_review = True
        item.review_reason = "Danh tính đã chỉnh; cần duyệt lại nội dung nguyện vọng."
    else:
        item.human_confirmed = True
        item.status = "IGNORED" if payload.ignore else "HUMAN_CONFIRMED" if lecturer else "NEEDS_REVIEW"
    commit(db)
    return {"id": item.id, "lecturer_id": item.lecturer_id, "status": item.status}


@router.get("/lecturers/review")
def review_lecturers(semester_id: int = Query(...), db: Session = Depends(get_db)):
    semester = db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(404, "Không tìm thấy kỳ học.")

    schedule_lecturers = set(db.scalars(
        select(ClassSection.assigned_lecturer_id).where(
            ClassSection.semester_id == semester_id,
            ClassSection.assigned_lecturer_id.is_not(None),
        )
    ))
    preference_lecturers = set(db.scalars(
        select(NormalizedPreferenceDraft.lecturer_id).where(
            NormalizedPreferenceDraft.semester_id == semester_id,
            or_(NormalizedPreferenceDraft.source_version_id.is_(None), NormalizedPreferenceDraft.source_version_id == semester.active_preference_source_id),
            NormalizedPreferenceDraft.lecturer_id.is_not(None),
        )
    ))
    history_lecturers = set(db.scalars(
        select(HistoricalEvidence.lecturer_id).where(
            HistoricalEvidence.semester_id == semester_id,
            HistoricalEvidence.lecturer_id.is_not(None),
        )
    ))

    profiles = {
        p.lecturer_id: p
        for p in db.scalars(select(LecturerSemesterProfile).where(LecturerSemesterProfile.semester_id == semester_id))
    }
    relevant_ids = schedule_lecturers | preference_lecturers | history_lecturers | set(profiles)
    all_lecturers = db.scalars(select(Lecturer).where(Lecturer.id.in_(relevant_ids)).order_by(Lecturer.id)).all()

    confirmed_caps = defaultdict(int)
    for c in db.scalars(select(LecturerCourseCapability).where(LecturerCourseCapability.allowed.is_(True), LecturerCourseCapability.confirmed.is_(True))):
        confirmed_caps[c.lecturer_id] += 1

    suggested_caps = defaultdict(int)
    for h in db.scalars(select(HistoricalEvidence).where(HistoricalEvidence.semester_id == semester_id, HistoricalEvidence.lecturer_id.is_not(None))):
        course_code = (h.evidence or {}).get("course_code")
        if course_code:
            suggested_caps[h.lecturer_id] += 1

    unlinked_drafts = db.scalars(
        select(NormalizedPreferenceDraft).where(
            NormalizedPreferenceDraft.semester_id == semester_id,
            or_(NormalizedPreferenceDraft.source_version_id.is_(None), NormalizedPreferenceDraft.source_version_id == semester.active_preference_source_id),
            NormalizedPreferenceDraft.lecturer_id.is_(None),
            NormalizedPreferenceDraft.draft_kind == "CONSTRAINT",
            NormalizedPreferenceDraft.status != "REJECTED",
        )
    ).all()

    norm_names = defaultdict(list)
    for l in all_lecturers:
        norm_names[normalize_identity(l.canonical_name)].append(l.id)
    human_linked = {d.lecturer_id for d in db.scalars(select(NormalizedPreferenceDraft).where(NormalizedPreferenceDraft.semester_id==semester_id)) if (d.target or {}).get('identity_confirmed')}

    lec_courses = defaultdict(list)
    lec_capabilities = defaultdict(list)
    for cap, crs in db.execute(
        select(LecturerCourseCapability, Course)
        .join(Course, Course.id == LecturerCourseCapability.course_id)
        .order_by(Course.name, Course.code)
    ):
        lec_capabilities[cap.lecturer_id].append({
            "course_id": crs.id, "course_code": crs.code, "course_name": crs.name,
            "allowed": cap.allowed, "confirmed": cap.confirmed, "source": cap.source,
        })
        if cap.allowed and cap.confirmed:
            lec_courses[cap.lecturer_id].append(crs.code)

    lec_aliases = defaultdict(list)
    for a in db.scalars(select(LecturerAlias)):
        lec_aliases[a.lecturer_id].append(a.alias_text)

    lecturer_items = []
    matched_count = 0
    needs_conf_count = 0

    for l in all_lecturers:
        sources = []
        if l.id in schedule_lecturers:
            sources.append("Lịch học")
        if l.id in preference_lecturers:
            sources.append("Nguyện vọng")
        if l.id in history_lecturers:
            sources.append("Lịch sử kỳ trước")
        if not sources:
            sources.append("Lecturer Master")
        source_text = " + ".join(sources)

        is_dup = len(norm_names[normalize_identity(l.canonical_name)]) > 1
        if l.id in human_linked or (l.id in schedule_lecturers and l.code):
            id_status='STANDARDIZED';id_label='Đã khớp';matched_count+=1
        elif is_dup:
            id_status = "POSSIBLE_DUPLICATE"
            id_label = "Có thể trùng"
            needs_conf_count += 1
        elif not l.code:
            id_status = "NEEDS_CONFIRMATION"
            id_label = "Cần xác nhận"
            needs_conf_count += 1
        elif not l.confirmed:
            id_status = "NEW_CANDIDATE"
            id_label = "Giảng viên mới"
            needs_conf_count += 1
        else:
            id_status = "STANDARDIZED"
            id_label = "Đã khớp"
            matched_count += 1

        profile = profiles.get(l.id)
        part_status = profile.participation_status if profile else "ACTIVE"
        part_label = {
            "ACTIVE": "Đang công tác",
            "ON_LEAVE": "Nghỉ phép / Đi học",
            "NOT_PARTICIPATING": "Không tham gia kỳ",
        }.get(part_status, "Đang công tác")

        c_count = confirmed_caps[l.id]
        s_count = suggested_caps[l.id]
        if c_count > 0 and s_count > 0:
            cap_summary = f"{c_count} môn đã xác nhận ({s_count} gợi ý)"
        elif c_count > 0:
            cap_summary = f"{c_count} môn đã xác nhận"
        elif s_count > 0:
            cap_summary = f"{s_count} môn gợi ý"
        else:
            cap_summary = "0 môn"

        aliases_list = sorted(set([*(l.aliases or []), *lec_aliases[l.id]]))
        sources_list = sources if sources else ["Lecturer Master"]

        lecturer_items.append({
            "id": l.id,
            "code": l.code,
            "name": l.canonical_name,
            "email": l.email,
            "department": l.department,
            "quota_min": profile.min_workload if profile and profile.min_workload is not None else 0,
            "quota_max": profile.max_workload if profile and profile.max_workload is not None else 24,
            "aliases": aliases_list,
            "sources": sources_list,
            "source": source_text,
            "identity_status": id_status,
            "identity_status_label": id_label,
            "participation_status": part_status,
            "participates": part_status == "ACTIVE",
            "participation_status_label": part_label,
            "capabilities_count": c_count,
            "courses_can_teach": lec_courses[l.id],
            "capabilities": lec_capabilities[l.id],
            "suggested_capabilities_count": s_count,
            "capability_summary": cap_summary,
            "note": l.note,
            "status": l.status,
            "active": l.status == "ACTIVE",
            "has_blocking_conflict": is_dup,
        })

    unresolved_items = []
    blocker_messages = []
    seen_unresolved_aliases = set()
    for d in unlinked_drafts:
        alias_text = d.lecturer_alias or d.raw_text or f"Dòng {d.source_row}"
        if alias_text not in seen_unresolved_aliases:
            seen_unresolved_aliases.add(alias_text)
            blocker_messages.append(f"Nguyện vọng '{alias_text}' ({d.source_file}) chưa gắn với giảng viên cụ thể.")
        unresolved_items.append({
            "id": d.id,
            "source_text": alias_text,
            "source_file": d.source_file,
            "source_cell": d.source_cell,
            "raw_text": d.raw_text,
            "constraint_type": d.constraint_type,
        })

    for item in lecturer_items:
        if not item["code"] and item['identity_status']!='STANDARDIZED':
            msg = f"Giảng viên '{item['name']}' chưa có mã GV chuẩn."
            if msg not in blocker_messages:
                blocker_messages.append(msg)
        elif item["identity_status"] == "POSSIBLE_DUPLICATE":
            msg = f"Giảng viên '{item['name']}' ({item['code']}) có nghi ngờ trùng lặp tên."
            if msg not in blocker_messages:
                blocker_messages.append(msg)

    blockers = len(blocker_messages)
    can_continue = blockers == 0

    return {
        "semester_id": semester_id,
        "semester_name": semester.name,
        "total_lecturers": len(lecturer_items),
        "ready_count": matched_count,
        "needs_review_count": needs_conf_count + len(seen_unresolved_aliases),
        "blockers_count": blockers,
        "blockers": blocker_messages,
        "items": lecturer_items,
        "lecturers": lecturer_items,
        "unresolved_drafts": unresolved_items,
        "readiness": {
            "total_lecturers": len(lecturer_items),
            "matched": matched_count,
            "needs_confirmation": needs_conf_count + len(seen_unresolved_aliases),
            "blockers": blockers,
            "can_continue": can_continue,
            "blocker_messages": blocker_messages,
        },
    }


def _parse_import_workbook(file: UploadFile) -> list[list]:
    file.file.seek(0)
    try:
        wb = openpyxl.load_workbook(file.file, data_only=True)
        sheet = wb.active
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    except Exception:
        import xlrd
        file.file.seek(0)
        content = file.file.read()
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(r) for r in range(sheet.nrows)]


@router.post("/lecturers/import/preview")
def import_lecturers_preview(
    file: UploadFile = File(...),
    semester_id: int = Query(...),
    db: Session = Depends(get_db),
):
    semester = db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(404, "Không tìm thấy kỳ học.")

    try:
        rows = _parse_import_workbook(file)
    except Exception as error:
        raise HTTPException(422, {"code": "INVALID_LECTURER_WORKBOOK", "message": "Không đọc được workbook giảng viên."}) from error
    if not rows:
        raise HTTPException(422, "File Excel trống.")

    headers = [normalize_identity(str(cell or "")) for cell in rows[0]]
    col_alias = next((i for i, h in enumerate(headers) if any(k in h for k in ("alias", "danh xung", "ten thuong goi", "biet danh"))), None)
    col_code = next((i for i, h in enumerate(headers) if any(k in h for k in ("ma gv", "ma giang vien", "magv", "code"))), None)
    col_name = next((i for i, h in enumerate(headers) if i != col_alias and any(k in h for k in ("ho va ten", "ho ten", "canonical", "tengv", "ten", "name"))), None)
    col_part = next((i for i, h in enumerate(headers) if any(k in h for k in ("tham gia", "participation"))), None)
    col_status = next((i for i, h in enumerate(headers) if any(k in h for k in ("trang thai", "status"))), None)

    if col_name is None and col_code is None:
        raise HTTPException(422, "Không tìm thấy cột Mã giảng viên hoặc Tên giảng viên.")

    all_lecturers = db.scalars(select(Lecturer)).all()
    by_code = {l.code: l for l in all_lecturers if l.code}
    by_normalized_name = defaultdict(list)
    for l in all_lecturers:
        by_normalized_name[normalize_identity(l.canonical_name)].append(l)

    confirmed_aliases = {
        a.normalized_alias: db.get(Lecturer, a.lecturer_id)
        for a in db.scalars(select(LecturerAlias).where(LecturerAlias.confirmed.is_(True)))
    }

    to_add = []
    to_update = []
    needs_confirmation = []
    possible_duplicates = []
    skipped = []

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        code_val = str(row[col_code] or "").strip() if col_code is not None and col_code < len(row) and row[col_code] is not None else None
        name_val = str(row[col_name] or "").strip() if col_name is not None and col_name < len(row) and row[col_name] is not None else ""
        alias_val = str(row[col_alias] or "").strip() if col_alias is not None and col_alias < len(row) and row[col_alias] is not None else None
        part_val = normalize_identity(str(row[col_part] or "")).upper() if col_part is not None and col_part < len(row) and row[col_part] is not None else "ACTIVE"
        status_val = normalize_identity(str(row[col_status] or "")).upper() if col_status is not None and col_status < len(row) and row[col_status] is not None else "ACTIVE"

        if not name_val and not code_val:
            continue
        if not name_val:
            skipped.append({
                "row_index": row_idx,
                "code": code_val,
                "name": "",
                "alias": alias_val,
                "participation_status": "ACTIVE",
                "status": "ACTIVE",
                "reason": "Thiếu thông tin họ và tên giảng viên.",
            })
            continue

        norm_name = normalize_identity(name_val)
        norm_alias = normalize_identity(alias_val) if alias_val else None

        matched_lecturer = by_code.get(code_val) if code_val else None
        match_type = "EXACT_CODE" if matched_lecturer else None

        if not matched_lecturer and norm_alias and norm_alias in confirmed_aliases:
            matched_lecturer = confirmed_aliases[norm_alias]
            match_type = "CONFIRMED_ALIAS"
        if not matched_lecturer and norm_name in confirmed_aliases:
            matched_lecturer = confirmed_aliases[norm_name]
            match_type = "CONFIRMED_ALIAS"

        if not matched_lecturer:
            exact_matches = by_normalized_name.get(norm_name, [])
            if len(exact_matches) == 1:
                matched_lecturer = exact_matches[0]
                match_type = "EXACT_NAME"
            elif len(exact_matches) > 1:
                match_type = "POSSIBLE_DUPLICATE"

        parsed_part = "ON_LEAVE" if any(k in part_val for k in ("NGHI", "LEAVE", "PHEP")) else "NOT_PARTICIPATING" if any(k in part_val for k in ("KHONG", "NOT")) else "ACTIVE"
        parsed_status = "INACTIVE" if any(k in status_val for k in ("INACTIVE", "NGUNG")) else "ACTIVE"

        row_data = {
            "row_index": row_idx,
            "code": code_val,
            "name": name_val,
            "alias": alias_val,
            "participation_status": parsed_part,
            "status": parsed_status,
        }

        conflicting_code = matched_lecturer and code_val and matched_lecturer.code and matched_lecturer.code != code_val
        if conflicting_code:
            needs_confirmation.append({**row_data, "reason": "Mã mới khác mã của danh tính đã khớp; cần xác nhận."})
        elif match_type == "POSSIBLE_DUPLICATE":
            possible_duplicates.append({**row_data, "reason": "Nhiều giảng viên có cùng tên chuẩn hóa trong hệ thống."})
        elif matched_lecturer:
            to_update.append({**row_data, "target_id": matched_lecturer.id, "existing_name": matched_lecturer.canonical_name})
        else:
            to_add.append(row_data)

    total_rows = len(to_add) + len(to_update) + len(needs_confirmation) + len(possible_duplicates) + len(skipped)
    return {
        "total_rows": total_rows,
        "can_commit": not (possible_duplicates or needs_confirmation or skipped),
        "to_add": to_add,
        "to_update": to_update,
        "needs_confirmation": needs_confirmation,
        "possible_duplicates": possible_duplicates,
        "skipped": skipped,
        "summary": {
            "total_rows": total_rows,
            "to_add": len(to_add),
            "to_update": len(to_update),
            "needs_confirmation": len(needs_confirmation),
            "possible_duplicates": len(possible_duplicates),
            "skipped": len(skipped),
        },
    }


@router.post("/lecturers/import")
def commit_lecturer_import(
    file: UploadFile = File(...),
    semester_id: int = Query(...),
    db: Session = Depends(get_db),
):
    semester = db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(404, "Không tìm thấy kỳ học.")

    preview = import_lecturers_preview(file, semester_id, db)
    if not preview["can_commit"]:
        raise HTTPException(422, {"code": "LECTURER_IMPORT_REQUIRES_REVIEW", "message": "Cần xử lý toàn bộ dòng chưa rõ trước khi nhập.", "summary": preview["summary"]})

    try:
        created_count = 0
        updated_count = 0

        for item in preview["to_add"]:
            aliases = [item["alias"]] if item.get("alias") else []
            lecturer = Lecturer(
                code=item["code"],
                canonical_name=item["name"],
                aliases=aliases,
                status=item["status"],
                confirmed=bool(item["code"]),
                source_file=file.filename or "IMPORT_XLSX",
            )
            db.add(lecturer)
            db.flush()
            if item.get("alias"):
                confirm_alias(db, lecturer, item["alias"])
            db.add(LecturerSemesterProfile(
                lecturer_id=lecturer.id,
                semester_id=semester_id,
                participation_status=item["participation_status"],
            ))
            created_count += 1

        for item in preview["to_update"]:
            lecturer = db.get(Lecturer, item["target_id"])
            if lecturer:
                if item["code"] and not lecturer.code:
                    lecturer.code = item["code"]
                if item.get("alias"):
                    confirm_alias(db, lecturer, item["alias"])
                profile = db.scalar(select(LecturerSemesterProfile).where(
                    LecturerSemesterProfile.lecturer_id == lecturer.id,
                    LecturerSemesterProfile.semester_id == semester_id,
                ))
                if profile:
                    profile.participation_status = item["participation_status"]
                else:
                    db.add(LecturerSemesterProfile(
                        lecturer_id=lecturer.id,
                        semester_id=semester_id,
                        participation_status=item["participation_status"],
                    ))
                updated_count += 1

        commit(db)
        return {
            "success": True,
            "created": created_count,
            "added": created_count,
            "updated": updated_count,
            "summary": preview["summary"],
        }
    except Exception as error:
        db.rollback()
        raise HTTPException(422, f"Lỗi nhập danh sách giảng viên: {error}") from error


@router.get("/lecturers/export")
def export_lecturers(
    semester_id: int = Query(...),
    db: Session = Depends(get_db),
):
    semester = db.get(Semester, semester_id)
    if not semester:
        raise HTTPException(404, "Không tìm thấy kỳ học.")

    review_data = review_lecturers(semester_id, db)
    lecturers = review_data["lecturers"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách giảng viên"

    headers = [
        "Mã GV", "Tên GV", "Aliases", "Trạng thái hoạt động", "Tham gia kỳ này",
        "Trạng thái nhận diện", "Năng lực giảng dạy", "Nguồn dữ liệu", "Ghi chú",
    ]
    ws.append(headers)

    for l in lecturers:
        ws.append([
            l["code"] or "",
            l["name"],
            ", ".join(l["aliases"]) if l["aliases"] else "",
            l["status"],
            l["participation_status_label"],
            l["identity_status_label"],
            l["capability_summary"],
            l["source"],
            l["note"] or "",
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"danh-sach-giang-vien-ky-{semester_id}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
