from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl

from app.parsers.preferences import ParsedPreference, interpret_raw_preference


# Rows map the five sessions in the form to the 15-period solver calendar.
PERIOD_MAP = {14: (1, 3), 15: (4, 6), 16: (7, 9), 17: (10, 12), 18: (13, 15)}
DAY_MAP = {3: (2, "T2"), 4: (3, "T3"), 5: (4, "T4"), 6: (5, "T5"), 7: (6, "T6"), 8: (7, "T7"), 9: (8, "CN")}
FORM_TAG = "[HUCE-FORM-PREF-V3.0]"


def _text(value: object) -> str:
    return str(value or "").strip()


def _plain(value: object) -> str:
    normalized = unicodedata.normalize("NFD", _text(value).casefold())
    return "".join(char for char in normalized if unicodedata.category(char) != "Mn").replace("đ", "d")


def _is_form_sheet(ws) -> bool:
    c1 = _text(ws.cell(1, 1).value)
    c2 = _text(ws.cell(2, 1).value)
    return (
        FORM_TAG in c1
        or "phieu nguyen vong" in _plain(c1)
        or "phieu nguyen vong" in _plain(c2)
        or "Phieu_Nguyen_Vong" in ws.title
    )


def _title_alias(sheet_title: str) -> str:
    """Return a human-readable alias from a numbered individual form sheet."""
    cleaned = re.sub(r"^\s*\d+[\s_\-.]+\s*", "", sheet_title).strip()
    return cleaned.replace("_", " ")


def _is_help_text(value: object) -> bool:
    """Do not mistake in-cell instructions for an answer selected by the lecturer."""
    return _plain(value).startswith(("bo trong", "chon ", "nhap ", "vi du", "vd"))


def _answer_cell(ws, row: int) -> tuple[str, str]:
    """Read either supported answer column; V3.0 files use E or F by version."""
    for column in (5, 6):
        value = ws.cell(row, column).value
        if _text(value) and not _is_help_text(value):
            return _text(value), f"{openpyxl.utils.get_column_letter(column)}{row}"
    return "", f"E{row}:F{row}"


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(_text(value), fmt).date()
        except ValueError:
            pass
    return None


def is_grid_preference_template(path: Path) -> bool:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return any(_is_form_sheet(workbook[name]) for name in workbook.sheetnames)
    except Exception:
        return False


def parse_grid_preferences(path: Path, *, semester_start: date | None = None,
                           semester_end: date | None = None) -> list[ParsedPreference]:
    """Parse every V3.0 individual form deterministically.

    A worksheet title is only an identity hint. If the form has neither a
    lecturer code nor a full name, each parsed rule stays review-required until
    the import workflow resolves that hint in the selected semester.
    """
    workbook = openpyxl.load_workbook(path, data_only=True)
    drafts: list[ParsedPreference] = []

    summary_identities = {}
    summary_sheet_name = next((s for s in workbook.sheetnames if _plain(s) in {"tong hop", "tong_hop", "summary"}), None)
    if summary_sheet_name:
        summary_sheet = workbook[summary_sheet_name]
        for row in range(2, summary_sheet.max_row + 1):
            stt = _text(summary_sheet.cell(row, 1).value)
            name = _text(summary_sheet.cell(row, 2).value)
            if stt and name:
                summary_identities[stt] = name
                summary_identities[stt.lstrip("0") or "0"] = name

    for worksheet in (workbook[name] for name in workbook.sheetnames):
        if not _is_form_sheet(worksheet):
            continue

        lecturer_code = _text(worksheet.cell(8, 1).value)
        entered_name = _text(worksheet.cell(8, 5).value)

        fallback_name = None
        match = re.match(r"^\s*0*(\d+)[\s_\-.]*", worksheet.title)
        if match and match.group(1) in summary_identities:
            fallback_name = summary_identities[match.group(1)]

        canonical_name = entered_name or fallback_name or _title_alias(worksheet.title)
        if not lecturer_code and not canonical_name:
            continue

        title_identity_only = not lecturer_code and not entered_name and not fallback_name
        review_reason = (
            "Chưa có mã hoặc họ tên giảng viên trong biểu mẫu; tên sheet chỉ được dùng làm gợi ý nhận diện."
            if title_identity_only else None
        )
        status = "NEEDS_REVIEW" if title_identity_only else "CONFIRMED"

        def draft(**values) -> ParsedPreference:
            return ParsedPreference(
                lecturer_alias=canonical_name, canonical_name=canonical_name, lecturer_code=lecturer_code or None,
                confidence=1.0, confidence_label="HIGH", needs_review=title_identity_only,
                review_reason=review_reason, status=status, source_sheet=worksheet.title,
                context_type="TEACHING", **values,
            )

        lecturer_drafts_count = 0

        # Weekly availability and period preferences.
        for row, (start_period, end_period) in PERIOD_MAP.items():
            for column, (weekday, day_scope) in DAY_MAP.items():
                value = _plain(worksheet.cell(row, column).value)
                if not value:
                    continue
                target = {"weekday": weekday, "start_period": start_period, "end_period": end_period,
                          "periods": list(range(start_period, end_period + 1))}
                cell = f"{openpyxl.utils.get_column_letter(column)}{row}"
                if "ban cung" in value or "xin nghi" in value:
                    drafts.append(draft(
                        weekday=weekday, day_scope=day_scope, start_period=start_period, end_period=end_period,
                        constraint_type="UNAVAILABLE", hardness="hard", weight=1.0, target=target,
                        raw_text=f"{day_scope} tiết {start_period}-{end_period}: Bận cứng", source_cell=cell, source_row=row,
                    ))
                    lecturer_drafts_count += 1
                elif "uu tien" in value:
                    drafts.append(draft(
                        weekday=weekday, day_scope=day_scope, start_period=start_period, end_period=end_period,
                        constraint_type="PREFERRED_PERIOD", hardness="soft", weight=0.8, target=target,
                        raw_text=f"{day_scope} tiết {start_period}-{end_period}: Ưu tiên dạy", source_cell=cell, source_row=row,
                    ))
                    lecturer_drafts_count += 1
                elif "han che" in value or "tranh" in value:
                    drafts.append(draft(
                        weekday=weekday, day_scope=day_scope, start_period=start_period, end_period=end_period,
                        constraint_type="AVOID_PERIOD", hardness="soft", weight=0.45, target=target,
                        raw_text=f"{day_scope} tiết {start_period}-{end_period}: Hạn chế nếu có thể", source_cell=cell, source_row=row,
                    ))
                    lecturer_drafts_count += 1

        # The scheduling-style fields are opt-in. Guidance such as “Chọn Rất
        # mong muốn ...” is never converted to a constraint.
        max_days, cell = _answer_cell(worksheet, 21)
        match = re.search(r"\b([1-5])\b", _plain(max_days))
        if match:
            number = int(match.group(1))
            drafts.append(draft(
                weekday=None, day_scope="ALL_WEEKDAYS", constraint_type="MAX_DAYS_PER_WEEK",
                hardness="soft", weight={2: 0.85, 3: 0.75, 4: 0.65, 5: 0.5}.get(number, 0.65),
                numeric_value=number, target={"max": number, "value": number},
                raw_text=f"Tối đa {number} ngày lên trường mỗi tuần", source_cell=cell, source_row=21,
            ))
            lecturer_drafts_count += 1

        consecutive, cell = _answer_cell(worksheet, 22)
        if _plain(consecutive) == "rat mong muon":
            drafts.append(draft(
                weekday=None, day_scope="ALL_WEEKDAYS", constraint_type="PREFER_COMPACT_SCHEDULE",
                hardness="soft", weight=0.7, numeric_value=6, target={"numeric_value": 6, "value": 6},
                raw_text="Rất mong muốn lịch dạy gọn, hạn chế các tiết trống xen giữa trong một ngày",
                source_cell=cell, source_row=22,
            ))
            lecturer_drafts_count += 1

        free_morning, cell = _answer_cell(worksheet, 23)
        if _plain(free_morning) == "co":
            drafts.append(draft(
                weekday=None, day_scope="ALL_WEEKDAYS", constraint_type="MIN_FREE_MORNING_PER_WEEK",
                hardness="soft", weight=0.65, numeric_value=1, target={"numeric_value": 1, "value": 1},
                raw_text="Mong muốn có ít nhất một buổi sáng trống mỗi tuần", source_cell=cell, source_row=23,
            ))
            lecturer_drafts_count += 1

        # Date exceptions cover the complete 15-period day. They are hard only
        # when the lecturer explicitly selected a mandatory level.
        for row in range(28, 31):
            start_date, end_date = _parse_date(worksheet.cell(row, 2).value), _parse_date(worksheet.cell(row, 3).value)
            if not start_date or not end_date:
                continue
            reason, level = _text(worksheet.cell(row, 4).value), _plain(worksheet.cell(row, 5).value)
            hard = "bat buoc" in level or "hard" in level
            drafts.append(draft(
                weekday=None, day_scope="ALL_DAYS", start_period=1, end_period=15, start_date=start_date, end_date=end_date,
                constraint_type="UNAVAILABLE", hardness="hard" if hard else "soft", weight=1.0,
                target={"start_date": start_date.isoformat(), "end_date": end_date.isoformat(), "periods": list(range(1, 16))},
                raw_text=f"Bận từ {start_date.isoformat()} đến {end_date.isoformat()}: {reason}".rstrip(": "),
                source_cell=f"B{row}:E{row}", source_row=row,
            ))
            lecturer_drafts_count += 1

        # Interpret the free-text once into atomic rules.  The original text is
        # still retained below as provenance, but it never reaches the solver.
        original_wish = _text(worksheet.cell(25, 1).value)
        if original_wish:
            structured = interpret_raw_preference(original_wish, semester_start=semester_start, semester_end=semester_end)
            for index, semantic in enumerate(structured, 1):
                target = semantic["target"]
                periods = target.get("periods") or []
                drafts.append(draft(
                    weekday=target.get("weekday"), day_scope=target.get("day_scope"),
                    start_period=min(periods) if periods else None, end_period=max(periods) if periods else None,
                    start_date=_parse_date(target.get("start_date")), end_date=_parse_date(target.get("end_date")),
                    constraint_type=semantic["constraint_type"], hardness=semantic["hardness"], weight=semantic["weight"],
                    target=target, raw_text=original_wish, source_cell=f"A25#{index}", source_row=25,
                ))
                lecturer_drafts_count += 1
            drafts.append(ParsedPreference(
                lecturer_alias=canonical_name, canonical_name=canonical_name, lecturer_code=lecturer_code or None,
                weekday=None, day_scope=None, constraint_type="RAW_NOTE", target={
                    "source": "original_wish",
                    "interpreted_rules": [
                        {"constraint_type": rule["constraint_type"], "target": rule["target"],
                         "hardness": rule["hardness"], "weight": rule["weight"]}
                        for rule in structured
                    ],
                },
                raw_text=original_wish, confidence=1.0, confidence_label="HIGH", source_row=25,
                source_sheet=worksheet.title, source_cell="A25", hardness="soft", weight=0.0,
                needs_review=False,
                review_reason=("Đã diễn giải thành các quy tắc có cấu trúc ở cùng ô nguồn."
                               if structured else "Không có nguyện vọng đặc biệt hoặc chưa có semantics để tối ưu."),
                status="INTERPRETED" if structured else "INTERPRETED", context_type="TEACHING",
            ))
            lecturer_drafts_count += 1
            
        if lecturer_drafts_count == 0:
            drafts.append(ParsedPreference(
                lecturer_alias=canonical_name, canonical_name=canonical_name, lecturer_code=lecturer_code or None,
                weekday=None, day_scope=None, constraint_type="RAW_NOTE",
                target={"source": "empty_grid", "interpreted_rules": []},
                raw_text="Giảng viên không có nguyện vọng đặc biệt trên biểu mẫu",
                confidence=1.0, confidence_label="HIGH", source_row=12,
                source_sheet=worksheet.title, source_cell="A12", hardness="soft", weight=0.0,
                needs_review=False, review_reason="Không có nguyện vọng đặc biệt; không tạo ràng buộc solver.",
                status="INTERPRETED", context_type="TEACHING",
            ))

    # Free text is the authoritative semantic explanation when it clearly
    # qualifies a grid mark (for example “có thể nhận thêm 7-9” must not remain
    # an AVOID_PERIOD). Remove covered or contradictory grid rules first.
    semantic = [d for d in drafts if d.source_cell.startswith("A25#")]
    def covered(grid_item: ParsedPreference, raw_item: ParsedPreference) -> bool:
        def days(item):
            target=item.target or {}
            if target.get("weekdays"): return set(target["weekdays"])
            if target.get("weekday"): return {target["weekday"]}
            if item.day_scope == "ALL_WEEKDAYS": return set(range(2,7))
            if item.day_scope == "ALL_DAYS": return set(range(2,9))
            if item.day_scope and item.day_scope.startswith("T"): return {int(item.day_scope[1:])}
            return set(range(2,9))
        gp=set((grid_item.target or {}).get("periods") or [])
        rp=set((raw_item.target or {}).get("periods") or [])
        return days(grid_item) <= days(raw_item) and (not rp or gp <= rp)
    reconciled=[]
    for item in drafts:
        if not item.source_cell.startswith("A25#") and item.constraint_type != "RAW_NOTE":
            qualifying = next((raw for raw in semantic if covered(item, raw) and (
                raw.constraint_type == item.constraint_type or
                (item.constraint_type == "AVOID_PERIOD" and raw.constraint_type == "PREFERRED_PERIOD")
            )), None)
            if qualifying:
                continue
        reconciled.append(item)
    drafts=reconciled

    # One source cell can be represented by the grid and by its free-text
    # explanation.  Keep a single identical semantic rule so its weight is
    # never counted twice by the solver; RAW_NOTE provenance is intentionally
    # excluded from this de-duplication.
    unique: list[ParsedPreference] = []
    seen: set[tuple] = set()
    for item in drafts:
        if item.constraint_type == "RAW_NOTE":
            unique.append(item); continue
        key = (item.lecturer_code or item.lecturer_alias, item.constraint_type,
               tuple(sorted((item.target or {}).get("weekdays") or [])), item.day_scope,
               tuple((item.target or {}).get("periods") or []), item.start_date, item.end_date,
               item.hardness)
        if key not in seen:
            seen.add(key); unique.append(item)
    return unique
