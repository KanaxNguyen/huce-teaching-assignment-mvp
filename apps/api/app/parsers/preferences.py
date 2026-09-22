from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from app.parsers.schedule import clean_text

CONSTRAINT_TYPES = {
    "UNAVAILABLE", "AVOID_PERIOD", "PREFERRED_PERIOD", "PREFER_CONSECUTIVE_PERIODS",
    "MIN_CLASSES", "MAX_CLASSES", "MAX_SESSIONS_PER_DAY", "MAX_DAYS_PER_WEEK",
    "MIN_FREE_MORNING_PER_WEEK", "REQUIRED_ASSIGNMENT", "FORBIDDEN_ASSIGNMENT", "RAW_NOTE",
    "SEMINAR_COMMITMENT", "SEMINAR_NOTE",
    "PREFERRED_DAYS", "AVOID_DAYS", "PREFER_LOW_WORKLOAD", "PREFER_COMPACT_SCHEDULE",
}
PREFERENCE_WEIGHTS = {
    "STRONG": 0.80, "VERY_STRONG": 0.88, "MEDIUM": 0.60,
    "WEAK": 0.30, "DAY": 0.70, "AVOID_DAY": 0.50,
}
DAY_SCOPES = {"T2", "T3", "T4", "T5", "T6", "T7", "CN", "ALL_WEEKDAYS", "ALL_DAYS"}
HARDNESS_VALUES = {"SOFT", "HARD"}
STATUS_VALUES = {"DRAFT", "CONFIRMED", "NEEDS_REVIEW", "REJECTED"}
V2_PREFERENCE_HEADERS = [
    "Mã GV*", "Họ tên GV", "Ngữ cảnh*", "Loại ràng buộc*", "Thứ/Phạm vi*", "Tiết bắt đầu",
    "Tiết kết thúc", "Ngày bắt đầu", "Ngày kết thúc", "Độ cứng*", "Trọng số",
    "Giá trị số", "Ghi chú / Nguyên văn", "Trạng thái",
]


@dataclass
class ParsedPreference:
    lecturer_alias: str
    canonical_name: str
    weekday: int | None
    constraint_type: str
    target: dict
    raw_text: str
    confidence: float
    source_row: int
    lecturer_code: str | None = None
    source_sheet: str = ""
    source_cell: str = ""
    day_scope: str | None = None
    start_period: int | None = None
    end_period: int | None = None
    start_date: date | None = None
    end_date: date | None = None
    hardness: str = "soft"
    weight: float = 0.8
    numeric_value: float | None = None
    confidence_label: str = "LOW"
    needs_review: bool = True
    review_reason: str | None = None
    status: str = "NEEDS_REVIEW"
    context_type: str = "TEACHING"
    context_confidence: str = "HIGH"
    context_confirmed: bool = False
    seminar_link: str | None = None


@dataclass
class ParsedSharedSeminar:
    seminar_code: str
    name: str
    participant_codes: list[str]
    day_scopes: list[str]
    period_blocks: list[list[int]]
    hardness: str
    weight: float
    raw_text: str
    source_sheet: str
    source_row: int
    source_cell: str
    status: str
    needs_review: bool = False
    review_reason: str | None = None


@dataclass
class PreferenceParseResult:
    format: str
    drafts: list[ParsedPreference] = field(default_factory=list)
    seminars: list[ParsedSharedSeminar] = field(default_factory=list)
    invalid_rows: int = 0
    raw_clauses: int = 0
    dropped_clauses: int = 0
    warnings: list[dict] = field(default_factory=list)


def _plain(value: str) -> str:
    normalized = unicodedata.normalize("NFD", value.casefold())
    plain = "".join(char for char in normalized if unicodedata.category(char) != "Mn").replace("đ", "d")
    for word, number in {'hai':2,'ba':3,'tu':4,'nam':5,'sau':6,'bay':7}.items():
        plain = re.sub(rf'\bthu\s+{word}\b',f'thu {number}',plain)
    return re.sub(r'\s+',' ',plain).strip()


def _normalized_header(value: object) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize('NFC',clean_text(value))).strip().casefold()


def _uncertain(text: str, field: str) -> bool:
    subject={'period':r'(?:tiet|gio|thoi gian)', 'day':r'(?:ngay|thu)', 'identity':r'(?:giang vien|giao vien|gv)'}[field]
    uncertain=r'(?:chua\s+(?:xac dinh|ro|chot)|khong ro)'
    return bool(re.search(rf'{uncertain}\s+(?:duoc\s+)?{subject}|{subject}\s+{uncertain}',_plain(text)))


def _periods_from_text(text: str) -> list[int]:
    """Extract complete inclusive period ranges without treating weekdays as periods."""
    plain = _plain(text)
    if _uncertain(text,'period'):
        return []
    # Remove weekday ranges before extracting numeric period ranges.
    plain = re.sub(r"(?:thu|\bt)\s*[2-7]\s*(?:-|–|den)\s*(?:(?:thu|t)\s*)?[2-7]", "", plain)
    ranges=re.findall(r'(?:tiet\s*)?(\d{1,2})\s*[-–]\s*(\d{1,2})',plain)
    if any(not 1<=int(a)<=int(b)<=15 for a,b in ranges): return []
    open_end = re.search(r"tu\s+tiet\s*(\d{1,2})(?:\s+tro\s+di|\s*$)", plain)
    if open_end and 1 <= int(open_end.group(1)) <= 15:
        return list(range(int(open_end.group(1)), 16))
    periods: list[int] = []
    compact = re.search(r"(?:tiet|ca|block)\s*([1-9]{2,3})(?!\d)", plain)
    if compact and compact.group(1) in {"123", "456", "789"}:
        periods.extend(int(value) for value in compact.group(1))
    for start_text, end_text in re.findall(
        r"(?:tiet|ca|block)?\s*(\d{1,2})\s*(?:-|–|den|toi)\s*(?:tiet\s*)?(\d{1,2})", plain
    ):
        start, end = int(start_text), int(end_text)
        if 1 <= start <= end <= 15:
            periods.extend(range(start, end + 1))
    if periods:
        return sorted(set(periods))
    match = re.search(r"(?:tiet|ca|block)\s*(\d{1,2})(?!\s*(?:-|–|den|toi))", plain)
    if match and 1 <= int(match.group(1)) <= 15:
        return [int(match.group(1))]
    return []


def _weekday_from_text(text: str) -> int | None:
    plain = _plain(text)
    if _uncertain(text,'day'): return None
    if re.search(r"chu\s*nhat|\bcn\b", plain):
        return 8
    match = re.search(r"(?:thu|\bt)\s*([2-7])", plain)
    return int(match.group(1)) if match else None


def _day_scope(weekday: int | None, text: str = "") -> str | None:
    plain = _plain(text)
    if _uncertain(text,'day'): return None
    if re.search(r'tat ca\s+(?:cac\s+)?ngay|\bmoi ngay\b',plain):
        return "ALL_DAYS"
    if "cac buoi" in plain and "trong tuan" not in plain and weekday is None:
        return "ALL_DAYS"
    day_range = re.search(r"(?:thu|\bt)\s*([2-7])\s*(?:-|–|den)\s*(?:(?:thu|t)\s*)?([2-7])", plain)
    if day_range and day_range.groups() == ("2", "6"):
        return "ALL_WEEKDAYS"
    if any(token in plain for token in ("trong tuan", "cac buoi", "cac ngay", "t2 den t6")):
        return "ALL_WEEKDAYS"
    return "CN" if weekday == 8 else (f"T{weekday}" if weekday else None)


def _periods_for_part_of_day(text: str) -> list[int]:
    plain = _plain(text)
    if _uncertain(text,'period'):
        return []
    if "3 tiet dau sang" in plain:
        return [1, 2, 3]
    if "3 tiet cuoi sang" in plain:
        return [4, 5, 6]
    if "3 tiet dau chieu" in plain:
        return [7, 8, 9]
    if "3 tiet cuoi chieu" in plain:
        return [10, 11, 12]
    explicit = _periods_from_text(text)
    if explicit:
        return explicit
    if re.search(r'tiet\s*\d',plain): return []
    if "ca ngay" in plain:
        return list(range(1, 16))
    if "sang" in plain:
        return list(range(1, 7))
    if "chieu" in plain:
        return list(range(7, 13))
    if "toi" in plain:
        return list(range(13, 16))
    return []


def _target(day_scope: str | None, periods: list[int], numeric_value: float | None = None) -> dict:
    target: dict = {"day_scope": day_scope, "periods": periods}
    if day_scope in {"T2", "T3", "T4", "T5", "T6", "T7"}:
        target["weekday"] = int(day_scope[1:])
    elif day_scope == "CN":
        target["weekday"] = 8
    elif day_scope in {"ALL_WEEKDAYS", "ALL_DAYS"}:
        weekdays = range(2, 7) if day_scope == "ALL_WEEKDAYS" else range(2, 9)
        target["slots"] = [{"weekday": weekday, "periods": periods} for weekday in weekdays]
    if numeric_value is not None:
        target["value"] = numeric_value
        target["max"] = numeric_value
    return target


def interpret_raw_preference(text: str, *, semester_start: date | None = None,
                            semester_end: date | None = None) -> list[dict]:
    """Deterministic, atomic semantics for the free-text field in V3 forms.

    This deliberately returns no RAW_NOTE rule: provenance is retained by the
    caller, while only the returned structured rules may become solver-visible.
    """
    plain = _plain(text)
    if not plain or "khong co nguyen vong dac biet" in plain:
        return []
    rules: list[dict] = []
    def add(kind, target, hardness="soft", weight=None, **extra):
        rules.append({"constraint_type": kind, "target": target, "hardness": hardness,
                      "weight": PREFERENCE_WEIGHTS[weight] if isinstance(weight, str) else (weight if weight is not None else PREFERENCE_WEIGHTS["MEDIUM"]), **extra})
    def add_blocks(kind, scope, blocks, hardness="soft", weight=None):
        for periods in blocks:
            add(kind, _target(scope, periods), hardness, weight)

    # Explicit seminar/busy slots are hard only where both time and day are stated.
    for m in re.finditer(r"seminar\s+(sang|chieu|toi|(?:tiet\s*)?\d+\s*[-–]\s*\d+)\s+(?:thu|t)\s*([2-7])", plain):
        periods = _periods_for_part_of_day(m.group(1))
        if periods: add("UNAVAILABLE", _target(f"T{m.group(2)}", periods), "hard", 1.0)
    # “sáng seminar (có thể đến 12h30)” covers the first three HUCE blocks.
    for m in re.finditer(r"(?:t|thu)\s*([2-7])\s+sang\s+seminar[^.;]*12h30", plain):
        add_blocks("UNAVAILABLE", f"T{m.group(1)}", ([1,2,3], [4,5,6], [7,8,9]), "hard", 1.0)
    if re.search(r"xin nghi\s+sang\s+(?:thu|t)\s*([2-7])", plain):
        d=re.search(r"xin nghi\s+sang\s+(?:thu|t)\s*([2-7])", plain).group(1)
        add_blocks("UNAVAILABLE", f"T{d}", ([1,2,3], [4,5,6]), "hard", 1.0)
    if re.search(r"khong day\s+tiet\s*(\d+)\s*[-–]\s*(\d+)\s+(?:thu|t)\s*([2-7])", plain):
        m=re.search(r"khong day\s+tiet\s*(\d+)\s*[-–]\s*(\d+)\s+(?:thu|t)\s*([2-7])", plain)
        add("UNAVAILABLE", _target(f"T{m.group(3)}", list(range(int(m.group(1)),int(m.group(2))+1))), "hard", 1.0)

    # Compactness and weekday intentions are separate, normalized semantic rules.
    day_matches = re.search(r"(?:cac\s+ngay|ngay)\s+t?([2-7])\s*,?\s*(?:t|thu)?([2-7])\s*,?\s*(?:t|thu)?([2-7])\s*,?\s*(?:t|thu)?([2-7])", plain)
    if ("day gon" in plain or "lich day gon" in plain) and day_matches:
        preferred = sorted({int(x) for x in day_matches.groups()})
        add("PREFERRED_DAYS", {"weekdays": preferred}, "soft", "DAY")
        avoided = sorted(set(range(2, 8)) - set(preferred))
        if avoided: add("AVOID_DAYS", {"weekdays": avoided}, "soft", "AVOID_DAY")
        add("PREFER_COMPACT_SCHEDULE", {"day_scope": "ALL_DAYS"}, "soft", "VERY_STRONG")
    elif "day gon" in plain or "day lien" in plain or "lien 6 tiet" in plain:
        add("PREFER_COMPACT_SCHEDULE", {"day_scope": "ALL_DAYS"}, "soft", "VERY_STRONG")

    morning = list(range(1, 7)); early_afternoon = list(range(7, 10)); afternoon = list(range(7, 13))
    if re.search(r"(?:dang ky|dang ki|xin|duoc)\s+day\s+sang|day\s+cac\s+buoi\s+sang", plain):
        add("PREFERRED_PERIOD", _target("ALL_WEEKDAYS", morning), "soft", "STRONG")
    if "co the nhan them tiet 7-9" in plain or "co the keo dai den tiet 7-9" in plain:
        add("PREFERRED_PERIOD", _target("ALL_WEEKDAYS", early_afternoon), "soft", "WEAK")
    if "co the day vai buoi chieu" in plain:
        add("PREFERRED_PERIOD", _target("ALL_WEEKDAYS", afternoon), "soft", "WEAK")
    if "xin tranh tiet 1-3" in plain:
        add("AVOID_PERIOD", _target("ALL_DAYS", [1,2,3]), "soft", "MEDIUM")
    if "3 tiet cuoi sang" in plain or "3 tiet dau chieu" in plain:
        add("PREFERRED_PERIOD", _target("ALL_WEEKDAYS", [4,5,6]), "soft", "STRONG")
        add("PREFERRED_PERIOD", _target("ALL_WEEKDAYS", [7,8,9]), "soft", "STRONG")
    if re.search(r"xin day it|hoac day it", plain):
        target={"day_scope":"ALL_DAYS"}
        date_range=_extract_date_unavailable(text, semester_start=semester_start, semester_end=semester_end)
        if date_range:
            start,end=date_range; target.update({"start_date":start.isoformat() if start else None,"end_date":end.isoformat() if end else None})
        add("PREFER_LOW_WORKLOAD", target, "soft", "MEDIUM")
    return rules


def describe_preference(constraint_type: str, target: dict) -> str:
    if target.get("period_alternatives"):
        alternatives = " HOẶC ".join(f"tiết {min(p)}–{max(p)}" for p in target["period_alternatives"])
        return f"Ưu tiên {alternatives} · {target.get('day_scope') or 'theo các ngày đã chọn'}"
    normalized = constraint_type.upper()
    normalized = {"PREFER_PERIOD": "PREFERRED_PERIOD", "AVOID": "AVOID_PERIOD", "RAW_PREFERENCE": "RAW_NOTE"}.get(normalized, normalized)
    labels = {
        "UNAVAILABLE": "Không xếp lịch", "AVOID_PERIOD": "Ưu tiên tránh",
        "PREFERRED_PERIOD": "Ưu tiên xếp lịch", "PREFER_CONSECUTIVE_PERIODS": "Ưu tiên dạy liền",
        "PREFERRED_DAYS": "Ưu tiên ngày dạy", "AVOID_DAYS": "Ưu tiên tránh ngày dạy",
        "PREFER_LOW_WORKLOAD": "Ưu tiên tải dạy thấp", "PREFER_COMPACT_SCHEDULE": "Ưu tiên lịch dạy gọn",
        "MIN_CLASSES": "Số lớp tối thiểu", "MAX_CLASSES": "Số lớp tối đa",
        "MAX_SESSIONS_PER_DAY": "Số buổi tối đa mỗi ngày", "MAX_DAYS_PER_WEEK": "Số ngày tối đa mỗi tuần",
        "MIN_FREE_MORNING_PER_WEEK": "Số buổi sáng nghỉ tối thiểu", "REQUIRED_ASSIGNMENT": "Bắt buộc phân công",
        "FORBIDDEN_ASSIGNMENT": "Không được phân công", "RAW_NOTE": "Cần trưởng bộ môn diễn giải",
        "AVAILABLE": "Có thể xếp lịch", "SEMINAR": "Giữ lịch seminar",
    }
    slots = target.get("slots") or []
    if slots:
        rendered = []
        for slot in slots:
            weekday = slot.get("weekday")
            periods = slot.get("periods") or []
            day = "Chủ Nhật" if weekday == 8 else f"Thứ {weekday}"
            rendered.append(f"{day}{f', tiết {min(periods)}–{max(periods)}' if periods else ''}")
        return f"{labels.get(normalized, 'Ràng buộc')} · {'; '.join(rendered)}"
    scope = target.get("day_scope")
    weekday = target.get("weekday")
    day = {"ALL_WEEKDAYS": "Thứ 2–Thứ 6", "ALL_DAYS": "mọi ngày", "CN": "Chủ Nhật"}.get(
        scope, "Chủ Nhật" if weekday == 8 else (f"Thứ {weekday}" if weekday else "mọi ngày")
    )
    periods = target.get("periods") or target.get("period_range") or []
    suffix = f", tiết {min(periods)}–{max(periods)}" if periods else ""
    return f"{labels.get(normalized, 'Ràng buộc')} · {day}{suffix}"


def _extract_date_unavailable(text: str, default_year: int | None = None, *, semester_start: date | None = None, semester_end: date | None = None) -> tuple[date | None, date | None] | None:
    plain = _plain(text)
    # Check "truoc" pattern first (e.g. khong the day truoc 12/10 -> ends 11/10 inclusive)
    match_before = re.search(
        r"(?:khong\s*(?:the\s*)?day|nghi|ban)\s+truoc(?:\s+ngay)?\s*(\d{1,2})\s*(?:/|thang\s*)(\d{1,2})(?:\s*(?:/|nam\s*)(\d{4}))?",
        plain,
    )
    if match_before:
        d = int(match_before.group(1))
        m = int(match_before.group(2))
        y = int(match_before.group(3)) if match_before.group(3) else _resolve_year(d, m, default_year, semester_start, semester_end)
        if y is None:
            return None
        try:
            target_dt = date(y, m, d)
            end_dt = target_dt - timedelta(days=1)
            return semester_start, end_dt
        except ValueError:
            pass

    # "khong day den", "nghi den", "xin nghi den", "ban den"
    match_until = re.search(
        r"(?:khong\s*(?:the\s*)?day|nghi|xin\s+nghi|ban)\s+den(?:\s+ngay)?[:\s]*(\d{1,2})\s*(?:/|thang\s*)(\d{1,2})(?:\s*(?:/|nam\s*)(\d{4}))?",
        plain,
    )
    if match_until:
        d = int(match_until.group(1))
        m = int(match_until.group(2))
        y = int(match_until.group(3)) if match_until.group(3) else _resolve_year(d, m, default_year, semester_start, semester_end)
        if y is None:
            return None
        try:
            end_dt = date(y, m, d)
            return semester_start, end_dt
        except ValueError:
            pass
    return None


def _resolve_year(day, month, default_year, start, end):
    if default_year is not None:
        return default_year
    if start and end:
        candidates = []
        for year in range(start.year, end.year + 1):
            try:
                candidate = date(year, month, day)
            except ValueError:
                continue
            if start <= candidate <= end:
                candidates.append(year)
        if len(candidates) == 1:
            return candidates[0]
    return None


def _legacy_rule_impl(text: str, column_weekday: int | None, *, semester_start: date | None = None, semester_end: date | None = None) -> tuple[str, dict, str, str | None]:
    plain = _plain(text)
    if re.search(r'(?:thu|\bt)\s*[2-7]\s*(?:,|hoac|va)',plain):
        return 'RAW_NOTE', {'day_alternatives':re.findall(r'(?:thu|\bt)\s*([2-7])',plain)}, 'LOW', 'Danh sách/lựa chọn ngày cần rà soát; không tự chọn một ngày.'
    if any(token in plain for token in ("con lai", "vai buoi", "mot so buoi", "hoac day it", "12h30")):
        return "RAW_NOTE", {}, "LOW", "Điều kiện bổ sung chưa thể biểu diễn đầy đủ; cần xác nhận nguyên văn."
    date_unavail = _extract_date_unavailable(text, semester_start=semester_start, semester_end=semester_end)
    if date_unavail:
        start_d, end_d = date_unavail
        target = {
            "start_date": start_d.isoformat() if start_d else None,
            "end_date": end_d.isoformat() if end_d else None,
            "periods": list(range(1, 16)),
            "day_scope": "ALL_DAYS",
        }
        return "UNAVAILABLE", target, "HIGH", None
    if re.search(r"(?:nghi|ban|day)\s+(?:den|truoc)", plain):
        return "RAW_NOTE", {}, "LOW", "Không xác định chắc chắn ngày/năm từ kỳ học; cần xác nhận."

    explicit_day = _weekday_from_text(text)
    generic_part_of_day = explicit_day is None and any(token in plain for token in ("sang", "chieu", "toi"))
    day = _day_scope(explicit_day or column_weekday, text)
    periods = _periods_for_part_of_day(text)
    rule_day = day
    if "chieu:" in plain and periods and max(periods) > 12:
        return "RAW_NOTE", _target(rule_day, periods), "MEDIUM", "Câu nói buổi chiều nhưng mở từ tiết đến cuối ngày; cần xác nhận có bao gồm buổi tối."
    day_range = re.search(r"(?:thu|\bt)\s*([2-7])\s*(?:-|–|den)\s*(?:(?:thu|t)\s*)?([2-7])", plain)
    if day_range and day_range.groups() != ("2", "6"):
        return "RAW_NOTE", {"weekdays": list(range(int(day_range[1]), int(day_range[2]) + 1)), "periods": periods}, "MEDIUM", "Khoảng ngày cần xác nhận trước khi áp dụng."
    if "hoac" in plain:
        alternatives = [_periods_for_part_of_day(part) for part in re.split(r"\bhoac\b", plain)]
        if len(alternatives) == 2 and all(alternatives) and not any(t in plain for t in ("khong", "nghi", "tranh")):
            target = _target(rule_day, [])
            days = [slot["weekday"] for slot in target.get("slots", [])] or [target.get("weekday")]
            target["period_alternatives"] = alternatives
            target["slots"] = [{"weekday": d, "periods": p} for d in days for p in alternatives]
            return "PREFERRED_PERIOD", target, "HIGH", None
        return "RAW_NOTE", {}, "LOW", "Các lựa chọn HOẶC cần được xác nhận; không gộp thành khoảng liên tục."
    if "nghi it nhat" in plain and "buoi sang" in plain:
        match = re.search(r"(?:it nhat(?: la)?|toi thieu)\s*0?(\d+)", plain)
        if not match:
            return 'RAW_NOTE',{},'LOW','Chưa xác định số buổi sáng nghỉ tối thiểu.'
        value = float(match.group(1))
        return "MIN_FREE_MORNING_PER_WEEK", _target("ALL_WEEKDAYS", list(range(1, 7)), value), "HIGH", None
    consecutive = re.search(r"(?:lien|lien tuc)\s*0?(\d+)\s*tiet", plain)
    if consecutive:
        value = float(consecutive.group(1))
        return "PREFER_CONSECUTIVE_PERIODS", _target(rule_day, [], value), "HIGH", None
    if any(token in plain for token in ("xin nghi", "nghi day", "khong day", "khong the day", "khong xep")):
        if periods or explicit_day or "cac buoi" in plain:
            return "UNAVAILABLE", _target(rule_day, periods or list(range(1,16))), "HIGH", None
    if "tranh" in plain and periods:
        return "AVOID_PERIOD", _target(rule_day, periods), "HIGH", None
    if "3 tiet cuoi sang" in plain or "3 tiet dau chieu" in plain:
        return "RAW_NOTE", _target("ALL_WEEKDAYS", periods), "MEDIUM", "Khung '3 tiết cuối sáng/đầu chiều' cần xác nhận theo quy ước tiết của bộ môn."
    if any(token in plain for token in ("dang ki day", "dang ky day", "dk day", "xin day", "duoc day", "day cac", "co the")) or re.search(r"\bday\b", plain):
        if periods or any(token in plain for token in ("sang", "chieu", "toi")):
            return "PREFERRED_PERIOD", _target(rule_day, periods), "HIGH", None
    if any(token in plain for token in ("gon", "it ngay")):
        return "RAW_NOTE", _target(day, periods), "MEDIUM", "Chưa đủ cấu trúc để suy ra một quy tắc solver duy nhất."
    return "RAW_NOTE", _target(day, periods), "LOW", "Không nhận diện chắc chắn; cần trưởng bộ môn diễn giải."


def _legacy_rule(text, column_weekday, *, semester_start=None, semester_end=None):
    kind,target,confidence,reason = _legacy_rule_impl(text,column_weekday,semester_start=semester_start,semester_end=semester_end)
    if _uncertain(text,'period'):
        target={**target,'periods':[]}
        target.pop('period_alternatives',None);target.pop('slots',None)
        confidence='LOW';reason='Nguồn nói rõ chưa xác định tiết; cần xác nhận.'
    if kind in {'UNAVAILABLE','AVOID_PERIOD','PREFERRED_PERIOD'}:
        if not target.get('day_scope') and not (target.get('start_date') and target.get('end_date')):
            confidence='LOW';reason='Chưa xác định phạm vi ngày từ nguồn; cần xác nhận.'
        if not target.get('periods') and not target.get('period_alternatives'):
            confidence='LOW';reason='Chưa xác định tiết từ nguồn; cần xác nhận.'
    return kind,target,confidence,reason


def _infer(text: str, weekday: int | None) -> tuple[str, dict, float]:
    """Compatibility wrapper used by older callers and tests."""
    kind, target, confidence, _ = _legacy_rule(text, weekday)
    legacy_kind = {
        "UNAVAILABLE": "unavailable", "AVOID_PERIOD": "avoid", "PREFERRED_PERIOD": "prefer_period",
        "RAW_NOTE": "raw_preference", "PREFER_CONSECUTIVE_PERIODS": "compact_schedule",
    }.get(kind, kind)
    return legacy_kind, target, {"HIGH": 0.9, "MEDIUM": 0.7, "LOW": 0.35}[confidence]


def _as_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"Ngày không hợp lệ: {text}")


def _as_number(value: object, *, default: float | None = None) -> float | None:
    if value in (None, ""):
        return default
    number = float(value)
    return int(number) if number.is_integer() else number


def _split_legacy_clauses(text: str) -> list[str]:
    chunks = re.split(
        r"(?:\r?\n)+|(?<=[.!?])\s+|,\s*(?=(?:xin|em xin|cô |co |thầy |thay |seminar|có thể|co the)\b)|\s+[Vv]à\s+(?=xin\b)",
        text.strip(), flags=re.IGNORECASE,
    )
    return [chunk.strip(" ,.;") for chunk in chunks if chunk.strip(" ,.;")]


def detect_preference_format(path: Path) -> str:
    from app.parsers.grid_preferences import is_grid_preference_template
    if is_grid_preference_template(path):
        return 'GRID_V3'
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return 'STRUCTURED_V2' if _structured_layout(book) else 'LEGACY'
    finally:
        book.close()


def _structured_layout(book):
    for sheet in book:
        for number,row in enumerate(sheet.iter_rows(max_row=60),1):
            headers=[_normalized_header(c.value) for c in row]
            if all(_normalized_header(v) in headers for v in V2_PREFERENCE_HEADERS):
                return sheet,number,headers
    return None


def _invalid_structured_draft(sheet, row: int, values: list[object], reason: str) -> ParsedPreference:
    raw = clean_text(values[12] if len(values) > 12 else "") or " | ".join(clean_text(value) for value in values if value)
    return ParsedPreference(
        clean_text(values[1] if len(values) > 1 else ""), clean_text(values[1] if len(values) > 1 else ""),
        None, "RAW_NOTE", {}, raw, 0.0, row, lecturer_code=clean_text(values[0]) or None,
        source_sheet=sheet.title, source_cell=f"A{row}:N{row}", confidence_label="LOW",
        needs_review=True, review_reason=reason, status="NEEDS_REVIEW",
    )


def _parse_structured(book) -> PreferenceParseResult:
    result = PreferenceParseResult(format="STRUCTURED_V2")
    sheet,header_row,headers = _structured_layout(book)
    for row in range(header_row+1, sheet.max_row + 1):
        values = [sheet.cell(row,headers.index(_normalized_header(header))+1).value for header in V2_PREFERENCE_HEADERS]
        if not any(value not in (None, "") for value in values):
            continue
        result.raw_clauses += 1
        try:
            code, name, context, kind, scope = (clean_text(values[index]) for index in range(5))
            context = context.upper()
            hardness = clean_text(values[9]).upper()
            status = (clean_text(values[13]) or "DRAFT").upper()
            if not code:
                raise ValueError("Thiếu Mã GV; cần resolve identity trước khi xác nhận.")
            if kind.upper() not in CONSTRAINT_TYPES:
                raise ValueError(f"Loại ràng buộc không hợp lệ: {kind}")
            if context not in {"TEACHING", "SEMINAR", "MIXED"}:
                raise ValueError(f"Ngữ cảnh không hợp lệ: {context}")
            if scope.upper() not in DAY_SCOPES:
                raise ValueError(f"Thứ/Phạm vi không hợp lệ: {scope}")
            if hardness not in {"SOFT", "HARD"}:
                raise ValueError(f"Độ cứng không hợp lệ: {hardness}")
            if status not in {"DRAFT", "CONFIRMED", "NEEDS_REVIEW", "REJECTED"}:
                raise ValueError(f"Trạng thái không hợp lệ: {status}")
            start = int(values[5]) if values[5] not in (None, "") else None
            end = int(values[6]) if values[6] not in (None, "") else None
            if (start is None) != (end is None) or (start is not None and not 1 <= start <= end <= 15):
                raise ValueError("Khoảng tiết phải có đủ bắt đầu/kết thúc và nằm trong 1–15.")
            periods = list(range(start, end + 1)) if start is not None and end is not None else []
            weight = _as_number(values[10], default=0.8)
            if weight is None or not 0 <= weight <= 1:
                raise ValueError("Trọng số phải nằm trong 0–1.")
            numeric = _as_number(values[11])
            start_date, end_date = _as_date(values[7]), _as_date(values[8])
            if start_date and end_date and end_date < start_date:
                raise ValueError("Ngày kết thúc phải không trước ngày bắt đầu.")
            raw = clean_text(values[12])
            unsupported = False
            incomplete_assignment = kind.upper() in {"REQUIRED_ASSIGNMENT", "FORBIDDEN_ASSIGNMENT"}
            seminar_types = {"SEMINAR_COMMITMENT", "SEMINAR_NOTE"}
            incompatible_context = (
                (context == "TEACHING" and kind.upper() in seminar_types)
                or (context == "SEMINAR" and kind.upper() not in seminar_types)
            )
            needs_review = (
                status == "NEEDS_REVIEW"
                or kind.upper() in {"RAW_NOTE", "SEMINAR_NOTE"}
                or context == "MIXED"
                or unsupported
                or incomplete_assignment
                or incompatible_context
            )
            if context == "SEMINAR" and kind.upper() == "SEMINAR_COMMITMENT" and (not periods or not scope):
                needs_review = True
            parsed_target = _target(scope.upper(), periods, numeric)
            if kind.upper() == "MIN_CLASSES" and numeric is not None:
                parsed_target.pop("max", None); parsed_target["min"] = numeric
            if start_date: parsed_target["start_date"] = start_date.isoformat()
            if end_date: parsed_target["end_date"] = end_date.isoformat()
            parsed_target['_provenance']={field:{'origin':'SOURCE_EXPLICIT' if values[index] not in (None,'') else 'DEFAULT' if field=='weight' else 'UNKNOWN','sheet':sheet.title,'row':row}
                for field,index in {'lecturer_id':0,'context_type':2,'constraint_type':3,'day_scope':4,'periods':5,'start_date':7,'end_date':8,'hardness':9,'weight':10,'numeric_value':11}.items()}
            result.drafts.append(ParsedPreference(
                lecturer_alias=name or code, canonical_name=name or code, weekday=None,
                constraint_type=kind.upper(), target=parsed_target, raw_text=raw,
                confidence=1.0, source_row=row, lecturer_code=code, source_sheet=sheet.title,
                source_cell=f"A{row}:N{row}", day_scope=scope.upper(), start_period=start, end_period=end,
                start_date=start_date, end_date=end_date, hardness=hardness.lower(), weight=weight,
                numeric_value=numeric, confidence_label="HIGH", needs_review=needs_review,
                review_reason=("MIXED phải tách thành các rule nguyên tử trước khi Apply." if context == "MIXED" else
                               "Loại rule chưa phù hợp ngữ cảnh đã chọn." if incompatible_context else
                               "Rule đã được bảo toàn nhưng solver hiện chưa có semantics tương ứng." if unsupported else
                               "Cần chọn TeachingGroup cụ thể trước khi Apply." if incomplete_assignment else
                               "Ghi chú chưa có scheduling semantics cụ thể." if kind.upper() in {"RAW_NOTE", "SEMINAR_NOTE"} else
                               "Thiếu ngày hoặc tiết seminar." if needs_review else None),
                status="NEEDS_REVIEW" if needs_review else status, context_type=context,
                context_confidence="HIGH",
            ))
        except (TypeError, ValueError) as error:
            result.invalid_rows += 1
            result.drafts.append(_invalid_structured_draft(sheet, row, values, str(error)))
    if "Seminar_Shared" in book.sheetnames:
        seminar_sheet = book["Seminar_Shared"]
        for row in range(2, seminar_sheet.max_row + 1):
            values = [seminar_sheet.cell(row, column).value for column in range(1, 10)]
            if not any(value not in (None, "") for value in values):
                continue
            result.raw_clauses += 1
            try:
                code, name = clean_text(values[0]), clean_text(values[1])
                participants = [item.strip() for item in clean_text(values[2]).split(";") if item.strip()]
                days = [item.strip().upper() for item in clean_text(values[3]).split(";") if item.strip()]
                blocks = []
                for block in clean_text(values[4]).split(";"):
                    periods = _periods_from_text(f"tiết {block.strip()}")
                    if not periods:
                        raise ValueError(f"Khung tiết không hợp lệ: {block}")
                    blocks.append(periods)
                hardness = clean_text(values[5]).upper()
                weight = _as_number(values[6], default=0.8)
                status = (clean_text(values[8]) or "DRAFT").upper()
                if not code or not name or not participants:
                    raise ValueError("Thiếu mã, tên hoặc người tham gia seminar.")
                if not days or any(day not in DAY_SCOPES for day in days):
                    raise ValueError("Thứ cho phép của seminar không hợp lệ.")
                if hardness not in {"SOFT", "HARD"} or status not in {"DRAFT", "CONFIRMED", "NEEDS_REVIEW", "REJECTED"} or weight is None or not 0 <= weight <= 1:
                    raise ValueError("Hardness, weight hoặc status seminar không hợp lệ.")
                result.seminars.append(ParsedSharedSeminar(
                    code, name, participants, days, blocks, hardness.lower(), weight, clean_text(values[7]),
                    seminar_sheet.title, row, f"A{row}:I{row}", status,
                    needs_review=status == "NEEDS_REVIEW",
                ))
            except (TypeError, ValueError) as error:
                result.invalid_rows += 1
                result.seminars.append(ParsedSharedSeminar(
                    clean_text(values[0]), clean_text(values[1]) or "Seminar cần rà soát", [], [], [], "soft", 0.8,
                    clean_text(values[7]), seminar_sheet.title, row, f"A{row}:I{row}", "NEEDS_REVIEW", True, str(error),
                ))
    return result


def _legacy_layout(book):
    candidates=[]
    identity_headers={'ten','ten giang vien','ho ten','ho va ten','ho ten gv','giang vien','lecturer','lecturer name','ten gv'}
    code_headers={'ma gv','ma giang vien','lecturer code','code','ma cbgv'}
    for sheet in book:
        for row in sheet.iter_rows(max_row=min(sheet.max_row,60)):
            headers={c.column:_plain(str(c.value or '')).strip(' *:') for c in row}
            names=[c for c,h in headers.items() if h in identity_headers]
            codes=[c for c,h in headers.items() if h in code_headers]
            days={c:_weekday_from_text(h) for c,h in headers.items() if re.fullmatch(r'(?:thu\s*[2-7]|t[2-7]|cn|chu nhat)',h)}
            notes=[c for c,h in headers.items() if any(k in h for k in ('ghi chu','nguyen vong','yeu cau','preferences','notes'))]
            if len(names)<=1 and len(codes)<=1 and (names or codes) and (days or notes):
                candidates.append((len(days)+len(notes),sheet,row[0].row,names[0] if names else None,codes[0] if codes else None,days,notes))
    if not candidates:
        raise ValueError('PREFERENCE_STRUCTURE_UNRECOGNIZED: Không nhận diện chắc chắn cột giảng viên và ngày/nguyện vọng.')
    candidates.sort(key=lambda x:x[0],reverse=True)
    if len(candidates)>1 and candidates[0][0]==candidates[1][0]:
        raise ValueError('PREFERENCE_STRUCTURE_AMBIGUOUS: Có nhiều bảng nguyện vọng; cần chọn nguồn rõ ràng.')
    return candidates[0][1:]


def _parse_legacy(book, *, semester_start=None, semester_end=None) -> PreferenceParseResult:
    sheet,header,name_col,code_col,day_columns,note_columns = _legacy_layout(book)
    result = PreferenceParseResult(format="LEGACY")
    source_clauses = set()
    def cell_value(row,column):
        if column is None: return None
        cell=sheet.cell(row,column)
        if cell.value is not None: return cell.value
        for merged in sheet.merged_cells.ranges:
            if cell.coordinate in merged: return sheet.cell(merged.min_row,merged.min_col).value
        return None
    for row in range(header+1, sheet.max_row + 1):
        code=clean_text(cell_value(row,code_col)) or None
        alias = clean_text(cell_value(row,name_col)) or code or ''
        if not alias and not any(sheet.cell(row,c).value for c in set(day_columns)|set(note_columns)):
            continue
        if not alias: result.invalid_rows+=1
        for column in sorted(set(day_columns) | set(note_columns)):
            raw_cell = str(sheet.cell(row,column).value or '').strip()
            if not raw_cell:
                continue
            clauses = _split_legacy_clauses(raw_cell)
            result.raw_clauses += len(clauses)
            for clause_index, clause in enumerate(clauses, start=1):
                source_clauses.add(f"{sheet.cell(row, column).coordinate}#{clause_index}")
                weekday = day_columns.get(column)
                if any(sheet.cell(row,column).coordinate in m and m.min_col!=m.max_col for m in sheet.merged_cells.ranges):
                    weekday=None
                clause_plain = _plain(clause)
                segmented_teaching = re.search(
                    r"(?i)(?:sáng|chiều|tối|sang|chieu|toi)\s*:\s*(?:(?:có thể|co the|xin|muốn|muon)\s+)?dạy\b",
                    clause,
                )
                if "seminar" in clause_plain and segmented_teaching:
                    seminar_text = clause[:segmented_teaching.start()].strip()
                    teaching_text = clause[segmented_teaching.start():].strip()
                    seminar_day = _day_scope(_weekday_from_text(seminar_text) or weekday, seminar_text)
                    seminar_periods = _periods_for_part_of_day(seminar_text)
                    seminar_complete = bool(seminar_day and seminar_periods) and "12h30" not in _plain(seminar_text)
                    seminar_target = _target(seminar_day, seminar_periods)
                    teaching_kind, teaching_target, teaching_confidence, teaching_reason = _legacy_rule(teaching_text, weekday, semester_start=semester_start, semester_end=semester_end)
                    teaching_periods = teaching_target.get("periods") or []
                    teaching_review = teaching_confidence == "LOW" or teaching_kind == "RAW_NOTE"
                    common = dict(
                        lecturer_alias=alias, canonical_name=alias, raw_text=clause,
                        source_row=row, source_sheet=sheet.title, context_confidence="HIGH",
                    )
                    result.drafts.append(ParsedPreference(
                        weekday=seminar_target.get("weekday"),
                        constraint_type="SEMINAR_COMMITMENT" if seminar_complete else "SEMINAR_NOTE",
                        target=seminar_target, confidence=0.9 if seminar_complete else 0.35,
                        source_cell=f"{sheet.cell(row, column).coordinate}#{clause_index}A",
                        day_scope=seminar_day,
                        start_period=min(seminar_periods) if seminar_periods else None,
                        end_period=max(seminar_periods) if seminar_periods else None,
                        confidence_label="HIGH" if seminar_complete else "LOW",
                        needs_review=not seminar_complete,
                        review_reason=None if seminar_complete else "Thiếu ngày hoặc tiết seminar; không tự suy đoán.",
                        status="DRAFT" if seminar_complete else "NEEDS_REVIEW",
                        context_type="SEMINAR", **common,
                    ))
                    result.drafts.append(ParsedPreference(
                        weekday=teaching_target.get("weekday"), constraint_type=teaching_kind,
                        target=teaching_target,
                        confidence={"HIGH": 0.9, "MEDIUM": 0.7, "LOW": 0.35}[teaching_confidence],
                        source_cell=f"{sheet.cell(row, column).coordinate}#{clause_index}B",
                        day_scope=teaching_target.get("day_scope"),
                        start_period=min(teaching_periods) if teaching_periods else None,
                        end_period=max(teaching_periods) if teaching_periods else None,
                        numeric_value=teaching_target.get("value"), confidence_label=teaching_confidence,
                        needs_review=teaching_review, review_reason=teaching_reason,
                        status="NEEDS_REVIEW" if teaching_review else "DRAFT",
                        context_type="TEACHING", **common,
                    ))
                    continue
                is_mixed = "seminar" in clause_plain and any(token in clause_plain for token in (" vi ", " nen ")) and any(token in clause_plain for token in ("day", "gon"))
                if is_mixed:
                    before, after = re.split(r"\b(?:vi|nen)\b", clause_plain, maxsplit=1)
                    seminar_text, teaching_text = (before, after) if "seminar" in before else (after, before)
                    seminar_days = sorted({int(value) for value in re.findall(r"[2-7]", seminar_text)})
                    teaching_days = sorted({int(value) for value in re.findall(r"[2-7]", teaching_text)})
                    common = dict(
                        lecturer_alias=alias, canonical_name=alias, weekday=None, raw_text=clause,
                        confidence=0.7, source_row=row, source_sheet=sheet.title,
                        confidence_label="MEDIUM", needs_review=True, status="NEEDS_REVIEW",
                        context_confidence="HIGH",
                    )
                    result.drafts.append(ParsedPreference(
                        constraint_type="SEMINAR_NOTE", target={"weekdays": seminar_days, "periods": []},
                        source_cell=f"{sheet.cell(row, column).coordinate}#{clause_index}A",
                        review_reason="Đã tách phần seminar nhưng chưa có khung tiết cụ thể.",
                        context_type="SEMINAR", **common,
                    ))
                    result.drafts.append(ParsedPreference(
                        constraint_type="RAW_NOTE", target={"weekdays": teaching_days, "periods": []},
                        source_cell=f"{sheet.cell(row, column).coordinate}#{clause_index}B",
                        review_reason="Đã tách phần lịch dạy; cần chọn scheduling semantics cụ thể.",
                        context_type="TEACHING", **common,
                    ))
                    continue
                is_seminar_event = "seminar" in clause_plain and not any(token in clause_plain for token in ("xin nghi", "nghi day", "khong day"))
                if is_seminar_event:
                    seminar_day = _day_scope(_weekday_from_text(clause) or weekday, clause)
                    seminar_periods = _periods_for_part_of_day(clause)
                    complete = bool(seminar_day and seminar_periods) and "12h30" not in clause_plain
                    target = _target(seminar_day, seminar_periods)
                    result.drafts.append(ParsedPreference(
                        alias, alias, target.get("weekday"), "SEMINAR_COMMITMENT" if complete else "SEMINAR_NOTE",
                        target, clause, 0.9 if complete else 0.35, row, source_sheet=sheet.title,
                        source_cell=f"{sheet.cell(row, column).coordinate}#{clause_index}", day_scope=seminar_day,
                        start_period=min(seminar_periods) if seminar_periods else None,
                        end_period=max(seminar_periods) if seminar_periods else None,
                        confidence_label="HIGH" if complete else "LOW", needs_review=not complete,
                        review_reason=None if complete else "Thiếu ngày hoặc tiết seminar; không tự suy đoán.",
                        status="DRAFT" if complete else "NEEDS_REVIEW", context_type="SEMINAR",
                        context_confidence="HIGH",
                    ))
                    continue
                kind, target, confidence, reason = _legacy_rule(clause, weekday, semester_start=semester_start, semester_end=semester_end)
                numeric = target.get("value")
                confidence_value = {"HIGH": 0.9, "MEDIUM": 0.7, "LOW": 0.35}[confidence]
                needs_review = confidence == "LOW" or kind == "RAW_NOTE"
                periods = target.get("periods") or []
                start_date = date.fromisoformat(target["start_date"]) if target.get("start_date") else None
                end_date = date.fromisoformat(target["end_date"]) if target.get("end_date") else None
                result.drafts.append(ParsedPreference(
                    alias, alias, target.get("weekday"), kind, target, clause, confidence_value, row,
                    source_sheet=sheet.title, source_cell=f"{sheet.cell(row, column).coordinate}#{clause_index}",
                    day_scope=target.get("day_scope"), start_period=min(periods) if periods else None,
                    end_period=max(periods) if periods else None,
                    start_date=start_date, end_date=end_date,
                    numeric_value=numeric,
                    confidence_label=confidence, needs_review=needs_review, review_reason=reason,
                    status="NEEDS_REVIEW" if needs_review else "DRAFT", context_type="TEACHING",
                    context_confidence="HIGH",
                ))
    preserved = {re.sub(r"(?<=\d)[AB]$", "", draft.source_cell) for draft in result.drafts}
    result.dropped_clauses = len(source_clauses - preserved)
    for draft in result.drafts:
        draft.lecturer_code=clean_text(cell_value(draft.source_row,code_col)) or None
        if _uncertain(draft.raw_text,'identity'):
            draft.target={**draft.target,'identity_uncertain':True}
            draft.needs_review=True;draft.status='NEEDS_REVIEW';draft.confidence_label='LOW'
        if not draft.lecturer_alias:
            draft.needs_review=True;draft.status='NEEDS_REVIEW';draft.confidence_label='LOW'
        origins={
            'lecturer_id':'UNKNOWN', 'constraint_type':'SOURCE_INFERRED_DETERMINISTIC',
            'context_type':'SOURCE_INFERRED_DETERMINISTIC',
            'day_scope':'SOURCE_EXPLICIT' if draft.day_scope else 'UNKNOWN',
            'periods':('SOURCE_EXPLICIT' if _periods_from_text(draft.raw_text) else 'SOURCE_INFERRED_DETERMINISTIC') if draft.target.get('periods') or draft.target.get('period_alternatives') else 'UNKNOWN',
            'start_date':'SOURCE_INFERRED_DETERMINISTIC' if draft.start_date else 'UNKNOWN',
            'end_date':'SOURCE_INFERRED_DETERMINISTIC' if draft.end_date else 'UNKNOWN',
            'numeric_value':'SOURCE_EXPLICIT' if draft.numeric_value is not None else 'UNKNOWN',
            'hardness':'DEFAULT', 'weight':'DEFAULT',
        }
        draft.target={**draft.target,'_provenance':{key:{'origin':origin,'sheet':draft.source_sheet,'cell':draft.source_cell} for key,origin in origins.items()}}
    if semester_start:
        title=' '.join(str(c.value or '') for cells in sheet.iter_rows(max_row=header-1) for c in cells)
        years=re.findall(r'(20\d{2})\s*[-–]\s*(20\d{2})',title)
        if years and any(int(a)!=semester_start.year for a,b in years):
            result.warnings.append({'code':'SEMESTER_METADATA_MISMATCH','message':'Năm học trong tiêu đề khác kỳ học được chọn. Ngày được phân tích theo kỳ học đã chọn.','source_sheet':sheet.title})
    return result


def parse_preference_workbook(path: Path, *, semester_start=None, semester_end=None) -> PreferenceParseResult:
    fmt = detect_preference_format(path)
    if fmt == "GRID_V3":
        from app.parsers.grid_preferences import parse_grid_preferences
        drafts = parse_grid_preferences(path, semester_start=semester_start, semester_end=semester_end)
        return PreferenceParseResult(format="GRID_V3", drafts=drafts, seminars=[], invalid_rows=0)
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        result = _parse_structured(book) if fmt == "STRUCTURED_V2" else _parse_legacy(book, semester_start=semester_start, semester_end=semester_end)
        from app.services.preference_validation import validate_preference_draft
        for draft in result.drafts:
            # Provenance-only notes are already covered by sibling structured
            # rules and are deliberately never candidates for solver apply.
            if draft.constraint_type == "RAW_NOTE" and draft.status == "INTERPRETED":
                continue
            check=validate_preference_draft({'lecturer_id':1 if draft.lecturer_alias or draft.lecturer_code else None,'context_type':draft.context_type,'constraint_type':draft.constraint_type,'day_scope':draft.day_scope,'periods':draft.target.get('periods',[]),'start_date':draft.start_date,'end_date':draft.end_date,'target':draft.target,'numeric_value':draft.numeric_value,'hardness':draft.hardness,'weight':draft.weight,'status':draft.status})
            if not check['is_confirmable']:
                draft.needs_review=True
                if draft.status!='REJECTED': draft.status='NEEDS_REVIEW'
                if draft.confidence_label=='HIGH': draft.confidence_label='LOW';draft.confidence=.35
                draft.review_reason=draft.review_reason or '; '.join(e['message'] for e in check['validation_errors'])
        return result
    finally:
        book.close()


def parse_preferences(path: Path) -> list[ParsedPreference]:
    """Backward-compatible list API; new callers should use parse_preference_workbook."""
    return parse_preference_workbook(path).drafts
