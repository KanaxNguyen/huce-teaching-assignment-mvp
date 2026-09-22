from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlrd


@dataclass
class ParsedSession:
    weekday: int
    start_period: int
    end_period: int
    room: str
    start_date: date | None
    end_date: date | None
    raw_weeks: str
    active_weeks: list[int]
    source_row: int
    source_rows: list[int] = field(default_factory=list)

    def signature(self) -> tuple:
        return (
            self.weekday,
            self.start_period,
            self.end_period,
            self.room.casefold(),
            self.start_date.isoformat() if self.start_date else "",
            self.end_date.isoformat() if self.end_date else "",
            tuple(self.active_weeks),
        )


@dataclass
class ParsedClass:
    course_code: str
    course_name: str
    class_code: str
    credits: float
    lecturer_code: str | None
    lecturer_name: str | None
    locked_assignment: bool
    source_file: str
    source_sheet: str
    source_row: int
    raw_values: dict[str, Any]
    sessions: list[ParsedSession] = field(default_factory=list)
    merged_group_id: str | None = None
    lecturer_evidence: list[dict] = field(default_factory=list)

    @property
    def key(self) -> str:
        return f"{self.course_code}::{self.class_code}"


@dataclass
class MergedGroup:
    id: str
    class_keys: list[str]


@dataclass
class PartialMergeCandidate:
    left_key: str
    right_key: str
    matched_sessions: list[tuple]
    left_only_sessions: list[tuple]
    right_only_sessions: list[tuple]


@dataclass
class ParseIssue:
    severity: str
    code: str
    message: str
    source_file: str
    source_sheet: str
    source_row: int | None = None
    field: str | None = None
    raw_value: str | None = None
    suggestion: str | None = None


@dataclass
class ScheduleParseResult:
    classes: list[ParsedClass]
    merged_groups: list[MergedGroup]
    issues: list[ParseIssue]
    rows_accepted: int
    rows_rejected: int
    header_row: int
    partial_merge_candidates: list[PartialMergeCandidate] = field(default_factory=list)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    return re.sub(r"\s+", " ", text.replace("\r", " ").replace("\n", " ")).strip()


def clean_code(value: Any) -> str:
    text = clean_text(value)
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def is_unassigned_lecturer(value: Any) -> bool:
    """Return true for common placeholders meaning that no teacher is assigned.

    Source workbooks often contain a textual placeholder instead of an empty
    lecturer cell. Treating it as a name locks the class to a fake lecturer and
    prevents the optimiser from assigning a real person.
    """
    text = unicodedata.normalize("NFD", clean_text(value).casefold())
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = text.replace("đ", "d")
    normalized = re.sub(r"[^a-z0-9]+", " ", text).strip()
    compact = normalized.replace(" ", "")

    placeholders = {
        "",
        "-",
        "na",
        "n a",
        "none",
        "null",
        "tbd",
        "chua phan cong",
        "chua duoc phan cong",
        "chua xep giang vien",
        "chua xep gv",
        "chua xep gvien",
        "chua co giang vien",
        "chua co gv",
        "chua bo tri giang vien",
        "dang phan cong",
        "dang cap nhat",
        "chua cap nhat",
    }
    if normalized in placeholders or compact in {"chuaphancong", "chuaxepgiangvien", "chuacogiangvien"}:
        return True
    return normalized.startswith("chua phan") or normalized.startswith("chua xep")


def lecturer_identities(value: Any) -> list[dict]:
    """Keep separate coded identities and common explicit person separators."""
    raw = str(value or "").strip()
    if is_unassigned_lecturer(raw):
        return []
    # Bracketed codes are unambiguous boundaries even without punctuation.
    parts = re.split(r"[,;\n\r]+|\s*/\s*|(?=\[[^\]]+\])", raw)
    identities = []
    for part in parts:
        text = clean_text(part).strip(" ,;/")
        if not text or is_unassigned_lecturer(text):
            continue
        match = re.match(r"^\[([^\]]+)]\s*(.*)$", text)
        if match:
            code, name = clean_code(match[1]), clean_text(match[2])
        else:
            code, name = None, re.sub(r"^(thầy|cô)\s+", "", text, flags=re.IGNORECASE)
        identities.append({"code": code, "name": name or code})
    return identities


def parse_lecturer(value: Any) -> tuple[str | None, str | None]:
    identities = lecturer_identities(value)
    keys = {(i["code"] or clean_text(i["name"]).casefold()) for i in identities}
    if len(keys) != 1:
        return None, None
    return identities[0]["code"], identities[0]["name"]


def parse_periods(value: Any) -> tuple[int, int]:
    numbers = [int(item) for item in re.findall(r"\d+", clean_text(value))]
    if not numbers:
        raise ValueError("Không có tiết học")
    start, end = numbers[0], numbers[-1]
    if start < 1 or end < start or end > 15:
        raise ValueError("Khoảng tiết không hợp lệ")
    return start, end


def parse_weekday(value: Any) -> int:
    text = clean_text(value).casefold()
    if text in {"cn", "chủ nhật", "chu nhat"}:
        return 8
    numbers = re.findall(r"\d+", text)
    if not numbers or not 2 <= int(numbers[0]) <= 8:
        raise ValueError("Thứ phải nằm trong 2–8")
    return int(numbers[0])


def decode_weeks(value: Any) -> tuple[str, list[int]]:
    raw = "" if value is None else unicodedata.normalize("NFC", str(value))
    raw = raw.replace("\r", "").replace("\n", "")
    if not raw:
        raise ValueError("Thiếu tuần học")
    invalid = [char for char in raw if not (char.isspace() or char.isdigit())]
    if invalid:
        raise ValueError(f"Ký tự tuần học không hợp lệ: {''.join(sorted(set(invalid)))}")
    active = [index + 1 for index, char in enumerate(raw) if char.isdigit()]
    if not active:
        raise ValueError("Không giải mã được tuần học")
    return raw, active


def parse_date(value: Any, datemode: int = 0) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return xlrd.xldate_as_datetime(value, datemode).date()
    text = clean_text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Ngày không hợp lệ: {text}")


def _read_rows(path: Path) -> tuple[str, list[list[Any]], int]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        return sheet.name, [sheet.row_values(row)[:14] for row in range(sheet.nrows)], book.datemode
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = [
        [sheet.cell(row, col).value for col in range(1, 15)] for row in range(1, sheet.max_row + 1)
    ]
    return sheet.title, rows, 0


def _find_header(rows: list[list[Any]]) -> int:
    required = {"stt", "tên môn học"}
    for index, row in enumerate(rows[:40]):
        values = {clean_text(value).casefold() for value in row}
        has_code = "mã học phần" in values or "mã hp" in values
        has_class = "mã lớp học" in values
        if required.issubset(values) and has_code and has_class:
            return index
    raise ValueError("Không tìm thấy dòng tiêu đề lịch học")


def parse_schedule(path: Path) -> ScheduleParseResult:
    sheet_name, rows, datemode = _read_rows(path)
    header = _find_header(rows)
    issues: list[ParseIssue] = []
    grouped: dict[str, ParsedClass] = {}
    rejected = 0
    accepted = 0

    # Data may follow immediately after the header, or after a blank/helper
    # row. Validate each row rather than unconditionally skipping one.
    for index, row in enumerate(rows[header + 1 :], start=header + 2):
        padded = (row + [None] * 14)[:14]
        stt, course_code, course_name, class_code = padded[:4]
        if not clean_text(stt) and not any(clean_text(v) for v in padded[:4]):
            continue
        if not clean_code(course_code) or not clean_text(course_name) or not clean_code(class_code):
            if any(clean_text(v) for v in padded[:4]):
                rejected += 1
                issues.append(
                    ParseIssue(
                        "error",
                        "missing_required_field",
                        "Dòng không có đủ mã học phần, tên môn và mã lớp.",
                        path.name,
                        sheet_name,
                        index,
                        suggestion="Bổ sung trường bắt buộc hoặc loại dòng chữ ký.",
                    )
                )
            continue
        try:
            weekday = parse_weekday(padded[5])
            start_period, end_period = parse_periods(padded[6])
            start_date = parse_date(padded[10], datemode)
            end_date = parse_date(padded[11], datemode)
            raw_weeks, active_weeks = decode_weeks(padded[12])
            if start_date and end_date and end_date < start_date:
                raise ValueError("Ngày kết thúc trước ngày bắt đầu")
        except ValueError as error:
            rejected += 1
            issues.append(
                ParseIssue(
                    "error",
                    "invalid_schedule_value",
                    str(error),
                    path.name,
                    sheet_name,
                    index,
                    raw_value=str(padded[5:13]),
                    suggestion="Sửa dữ liệu nguồn rồi upload lại; hệ thống không tự đoán.",
                )
            )
            continue

        code = clean_code(course_code)
        class_id = clean_code(class_code)
        key = f"{code}::{class_id}"
        lecturer_code, lecturer_name = parse_lecturer(padded[13])
        raw_lecturer = clean_text(padded[13])
        credits_text = clean_text(padded[8])
        try:
            credits = float(credits_text) if credits_text else 0
        except ValueError:
            credits = 0
        parsed_session = ParsedSession(
            weekday=weekday,
            start_period=start_period,
            end_period=end_period,
            room=clean_text(padded[7]),
            start_date=start_date,
            end_date=end_date,
            raw_weeks=raw_weeks,
            active_weeks=active_weeks,
            source_row=index, source_rows=[index],
        )
        if key not in grouped:
            grouped[key] = ParsedClass(
                course_code=code,
                course_name=clean_text(course_name),
                class_code=class_id,
                credits=credits,
                lecturer_code=lecturer_code,
                lecturer_name=lecturer_name,
                # Imported teacher is an existing assignment, but only an
                # explicit manager action may lock it for future re-solves.
                locked_assignment=False,
                source_file=path.name,
                source_sheet=sheet_name,
                source_row=index,
                raw_values={"row": [str(v) if v is not None else "" for v in padded]},
            )
        else:
            current = grouped[key]
            normalized_course_name = clean_text(course_name)
            if normalized_course_name.casefold() != current.course_name.casefold() or credits != current.credits:
                issues.append(ParseIssue(
                    "error", "TEACHING_GROUP_INCONSISTENT",
                    f"TeachingGroup {key} có tên môn hoặc số tín chỉ không nhất quán.",
                    path.name, sheet_name, index, "Môn học/Số TC", str(padded[:4]),
                    "Rà soát dữ liệu nguồn trước khi tối ưu.",
                ))
        current = grouped[key]
        current.lecturer_evidence.append({"row": index, "raw": str(padded[13] or ""), "identities": lecturer_identities(padded[13])})
        existing = next((m for m in current.sessions if m.signature() == parsed_session.signature()), None)
        if existing:
            existing.source_rows.append(index)
        else:
            current.sessions.append(parsed_session)
        accepted += 1

    for item in grouped.values():
        identities = [identity for row in item.lecturer_evidence for identity in row["identities"]]
        keys = {identity["code"] or clean_text(identity["name"]).casefold() for identity in identities}
        if len(keys) > 1:
            item.lecturer_code = item.lecturer_name = None
            issues.append(ParseIssue("error", "MULTI_LECTURER_REVIEW",
                f"Lớp {item.class_code} có nhiều danh tính giảng viên; cần xác nhận.",
                path.name, sheet_name, item.source_row, "Giảng viên",
                "\n".join(row["raw"] for row in item.lecturer_evidence),
                "Chọn một giảng viên cho cả nhóm hoặc đánh dấu trường hợp cần hỗ trợ phân đoạn."))
        elif identities:
            item.lecturer_code, item.lecturer_name = identities[0]["code"], identities[0]["name"]

    signature_groups: dict[tuple, list[ParsedClass]] = defaultdict(list)
    for parsed_class in grouped.values():
        # A blank room is insufficient evidence for a merge.  It is common in
        # source exports and must remain a standalone class until reviewed.
        if not parsed_class.sessions or not all(session.room.strip() for session in parsed_class.sessions):
            continue
        full_signature = (
            parsed_class.course_code,
            tuple(sorted(session.signature() for session in parsed_class.sessions)),
        )
        signature_groups[full_signature].append(parsed_class)

    merged_groups = []
    counter = 1
    for candidates in signature_groups.values():
        if len(candidates) < 2:
            continue
        group_id = f"MG-{counter:03d}"
        counter += 1
        for item in candidates:
            item.merged_group_id = group_id
        merged_groups.append(MergedGroup(group_id, sorted(item.key for item in candidates)))

    partial_merge_candidates: list[PartialMergeCandidate] = []
    by_course: dict[str, list[ParsedClass]] = defaultdict(list)
    for item in grouped.values():
        by_course[item.course_code].append(item)
    for items in by_course.values():
        for index, left in enumerate(items):
            left_sessions = {session.signature() for session in left.sessions if session.room.strip()}
            if not left_sessions:
                continue
            for right in items[index + 1:]:
                right_sessions = {session.signature() for session in right.sessions if session.room.strip()}
                matched = left_sessions.intersection(right_sessions)
                if not matched or left_sessions == right_sessions:
                    continue
                candidate = PartialMergeCandidate(
                    left_key=left.key,
                    right_key=right.key,
                    matched_sessions=sorted(matched),
                    left_only_sessions=sorted(left_sessions - right_sessions),
                    right_only_sessions=sorted(right_sessions - left_sessions),
                )
                partial_merge_candidates.append(candidate)
                issues.append(ParseIssue(
                    "warning", "PARTIAL_MERGE_CANDIDATE",
                    f"{left.class_code} và {right.class_code} chỉ trùng một phần lịch; không tự ghép.",
                    path.name, sheet_name, suggestion="Mở mục Merge review để xác nhận từng buổi trùng.",
                ))

    return ScheduleParseResult(
        classes=list(grouped.values()),
        merged_groups=merged_groups,
        issues=issues,
        rows_accepted=accepted,
        rows_rejected=rejected,
        header_row=header + 1,
        partial_merge_candidates=partial_merge_candidates,
    )


def sessions_overlap(left: ParsedSession, right: ParsedSession) -> bool:
    if left.weekday != right.weekday:
        return False
    if left.end_period < right.start_period or right.end_period < left.start_period:
        return False
    if not set(left.active_weeks).intersection(right.active_weeks):
        return False
    if left.start_date and right.end_date and left.start_date > right.end_date:
        return False
    if right.start_date and left.end_date and right.start_date > left.end_date:
        return False
    return True
