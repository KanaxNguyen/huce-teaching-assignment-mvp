from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func

from app.db.session import Base, engine, SessionLocal
from app.main import app
from app.models.entities import (
    Assignment, ClassSection, ClassSession, Constraint, Course, HistoricalEvidence,
    Lecturer, LecturerCourseCapability, LecturerSemesterProfile,
    NormalizedPreferenceDraft, OutputTemplateProfile, Semester,
)
from app.optimization.solver import solve, _slot_match
from app.optimization.occurrences import meeting_occurrences
from app.parsers.preferences import _legacy_rule, parse_preference_workbook
from app.services.manual_assignment import check_assignment_change
from app.services.importer import _import_files_impl as import_files
from app.services.template_detector import detect_output_template
from app.services.historical_learning import learn_from_template_profile


@pytest.fixture(autouse=True)
def database():
    Base.metadata.create_all(engine)
    yield
    Base.metadata.drop_all(engine)


def semester(db, name="A"):
    item = Semester(name=name, department_name="Toán", head_name="UAT", start_date=date(2026,8,3), end_date=date(2027,1,31))
    db.add(item); db.flush()
    return item


def setup_solver(db):
    sem = semester(db)
    teacher = Lecturer(code="GV_A", canonical_name="Lecturer A", confirmed=True)
    course = Course(code="MATH", name="Math")
    db.add_all([teacher, course]); db.flush()
    db.add(LecturerCourseCapability(lecturer_id=teacher.id, course_id=course.id, allowed=True, confirmed=True))
    return sem, teacher, course


def group(db, sem, course, weekday, first, last, index, weeks=None, start=None, end=None):
    item = ClassSection(semester_id=sem.id, course_id=course.id, class_code=f"G{index}", source_file="test", source_sheet="S", source_row=index)
    db.add(item); db.flush()
    db.add(ClassSession(class_id=item.id, weekday=weekday, start_period=first, end_period=last, room="R", raw_weeks="", active_weeks=weeks if weeks is not None else [1], start_date=start, end_date=end, source_row=index))
    db.flush()
    return item


def workbook(rows):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Mã GV", "Họ và tên", "Tham gia kỳ này"])
    for row in rows: ws.append(row)
    stream = BytesIO(); wb.save(stream)
    return {"file": ("lecturers.xlsx", stream.getvalue())}


def test_review_has_only_semester_relevant_lecturers():
    with SessionLocal() as db:
        a, b = semester(db), semester(db, "B")
        la = Lecturer(code="A", canonical_name="A", confirmed=True)
        lb = Lecturer(canonical_name="B missing code")
        db.add_all([la,lb]);db.flush()
        db.add_all([LecturerSemesterProfile(semester_id=a.id,lecturer_id=la.id), LecturerSemesterProfile(semester_id=b.id,lecturer_id=lb.id)])
        db.commit(); aid,bid=a.id,b.id
    client=TestClient(app)
    left=client.get(f"/api/v1/lecturers/review?semester_id={aid}").json()
    right=client.get(f"/api/v1/lecturers/review?semester_id={bid}").json()
    assert [x['name'] for x in left['items']] == ['A']
    assert left['blockers_count'] == 0
    assert right['blockers_count'] == 1


@pytest.mark.parametrize("issue", ["duplicate", "missing_name", "conflicting_code"])
def test_import_gate_rejects_entire_batch(issue):
    with SessionLocal() as db:
        sem=semester(db)
        db.add(Lecturer(code="OLD",canonical_name="Same"))
        if issue == 'duplicate': db.add(Lecturer(code="OTHER",canonical_name="Same"))
        db.commit(); sid=sem.id
    bad = [None,'Same','Có'] if issue=='duplicate' else ['BAD',None,'Có'] if issue=='missing_name' else ['NEW','Same','Có']
    files=workbook([['SAFE','Safe','Có'],bad])
    client=TestClient(app)
    preview=client.post(f'/api/v1/lecturers/import/preview?semester_id={sid}',files=files)
    assert preview.status_code==200 and not preview.json()['can_commit']
    response=client.post(f'/api/v1/lecturers/import?semester_id={sid}',files=files)
    assert response.status_code==422
    with SessionLocal() as db:
        assert db.scalar(select(Lecturer).where(Lecturer.code=='SAFE')) is None
        assert db.scalar(select(func.count(LecturerSemesterProfile.id)))==0


@pytest.mark.parametrize('value',['Không','KHÔNG','Khong'])
def test_vietnamese_participation_import(value):
    with SessionLocal() as db:
        sem=semester(db);db.commit();sid=sem.id
    response=TestClient(app).post(f'/api/v1/lecturers/import?semester_id={sid}',files=workbook([['NEW','New',value]]))
    assert response.status_code==200
    with SessionLocal() as db:
        assert db.scalar(select(LecturerSemesterProfile)).participation_status=='NOT_PARTICIPATING'


def test_import_forced_commit_failure_rolls_back(monkeypatch):
    import app.api.lecturer_routes as routes
    with SessionLocal() as db:
        sem=semester(db);db.commit();sid=sem.id
    def fail(db):
        db.flush()
        assert db.scalar(select(Lecturer).where(Lecturer.code=='NEW')) is not None
        raise RuntimeError('forced transaction failure')
    monkeypatch.setattr(routes,'commit',fail)
    response=TestClient(app).post(f'/api/v1/lecturers/import?semester_id={sid}',files=workbook([['NEW','New','Có']]))
    assert response.status_code==422
    with SessionLocal() as db:
        assert db.scalar(select(func.count(Lecturer.id)))==0
        assert db.scalar(select(func.count(LecturerSemesterProfile.id)))==0


def test_rejected_must_restore_before_confirmation():
    with SessionLocal() as db:
        sem,teacher,_=setup_solver(db)
        draft=NormalizedPreferenceDraft(semester_id=sem.id,lecturer_id=teacher.id,constraint_type='UNAVAILABLE',context_type='TEACHING',day_scope='T2',periods=[1,2,3],source_file='test',source_sheet='S',source_row=1,source_cell='C1',status='DRAFT',needs_review=False)
        db.add(draft);db.commit();sid,did=sem.id,draft.id
    client=TestClient(app);url=f'/api/v1/preference-drafts/{did}?semester_id={sid}'
    assert client.patch(url,json={'status':'REJECTED','rejected_reason':'Wrong'}).status_code==200
    assert client.patch(url,json={'status':'CONFIRMED'}).status_code==409
    assert client.post(f'/api/v1/preference-drafts/apply?semester_id={sid}',json={'draft_ids':[did]}).status_code==422
    with SessionLocal() as db:
        draft=db.get(NormalizedPreferenceDraft,did)
        assert draft.status=='REJECTED' and draft.rejected_at and draft.rejected_reason=='Wrong'
    restored=client.patch(url,json={'status':'NEEDS_REVIEW'}).json()
    assert restored['status']=='NEEDS_REVIEW' and restored['rejected_at'] is None and restored['rejected_reason'] is None
    assert client.patch(url,json={'status':'CONFIRMED'}).status_code==200
    assert client.post(f'/api/v1/preference-drafts/apply?semester_id={sid}',json={'draft_ids':[did]}).status_code==200


@pytest.mark.parametrize('text',['T2-T6','Thứ 2 - Thứ 6','từ thứ 2 đến thứ 6'])
def test_weekday_range_not_a_period_range(text):
    kind,target,confidence,_=_legacy_rule(f'Xin dạy tiết 7-9 {text}',None)
    assert kind=='PREFERRED_PERIOD' and confidence=='HIGH'
    assert [s['weekday'] for s in target['slots']]==[2,3,4,5,6]
    assert target['periods']==[7,8,9]


def test_all_days_includes_weekend():
    kind,target,conf,_=_legacy_rule('Không dạy tiết 1-3 tất cả các ngày',None)
    assert kind=='UNAVAILABLE' and conf=='HIGH'
    assert target['day_scope']=='ALL_DAYS'
    assert [s['weekday'] for s in target['slots']]==list(range(2,9))


@pytest.mark.parametrize('start',[7,10])
def test_open_ended_periods(start):
    assert _legacy_rule(f'Có thể dạy từ tiết {start} trở đi',5)[1]['periods']==list(range(start,16))


def test_alternative_ranges_survive_confirm_and_apply():
    kind,target,conf,_=_legacy_rule('Xin dạy tiết 4-6 HOẶC 7-9 T2-T6',None)
    assert target['periods']==[] and target['period_alternatives']==[[4,5,6],[7,8,9]]
    with SessionLocal() as db:
        sem,teacher,_=setup_solver(db)
        draft=NormalizedPreferenceDraft(semester_id=sem.id,lecturer_id=teacher.id,constraint_type=kind,context_type='TEACHING',day_scope='ALL_WEEKDAYS',periods=[],target=target,source_file='test',source_sheet='S',source_row=1,source_cell='C1',status='DRAFT',needs_review=False)
        db.add(draft);db.commit();sid,did=sem.id,draft.id
    client=TestClient(app)
    response=client.patch(f'/api/v1/preference-drafts/{did}?semester_id={sid}',json={'status':'CONFIRMED'})
    assert response.status_code==200
    assert client.post(f'/api/v1/preference-drafts/apply?semester_id={sid}',json={'draft_ids':[did]}).status_code==200
    with SessionLocal() as db:
        saved=db.scalar(select(Constraint))
        assert len(saved.target['slots'])==10
        assert {tuple(s['periods']) for s in saved.target['slots']}=={(4,5,6),(7,8,9)}


@pytest.mark.parametrize('text',['Các thời gian còn lại, kể cả tối, T7 và CN','Có thể dạy vài buổi chiều','Nghỉ đến 11/10'])
def test_unsupported_never_high(text):
    kind,_,confidence,reason=_legacy_rule(text,None)
    assert kind=='RAW_NOTE' and confidence!='HIGH' and reason


@pytest.mark.parametrize('year',[2026,2028])
def test_year_comes_from_semester(year):
    for text in ['KHÔNG DẠY ĐẾN: 11/10','nghỉ đến 11/10','bận đến 11 tháng10',f'không thể dạy trước 12/10/{year}']:
        kind,target,_,_=_legacy_rule(text,None,semester_start=date(year,8,1),semester_end=date(year+1,1,31))
        assert kind=='UNAVAILABLE' and target['start_date']==f'{year}-08-01' and target['end_date']==f'{year}-10-11'


def test_occurrences_and_manual_check_agree():
    with SessionLocal() as db:
        sem,teacher,course=setup_solver(db)
        # Week 10 Sunday is 11 October; week 11 Monday is 12 October.
        a=group(db,sem,course,8,1,3,1,weeks=[10])
        b=group(db,sem,course,2,1,3,2,weeks=[11])
        phantom=group(db,sem,course,8,1,3,3,weeks=[10,15],start=date(2026,10,12),end=date(2026,11,30))
        target={'end_date':'2026-10-11','periods':[1,2,3]}
        db.add(Constraint(semester_id=sem.id,lecturer_id=teacher.id,name='Until',constraint_type='unavailable',hardness='hard',target=target,confirmed=True,active=True));db.commit()
        assert _slot_match(a.sessions[0],target,sem)
        assert not _slot_match(b.sessions[0],target,sem)
        assert not _slot_match(phantom.sessions[0],{'start_date':'2026-10-01','end_date':'2026-10-15'},sem)
        assert not check_assignment_change(db,sem.id,a.id,teacher.id)['valid']
        assert check_assignment_change(db,sem.id,b.id,teacher.id)['valid']
        run=solve(db,3,False,sem.id)
        assigned=set(db.scalars(select(Assignment.class_id).where(Assignment.run_id==run.id)))
        assert a.id not in assigned and b.id in assigned


@pytest.mark.parametrize('periods,weight,gap',[
    ([(1,3),(4,6)],.8,0), ([(1,2),(4,6)],.8,1),
    ([(1,3),(7,9)],.8,3), ([(7,9),(13,15)],.8,3),
    ([(1,3),(4,6),(10,12)],.8,3), ([(1,3)],.8,0),
    ([(1,5),(7,9)],.8,1), ([(2,6),(7,9)],.8,0),
    ([(1,3),(7,9)],0,0),
])
def test_exact_period_gap(periods,weight,gap):
    with SessionLocal() as db:
        sem,teacher,course=setup_solver(db)
        for i,(first,last) in enumerate(periods): group(db,sem,course,2,first,last,i+1)
        db.add(Constraint(semester_id=sem.id,lecturer_id=teacher.id,name='Compact',constraint_type='PREFER_CONSECUTIVE_PERIODS',hardness='soft',weight=weight,target={'value':6},active=True,confirmed=True));db.commit()
        run=solve(db,3,False,sem.id)
        assert run.score==gap*80
        assert db.scalar(select(func.count(Assignment.id)).where(Assignment.run_id==run.id))==len(periods)


@pytest.mark.parametrize('count,periods,invalid_dates,scoped_out,assigned',[
    (4,(1,3),False,False,4), (5,(1,3),False,False,4),
    (5,(7,12),False,False,5), (5,(1,3),True,False,5),
    (5,(1,3),False,True,5),
])
def test_free_morning_real_occurrences(count,periods,invalid_dates,scoped_out,assigned):
    with SessionLocal() as db:
        sem,teacher,course=setup_solver(db)
        for i in range(count):
            group(db,sem,course,i+2,*periods,i+1,start=date(2026,9,1) if invalid_dates else None)
        target={'value':1}
        if scoped_out: target['start_date']='2026-09-01'
        db.add(Constraint(semester_id=sem.id,lecturer_id=teacher.id,name='Free morning',constraint_type='MIN_FREE_MORNING_PER_WEEK',hardness='hard',target=target,confirmed=True,active=True));db.commit()
        run=solve(db,3,False,sem.id)
        assert db.scalar(select(func.count(Assignment.id)).where(Assignment.run_id==run.id))==assigned


ROOT=Path('/Users/mac/AI/Huce_timetable')
ORIGINAL=ROOT/'INPUT GỐC/Nguyện vọng TKB 2026-2027.xlsx'
SCHEDULE=ROOT/'INPUT GỐC/Phan_Cong_Giang_Day_HK1_2026_2027.xlsx'
HISTORY=ROOT/'phan-cong-giang-day-lich-hoc-to-bo-mon-03-09-2026.xls'


def test_exact_original_files_and_historical_scope():
    for path in (ORIGINAL,SCHEDULE,HISTORY):
        if not path.is_file(): pytest.skip(f'Exact private fixture unavailable: {path}')
    parsed=parse_preference_workbook(ORIGINAL,semester_start=date(2026,8,3),semester_end=date(2027,1,31))
    by_cell={d.source_cell:d for d in parsed.drafts}
    assert parsed.dropped_clauses==0 and parsed.raw_clauses==38 and len(parsed.drafts)==39
    assert by_cell['J4#3'].needs_review
    assert by_cell['C16#1'].target['period_alternatives']==[[4,5,6],[7,8,9]]
    assert by_cell['F6#2'].target['weekday']==5
    assert by_cell['F6#2'].needs_review # afternoon vs open evening range
    assert by_cell['C7#1A'].target['weekdays']==[2,4]
    assert by_cell['C7#1B'].target['weekdays']==[3,5,6,7]
    assert by_cell['C14#2'].needs_review # "hoặc dạy ít" cannot be silently dropped
    with SessionLocal() as db:
        sem=semester(db);db.commit()
        detection=detect_output_template(HISTORY)
        assert detection['header_start_row']==8 and detection['header_end_row']==9
        assert len(detection['mappings'])==14
        assert detection['historical_summary']['scope']=='DETECTION_SAMPLE'
        profile=OutputTemplateProfile(semester_id=sem.id,source_file=str(HISTORY),source_sheet=detection['source_sheet'],header_row=detection['header_row'],mappings=detection['mappings'])
        db.add(profile);db.commit()
        learned=learn_from_template_profile(db,profile,HISTORY)
        assert learned['workbook_rows']==335 and learned['evidence_count']==332
        assert learned['scope']=='FULL_FILE_HISTORICAL_LEARNING'
        assert learned['learned_fields']==['identity','course','class']
        assert db.scalar(select(func.count(ClassSection.id)))==0
        assert db.scalar(select(func.count(Assignment.id)))==0
        result=import_files(db,[SCHEDULE,ORIGINAL],semester_id=sem.id,schedule_paths=[SCHEDULE],preference_paths=[ORIGINAL])
        assert result['summary']['classes']==158 and result['summary']['sessions']==286
        assert db.scalar(select(func.count(ClassSection.id)).where(ClassSection.locked_assignment.is_(True)))==0


def test_unsupported_and_hypothetical_capabilities_cannot_enter_candidates():
    with SessionLocal() as db:
        sem = semester(db)
        teacher_valid = Lecturer(code="GV_VAL", canonical_name="GV Valid", confirmed=True)
        teacher_hypo = Lecturer(code="GV_HYP", canonical_name="GV Hypo", confirmed=True)
        teacher_unconf = Lecturer(code="GV_UNC", canonical_name="GV Unconfirmed", confirmed=True)
        teacher_disallow = Lecturer(code="GV_DIS", canonical_name="GV Disallowed", confirmed=True)
        course = Course(code="MATH_CAP", name="Math Cap")
        db.add_all([teacher_valid, teacher_hypo, teacher_unconf, teacher_disallow, course])
        db.flush()

        # Valid confirmed capability
        c_valid = LecturerCourseCapability(lecturer_id=teacher_valid.id, course_id=course.id, allowed=True, confirmed=True, source="SCHEDULE_IMPORT")
        # HYPOTHETICAL_ALL source
        c_hypo = LecturerCourseCapability(lecturer_id=teacher_hypo.id, course_id=course.id, allowed=True, confirmed=True, source="HYPOTHETICAL_ALL")
        # Unconfirmed capability
        c_unconf = LecturerCourseCapability(lecturer_id=teacher_unconf.id, course_id=course.id, allowed=True, confirmed=False, source="SCHEDULE_IMPORT")
        # Disallowed capability
        c_disallow = LecturerCourseCapability(lecturer_id=teacher_disallow.id, course_id=course.id, allowed=False, confirmed=True, source="SCHEDULE_IMPORT")
        db.add_all([c_valid, c_hypo, c_unconf, c_disallow])
        db.flush()

        # Class section
        grp = group(db, sem, course, 2, 1, 3, 1)

        from app.optimization.solver import eligible_teachers
        caps = {course.id: [c_valid, c_hypo, c_unconf, c_disallow]}
        candidates = eligible_teachers(grp, caps)
        assert candidates == {teacher_valid.id}
        assert teacher_hypo.id not in candidates
        assert teacher_unconf.id not in candidates
        assert teacher_disallow.id not in candidates

        # Solve and verify only teacher_valid is assigned
        run = solve(db, 3, False, sem.id)
        assert run.status in {"optimal", "feasible"}
        assigned = db.scalars(select(Assignment).where(Assignment.run_id == run.id)).all()
        assert len(assigned) == 1
        assert assigned[0].lecturer_id == teacher_valid.id

        # Verify that HYPOTHETICAL_ALL was disabled in the database
        db.refresh(c_hypo)
        assert c_hypo.confirmed is False
        assert c_hypo.allowed is False


def test_dashboard_alignment_real_dataset():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.services.importer import dashboard
    from app.services.source_authority import latest_current_run
    db_path = ROOT / "huce-teaching-assignment-mvp/storage/database/huce_test.db"
    if not db_path.is_file():
        pytest.skip(f"Real DB fixture not available: {db_path}")
    real_engine = create_engine(f"sqlite:///{db_path}")
    with Session(real_engine) as db:
        run = latest_current_run(db, 4)
        d = dashboard(db, 4)
        assert d["classes"] == 158
        if run:
            assigned_ids = set(db.scalars(select(Assignment.class_id).where(Assignment.run_id == run.id)).all())
            assert d["unassigned_classes"] == 158 - len(assigned_ids)
            assert d["optimization_status"] == run.status

        # Historical run 14 explicit check (2 unassigned)
        d14 = dashboard(db, 4, run_id=14)
        assert d14["classes"] == 158
        assert d14["unassigned_classes"] == 2
        assert d14["optimization_status"] == "feasible"


def test_export_completeness_includes_all_158_groups_and_supplementary_71kme(tmp_path):
    for path in (ORIGINAL, SCHEDULE, HISTORY):
        if not path.is_file():
            pytest.skip(f"Exact private fixture unavailable: {path}")
    with SessionLocal() as db:
        sem = semester(db)
        db.commit()
        import_files(db, [SCHEDULE, ORIGINAL], semester_id=sem.id, schedule_paths=[SCHEDULE], preference_paths=[ORIGINAL])
        detection = detect_output_template(HISTORY)
        profile = OutputTemplateProfile(
            semester_id=sem.id,
            source_file=str(HISTORY),
            source_sheet=detection["source_sheet"],
            header_row=detection["header_row"],
            mappings=detection["mappings"],
        )
        db.add(profile)
        db.commit()

        run = solve(db, 10, False, sem.id)
        assert run.status in {"optimal", "feasible"}

        from app.exporters.excel import export_latest
        export_path = export_latest(db, tmp_path / "exports", sem.id, mode="draft")
        wb = openpyxl.load_workbook(export_path)
        sheet = wb[profile.source_sheet] if profile.source_sheet in wb.sheetnames else wb.active

        class_col = detection["mappings"]["class_code"]["column_index"]
        course_col = detection["mappings"]["course_code"]["column_index"]

        all_db_groups = set(
            db.execute(
                select(Course.code, ClassSection.class_code)
                .join(Course, Course.id == ClassSection.course_id)
                .where(ClassSection.semester_id == sem.id)
            ).all()
        )
        assert len(all_db_groups) == 158

        template_groups = set()
        supplementary_groups = set()
        in_supplementary = False

        for r in range(profile.header_row + 1, sheet.max_row + 1):
            c1 = str(sheet.cell(r, 1).value or "")
            if "DANH SÁCH LỚP BỔ SUNG" in c1:
                in_supplementary = True
                continue
            c_val = sheet.cell(r, course_col).value
            cls_val = sheet.cell(r, class_col).value
            if c_val and cls_val:
                pair = (str(c_val).strip(), str(cls_val).strip())
                if pair in all_db_groups:
                    if in_supplementary:
                        supplementary_groups.add(pair)
                    else:
                        template_groups.add(pair)

        assert in_supplementary, "Header 'DANH SÁCH LỚP BỔ SUNG' must be present"
        assert ("390111", "71KME") in supplementary_groups, "Class 71KME must be present in supplementary export"
        assert len(template_groups) == 157
        assert len(supplementary_groups) == 1
        assert template_groups & supplementary_groups == set()
        assert (template_groups | supplementary_groups) == all_db_groups


def test_every_one_of_the_158_teaching_groups_appears_exactly_once_in_export(tmp_path):
    """Priority 3: Export must never silently omit a TeachingGroup and every group appears in detailed export."""
    for path in (ORIGINAL, SCHEDULE):
        if not path.is_file():
            pytest.skip(f"Exact private fixture unavailable: {path}")
    with SessionLocal() as db:
        sem = semester(db)
        db.commit()
        import_files(db, [SCHEDULE, ORIGINAL], semester_id=sem.id, schedule_paths=[SCHEDULE], preference_paths=[ORIGINAL])
        all_db_groups = set(
            db.execute(
                select(Course.code, ClassSection.class_code)
                .join(Course, Course.id == ClassSection.course_id)
                .where(ClassSection.semester_id == sem.id)
            ).all()
        )
        assert len(all_db_groups) == 158

        run = solve(db, 10, False, sem.id)
        assert run.status in {"optimal", "feasible"}

        from app.exporters.excel import export_latest
        export_path = export_latest(db, tmp_path / "detailed_exports", sem.id, mode="draft")
        wb = openpyxl.load_workbook(export_path)
        detail = wb["Phân công"]

        # In Phân công sheet: Col 2 is Mã học phần, Col 4 is Mã lớp
        exported_groups = set()
        for r in range(3, detail.max_row + 1):
            c_code = detail.cell(r, 2).value
            cls_code = detail.cell(r, 4).value
            if c_code and cls_code:
                pair = (str(c_code).strip(), str(cls_code).strip())
                if pair in all_db_groups:
                    exported_groups.add(pair)

        assert exported_groups == all_db_groups
        assert len(exported_groups) == 158
