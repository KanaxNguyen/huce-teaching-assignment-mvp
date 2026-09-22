from datetime import date
from pathlib import Path
import hashlib
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from app.db.session import Base, SessionLocal, engine
from app.main import app
from app.models.entities import (Assignment, ClassSection, ClassSession, Constraint, Lecturer,
    LecturerAlias, LecturerCourseCapability, NormalizedPreferenceDraft, OptimizationRun, Semester, SourceVersion, ValidationIssue)
from app.services.source_authority import (stage_source, preview_activation, activate_source, SourceError,
    latest_current_run, resolve_issue, meeting_diff)
from app.services import source_authority as authority
from app.storage.backends import LocalStorage
from app.optimization.solver import solve
from app.api.routes import get_assignments, get_dashboard


@pytest.fixture(autouse=True)
def database(tmp_path, monkeypatch):
    Base.metadata.drop_all(engine); Base.metadata.create_all(engine)
    monkeypatch.setattr(authority, 'get_storage_backend', lambda: LocalStorage(tmp_path/'objects'))
    yield
    with engine.begin() as connection:
        connection.execute(Semester.__table__.update().values(active_schedule_source_id=None, active_preference_source_id=None))
    Base.metadata.drop_all(engine)


def semester(db, name='S'):
    s=Semester(name=name,department_name='D',head_name='H',start_date=date(2026,8,3),end_date=date(2027,1,31))
    db.add(s);db.commit();return s


def workbook(path, teachers=('[A]Teacher Alpha',), rooms=None):
    wb=Workbook();ws=wb.active;ws.title='Schedule'
    ws.append(['STT','Mã học phần','Tên môn học','Mã lớp học','','Thứ','Tiết','Phòng','Số TC','','Bắt đầu','Kết thúc','Tuần học','Giảng viên'])
    for i, teacher in enumerate(teachers):
        ws.append([i+1,'M','Math','L','',2,'1-3',(rooms or ['R']*len(teachers))[i],3,'','03/08/2026','31/01/2027','123',teacher])
    wb.save(path);return path


def preferences(path, text='Xin nghỉ Thứ 2 tiết 1-3'):
    wb=Workbook();ws=wb.active;ws.append(['Họ và tên','Ghi chú']);ws.append(['Teacher Alpha',text]);wb.save(path);return path


def v3_preferences(path):
    wb=Workbook();ws=wb.active;ws.title='01_Teacher Alpha'
    ws['A1']='[HUCE-FORM-PREF-V3.0]'
    ws['A8']='A';ws['E8']='Teacher Alpha'
    ws['C14']='Bận cứng';ws['D14']='Ưu tiên';ws['E14']='Hạn chế'
    ws['F22']='Rất mong muốn'
    wb.save(path)
    return path


def activate(db, s, path, role='CURRENT_SCHEDULE'):
    source=stage_source(db,path,s.id,role);db.commit()
    preview=preview_activation(db,s.id,source.id,role)
    activate_source(db,s.id,source.id,role,confirmed=True,preview_token=preview['preview_token'])
    return source


def test_active_v3_source_creates_canonical_drafts_and_survives_reload(tmp_path):
    with SessionLocal() as db:
        s=semester(db)
        sid=s.id
        schedule=activate(db,s,workbook(tmp_path/'schedule.xlsx'))
        preference=activate(db,s,v3_preferences(tmp_path/'v3.xlsx'),'PREFERENCE')
        lecturer=db.scalar(select(Lecturer).where(Lecturer.code=='A'))
        assert db.query(ClassSection).filter_by(semester_id=sid).count()==1
        drafts=db.scalars(select(NormalizedPreferenceDraft).where(
            NormalizedPreferenceDraft.semester_id==sid,
            NormalizedPreferenceDraft.source_version_id==preference.id,
        )).all()
        assert {d.constraint_type for d in drafts}=={
                'UNAVAILABLE','PREFERRED_PERIOD','AVOID_PERIOD','PREFER_COMPACT_SCHEDULE',
        }
        assert all(d.lecturer_id==lecturer.id for d in drafts)
        assert db.get(Semester,sid).active_schedule_source_id==schedule.id
        assert db.get(Semester,sid).active_preference_source_id==preference.id

    with SessionLocal() as reloaded:
        active=reloaded.get(Semester,sid)
        assert (active.active_schedule_source_id,active.active_preference_source_id)==(schedule.id,preference.id)
        assert reloaded.query(NormalizedPreferenceDraft).filter_by(semester_id=sid,source_version_id=preference.id).count()==4
        assert get_dashboard(sid,reloaded)['classes']==1

    client=TestClient(app)
    sources=client.get(f'/api/v1/sources?semester_id={sid}').json()
    assert sources['active_schedule_source_id']==schedule.id
    assert sources['active_preference_source_id']==preference.id
    review=client.get(f'/api/v1/preference-drafts?semester_id={sid}').json()
    assert len(review)==4
    assert {item['lecturer_id'] for item in review}=={lecturer.id}


def test_upload_pair_activate_then_reload_keeps_preference_review_populated(tmp_path):
    """Regression: staging two uploads cannot make Preference Review look ready with zero drafts."""
    with SessionLocal() as db:
        sem = semester(db)
        sid = sem.id
    schedule_path = workbook(tmp_path / "schedule.xlsx")
    preference_path = v3_preferences(tmp_path / "preferences.xlsx")
    client = TestClient(app)
    upload = client.post(
        f"/api/v1/imports/upload-pair?semester_id={sid}",
        files={
            "schedule_file": (schedule_path.name, schedule_path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "preference_file": (preference_path.name, preference_path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        },
    )
    assert upload.status_code == 200
    staged = upload.json()["sources"]
    assert {item["source_type"] for item in staged} == {"CURRENT_SCHEDULE", "PREFERENCE"}

    before_activation = client.get(f"/api/v1/sources?semester_id={sid}").json()
    assert before_activation["active_schedule_source_id"] is None
    assert before_activation["active_preference_source_id"] is None

    for source in staged:
        preview = client.post(
            f"/api/v1/sources/{source['id']}/preview?semester_id={sid}",
            json={"source_type": source["source_type"]},
        )
        assert preview.status_code == 200
        activated = client.post(
            f"/api/v1/sources/{source['id']}/activate?semester_id={sid}",
            json={"source_type": source["source_type"], "confirmed": True, "preview_token": preview.json()["preview_token"]},
        )
        assert activated.status_code == 200

    reloaded_sources = client.get(f"/api/v1/sources?semester_id={sid}").json()
    assert reloaded_sources["active_schedule_source_id"]
    assert reloaded_sources["active_preference_source_id"]
    drafts = client.get(f"/api/v1/preference-drafts?semester_id={sid}").json()
    assert len(drafts) == 4
    assert all(item["lecturer_id"] is not None for item in drafts)


def test_hash_bytes_immutable_upload_order_and_roles(tmp_path):
    p=workbook(tmp_path/'same.xlsx')
    with SessionLocal() as db:
        s=semester(db)
        a=stage_source(db,p,s.id,'CURRENT_SCHEDULE');db.commit()
        identical=stage_source(db,p,s.id,'CURRENT_SCHEDULE');db.commit()
        assert identical.id == a.id and a.content_hash == hashlib.sha256(p.read_bytes()).hexdigest()
        original=p.read_bytes();workbook(p,rooms=['DIFFERENT'])
        b=stage_source(db,p,s.id,'CURRENT_SCHEDULE');db.commit()
        assert b.id!=a.id and b.content_hash!=a.content_hash
        assert s.active_schedule_source_id is None and db.query(ClassSection).count()==0
        with authority.get_storage_backend().materialize(a.storage_ref) as saved: assert saved.read_bytes()==original
        pref=stage_source(db,preferences(tmp_path/'pref.xlsx'),s.id,'PREFERENCE');db.commit()
        assert pref.source_type=='PREFERENCE' and s.active_preference_source_id is None


@pytest.mark.parametrize('role',['HISTORICAL','OUTPUT_TEMPLATE','REFERENCE_MATRIX'])
def test_reference_cannot_activate_as_schedule(tmp_path,role):
    with SessionLocal() as db:
        s=semester(db);src=stage_source(db,workbook(tmp_path/'history.xlsx'),s.id,role);db.commit()
        with pytest.raises(SourceError,match='Vai trò'): preview_activation(db,s.id,src.id,'CURRENT_SCHEDULE')
        with pytest.raises(SourceError): activate_source(db,s.id,src.id,role,confirmed=True,preview_token='x')
        assert s.active_schedule_source_id is None


def test_confirmation_semester_hash_and_stale_preview(tmp_path):
    with SessionLocal() as db:
        s=semester(db);other=semester(db,'Other');src=stage_source(db,workbook(tmp_path/'s.xlsx'),s.id,'CURRENT_SCHEDULE');db.commit()
        with pytest.raises(SourceError): preview_activation(db,other.id,src.id,'CURRENT_SCHEDULE')
        preview=preview_activation(db,s.id,src.id,'CURRENT_SCHEDULE')
        with pytest.raises(SourceError): activate_source(db,s.id,src.id,'CURRENT_SCHEDULE',confirmed=False,preview_token=preview['preview_token'])
        assert db.query(ClassSection).count()==0
        s.source_revision+=1;db.commit()
        with pytest.raises(SourceError): activate_source(db,s.id,src.id,'CURRENT_SCHEDULE',confirmed=True,preview_token=preview['preview_token'])
        with authority.get_storage_backend().materialize(src.storage_ref) as p: p.write_bytes(b'tamper')
        with pytest.raises(SourceError,match='hash'): preview_activation(db,s.id,src.id,'CURRENT_SCHEDULE')


def test_diff_all_categories_and_changed_identity_fields():
    base={'course':'M','class_code':'L','day':2,'period':[1,3],'room':'R','week_mask':[1], 'date_range':[None,None],'source_rows':[2]}
    old=[base,{**base,'class_code':'CHANGED'},{**base,'class_code':'REMOVE','day':4,'room':'X'}]
    new=[base,{**base,'class_code':'CHANGED','room':'NEW'},{**base,'class_code':'ADD','day':5,'room':'Y'}]
    diff=meeting_diff(old,new)
    assert diff['counts']=={'UNCHANGED':1,'CHANGED':1,'REMOVED':1,'ADDED':1}
    for field,value in [('course','N'),('class_code','NEW'),('day',3),('period',[4,6]),('room','NEW'),('week_mask',[2]),('date_range',['2026-09-01',None])]:
        result=meeting_diff([base],[{**base,field:value}])
        assert result['counts']['CHANGED']==1 and field in result['meetings'][0]['changed_fields']


@pytest.mark.parametrize('teacher',[
    '[A]Teacher Alpha, [B]Teacher Beta', '[A]Teacher Alpha\n[B]Teacher Beta',
    '[A]Teacher Alpha; [B]Teacher Beta', '[A]Teacher Alpha [B]Teacher Beta',
    '[A]Teacher Alpha / [B]Teacher Beta', 'Teacher Alpha\nTeacher Beta',
])
def test_multi_lecturer_quarantine_and_backend_solve_blocker(tmp_path,teacher):
    with SessionLocal() as db:
        s=semester(db);source=activate(db,s,workbook(tmp_path/'multi.xlsx',[teacher]))
        group=db.scalar(select(ClassSection));issue=db.scalar(select(ValidationIssue).where(ValidationIssue.code=='MULTI_LECTURER_REVIEW'))
        assert group.assigned_lecturer_id is None and not group.locked_assignment
        assert db.query(Assignment).count()==0 and db.query(LecturerCourseCapability).count()==0
        assert issue.source_version_id==source.id and issue.details['rows'][0]['raw']==teacher
        assert issue.details['meeting_ids'] and issue.details['source_sheet']=='Schedule'
        sid=s.id
        with pytest.raises(SourceError): solve(db,2,False,s.id)
    response=TestClient(app).post(f'/api/v1/optimization/run?semester_id={sid}',json={'time_limit_seconds':2})
    assert response.status_code==422 and response.json()['detail']['blockers'][0]['code']=='MULTI_LECTURER_REVIEW'


def test_cross_row_identity_and_duplicate_lineage(tmp_path):
    with SessionLocal() as db:
        s=semester(db);source=activate(db,s,workbook(tmp_path/'cross.xlsx',['[A]Same Name','[B]Same Name']))
        group=db.scalar(select(ClassSection));meeting=db.scalar(select(ClassSession))
        assert group.assigned_lecturer_id is None
        assert group.source_version_id==meeting.source_version_id==source.id
        assert meeting.source_rows==[2,3] and len(group.sessions)==1
        assert db.scalar(select(ValidationIssue.code))=='MULTI_LECTURER_REVIEW'


def test_same_canonical_identity_normal_import_and_reconcile(tmp_path):
    with SessionLocal() as db:
        s=semester(db);teacher=Lecturer(code='A',canonical_name='Teacher Alpha',confirmed=True);db.add(teacher);db.flush()
        db.add(LecturerAlias(lecturer_id=teacher.id,alias_text='T Alpha',normalized_alias='t alpha',confirmed=True));db.commit()
        source=activate(db,s,workbook(tmp_path/'normal.xlsx',['[A]Teacher Alpha','T Alpha']))
        group=db.scalar(select(ClassSection))
        assert group.assigned_lecturer_id==teacher.id and not group.locked_assignment
        assert db.scalar(select(LecturerCourseCapability.confirmed)) is True
        assert not authority.source_blockers(db,s.id)
        assert group.sessions[0].source_rows==[2,3]
        pref=activate(db,s,preferences(tmp_path/'pref.xlsx'),'PREFERENCE')
        draft=db.scalar(select(NormalizedPreferenceDraft));draft.raw_text='Human edit';draft.status='REJECTED';draft.rejected_reason='Keep';db.commit()
        again=activate(db,s,tmp_path/'pref.xlsx','PREFERENCE')
        assert again.id==pref.id and db.query(NormalizedPreferenceDraft).count()==1
        assert draft.raw_text=='Human edit' and draft.rejected_reason=='Keep'


def test_human_normalization_and_deferred_split(tmp_path):
    with SessionLocal() as db:
        s=semester(db);activate(db,s,workbook(tmp_path/'multi.xlsx',['[A]Alpha;[B]Beta']))
        teacher=Lecturer(code='A',canonical_name='Alpha');db.add(teacher);db.commit()
        issue=db.scalar(select(ValidationIssue))
        resolve_issue(db,s.id,issue.id,action='DEFER_SPECIAL',note='Two actual teachers',actor='Reviewer')
        assert issue.code=='SPLIT_ASSIGNMENT_REQUIRES_SEGMENT_SUPPORT'
        with pytest.raises(SourceError):solve(db,2,False,s.id)
        resolve_issue(db,s.id,issue.id,action='NORMALIZE_SINGLE',lecturer_id=teacher.id,note='Corrected source interpretation',actor='Reviewer')
        assert db.scalar(select(ClassSection.assigned_lecturer_id))==teacher.id
        assert issue.resolution_status=='RESOLVED' and len(issue.details['audit'])==2
        assert not authority.source_blockers(db,s.id)
        assert db.query(LecturerCourseCapability).count()==0  # separate capability approval


def test_atomic_replacement_archives_runs_and_invalidates_references(tmp_path,monkeypatch):
    with SessionLocal() as db:
        s=semester(db);first=activate(db,s,workbook(tmp_path/'a.xlsx'))
        old=db.scalar(select(ClassSection));old_id=old.id;teacher_id=old.assigned_lecturer_id
        old.locked_assignment=True;old.assignment_source='MANUAL';db.commit()
        run=solve(db,2,False,s.id);run_id=run.id
        constraint=Constraint(semester_id=s.id,name='Old required',constraint_type='required_assignment',lecturer_id=teacher_id,target={'class_id':old.id},active=True,confirmed=True)
        db.add(constraint);db.flush()
        draft=NormalizedPreferenceDraft(semester_id=s.id,lecturer_id=teacher_id,constraint_type='REQUIRED_ASSIGNMENT',source_file='MANUAL',source_sheet='S',source_row=1,source_cell='A1',target={'class_id':old.id},status='CONFIRMED',applied_constraint_id=constraint.id)
        db.add(draft);db.commit()
        candidate=stage_source(db,workbook(tmp_path/'b.xlsx',rooms=['NEW']),s.id,'CURRENT_SCHEDULE');db.commit()
        preview=preview_activation(db,s.id,candidate.id,'CURRENT_SCHEDULE');assert preview['diff']['counts']['CHANGED']==1
        original_commit=db.commit
        monkeypatch.setattr(db,'commit',lambda: (_ for _ in ()).throw(RuntimeError('commit failure')))
        with pytest.raises(RuntimeError):activate_source(db,s.id,candidate.id,'CURRENT_SCHEDULE',confirmed=True,preview_token=preview['preview_token'])
        monkeypatch.setattr(db,'commit',original_commit)
        db.expire_all()
        assert db.get(Semester,s.id).active_schedule_source_id==first.id
        assert db.get(ClassSection,old_id).sessions[0].room=='R'
        assert db.get(OptimizationRun,run_id).summary.get('archived_assignments') is None
        activate_source(db,s.id,candidate.id,'CURRENT_SCHEDULE',confirmed=True,preview_token=preview['preview_token'])
        db.expire_all();current=db.scalar(select(ClassSection))
        assert current.id!=old_id and current.sessions[0].room=='NEW' and not current.locked_assignment
        assert not constraint.active and not constraint.confirmed and constraint.target['_stale_source_reference']
        assert draft.needs_review and draft.applied_constraint_id is None
        historical=db.get(OptimizationRun,run_id)
        assert historical.summary['archived_assignments'][0]['meetings'][0]['room']=='R'
        assert latest_current_run(db,s.id) is None and get_assignments(s.id,db)==[]
        assert get_dashboard(s.id,db)['optimization_status']=='not_run'
        assert any(i['code']=='SOURCE_CHANGED_REVIEW_REQUIRED' for i in authority.source_blockers(db,s.id))


def test_preference_switch_does_not_apply_previous_rules(tmp_path):
    with SessionLocal() as db:
        s=semester(db);activate(db,s,workbook(tmp_path/'s.xlsx'))
        previous=activate(db,s,preferences(tmp_path/'p.xlsx'),'PREFERENCE')
        draft=db.scalar(select(NormalizedPreferenceDraft));teacher=db.scalar(select(Lecturer))
        constraint=Constraint(semester_id=s.id,source_version_id=previous.id,name='Old',constraint_type='unavailable',lecturer_id=teacher.id,active=True,confirmed=True,target={})
        db.add(constraint);db.flush();draft.applied_constraint_id=constraint.id;draft.status='CONFIRMED';db.commit()
        new=activate(db,s,preferences(tmp_path/'p.xlsx','Xin nghỉ Thứ 3 tiết 4-6'),'PREFERENCE')
        assert new.id!=previous.id and s.active_preference_source_id==new.id
        assert not constraint.active and draft.applied_constraint_id is None
        assert db.query(NormalizedPreferenceDraft).count()==2
        assert all(d.applied_constraint_id is None for d in db.scalars(select(NormalizedPreferenceDraft)))
        assert not db.scalar(select(Constraint.id).where(Constraint.active.is_(True)))


def test_http_upload_is_staging_only_and_explicit_role(tmp_path):
    with SessionLocal() as db:s=semester(db);sid=s.id
    data=workbook(tmp_path/'schedule.xlsx').read_bytes();client=TestClient(app)
    response=client.post(f'/api/v1/sources/upload?semester_id={sid}&source_type=CURRENT_SCHEDULE',files={'file':('same.xlsx',data)})
    assert response.status_code==200 and not response.json()['active']
    source_id=response.json()['id']
    with SessionLocal() as db:assert db.query(ClassSection).count()==0
    response=client.post(f'/api/v1/sources/{source_id}/activate?semester_id={sid}',json={'source_type':'CURRENT_SCHEDULE'})
    assert response.status_code==422
    assert client.post(f'/api/v1/sources/upload?semester_id={sid}',files={'file':('same.xlsx',data)}).status_code==422


def test_real_files_read_only_dry_run(tmp_path):
    root=Path('/Users/mac/AI/Huce_timetable')
    files=[(root/'INPUT GỐC/Phan_Cong_Giang_Day_HK1_2026_2027.xlsx','CURRENT_SCHEDULE'),
           (root/'INPUT GỐC/Nguyện vọng TKB 2026-2027.xlsx','PREFERENCE'),
           (root/'phan-cong-giang-day-lich-hoc-to-bo-mon-03-09-2026.xls','HISTORICAL'),
           (root/'TKB-NCM Toan-Ky1-2026-2027.xlsx','REFERENCE_MATRIX')]
    if not all(p.exists() for p,_ in files):pytest.skip('Private integration files unavailable')
    before={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p,_ in files}
    with SessionLocal() as db:
        s=semester(db);sources=[stage_source(db,p,s.id,role) for p,role in files];db.commit()
        assert s.active_schedule_source_id is None and db.query(ClassSection).count()==0
        for source in sources[:2]:
            preview=preview_activation(db,s.id,source.id,source.source_type)
            assert preview['can_activate']
            activate_source(db,s.id,source.id,source.source_type,confirmed=True,preview_token=preview['preview_token'])
        assert all(g.source_version_id==sources[0].id for g in db.scalars(select(ClassSection)))
        assert all(m.source_rows and m.source_version_id==sources[0].id for m in db.scalars(select(ClassSession)))
        assert sum(len(m.source_rows) for m in db.scalars(select(ClassSession)))==sources[0].parse_summary['rows_accepted']
        assert all(d.source_version_id==sources[1].id for d in db.scalars(select(NormalizedPreferenceDraft)))
        with pytest.raises(SourceError): preview_activation(db,s.id,sources[2].id,'CURRENT_SCHEDULE')
        report={'sources':[authority.source_payload(src,s) for src in sources], 'groups':db.query(ClassSection).count(), 'meetings':db.query(ClassSession).count(), 'drafts':db.query(NormalizedPreferenceDraft).count(), 'blockers':authority.source_blockers(db,s.id)}
        print('REAL_SOURCE_DRY_RUN',json.dumps(report,ensure_ascii=False,default=str))
    assert before=={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p,_ in files}


def test_0008_upgrade_preserves_0007_rows_and_adds_foreign_keys(tmp_path):
    import sqlite3
    from test_rc_migration import legacy_at_0002, upgrade, snapshot, assert_preserved
    path=tmp_path/'upgrade-0007.db'
    legacy_at_0002(path);upgrade(path,'0007_preference_draft_rejection')
    before=snapshot(path)
    upgrade(path);upgrade(path)
    assert_preserved(path,before)
    with sqlite3.connect(path) as db:
        assert db.execute('SELECT COUNT(*) FROM source_versions').fetchone()==(0,)
        assert db.execute('SELECT active_schedule_source_id, active_preference_source_id FROM semesters').fetchone()==(None,None)
        for table,column in [('classes','source_version_id'),('sessions','source_version_id'),('historical_evidence','source_version_id'),('semesters','active_schedule_source_id'),('optimization_runs','schedule_source_id')]:
            assert column in {r[3] for r in db.execute(f'PRAGMA foreign_key_list({table})')}
