"""Regression cases from the full audit; original expected values are preserved."""
import os
os.environ['DATABASE_URL']='sqlite://'
from datetime import date
from pathlib import Path
import pytest
from sqlalchemy import create_engine,select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from fastapi.testclient import TestClient
from openpyxl import Workbook
from app.db.session import Base,get_db
from app.main import app
from app.models.entities import Semester,Course,Lecturer,LecturerCourseCapability,ClassSection,ClassSession,Assignment,Constraint,Seminar
from app.optimization.solver import solve
from app.services.importer import _summary
from app.services.manual_assignment import check_assignment_change,apply_manual_assignment
from app.services.unassigned_diagnostics import diagnose_unassigned_classes,get_candidate_analysis
from app.services.source_authority import stage_source,preview_activation,activate_source
from app.parsers.schedule import parse_schedule

@pytest.fixture
def db():
    engine=create_engine('sqlite://',poolclass=StaticPool,connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine)
    with Session(engine,expire_on_commit=False) as session: yield session
    engine.dispose()

def setup(db,n=1,teachers=1,credits=None,merged=False,locked=False):
    s=Semester(name='QA anonymous',department_name='QA',head_name='QA',start_date=date(2026,8,3),end_date=date(2027,1,31))
    c=Course(code='QA101',name='QA course'); ls=[Lecturer(code=f'{i:05}',canonical_name=f'Lecturer {i}',confirmed=True) for i in range(teachers)]
    db.add_all([s,c,*ls]);db.flush()
    db.add_all([LecturerCourseCapability(lecturer_id=l.id,course_id=c.id,allowed=True,confirmed=True,source='MANUAL') for l in ls])
    gs=[]
    for i in range(n):
        g=ClassSection(semester_id=s.id,course_id=c.id,class_code=f'QA{i}',credits=credits[i] if credits else 3,source_file='qa',source_sheet='qa',source_row=i+1,raw_values={},locked_assignment=locked,assigned_lecturer_id=ls[i%teachers].id if locked else None,assignment_source='MANUAL' if locked else None,merged_group_id='QA-MERGE' if merged else None,merge_status='candidate' if merged else 'single')
        db.add(g);db.flush();db.add(ClassSession(class_id=g.id,weekday=2,start_period=1,end_period=3,room='QA',active_weeks=[1],raw_weeks='1',start_date=date(2026,8,3),end_date=date(2026,8,3),source_row=i+1))
        gs.append(g)
    db.commit(); return s,c,ls,gs

def workbook(path,teacher='[00187] Lecturer A',room='R'):
    w=Workbook();s=w.active
    s.append(['STT','Mã học phần','Tên môn học','Mã lớp học','','Thứ','Tiết','Phòng','Số TC','','Bắt đầu','Kết thúc','Tuần học','Giảng viên'])
    s.append([1,'QA101','QA course','QA0','',2,'1-3',room,3,'','03/08/2026','03/08/2026','1',teacher]);w.save(path);return path

def test_import_summary_counts_assignments_separately_from_locks(tmp_path):
    parsed=parse_schedule(workbook(tmp_path/'assigned.xlsx'))
    assert len(parsed.classes)==1 and parsed.classes[0].lecturer_code=='00187'
    assert not parsed.classes[0].locked_assignment
    assert _summary(parsed)['unassigned_classes']==0

def test_workload_objective_uses_each_class_credits(db):
    s,c,ls,gs=setup(db,2,2,credits=[1,3],locked=True)
    run=solve(db,2,False,s.id)
    assert run.status=='optimal'
    # Two fixed teachers with 1 vs 3 credits: scaled disparity (30-10)*1000.
    assert run.summary['experiment']['objective_components']['workload_disparity']==20000

def test_seven_candidate_diagnostic_respects_unconfirmed_merge(db):
    s,c,ls,gs=setup(db,8,7,merged=True)
    run=solve(db,2,False,s.id)
    assert run.status=='optimal'
    assert len(list(db.scalars(select(Assignment).where(Assignment.run_id==run.id))))==7
    diagnosis=diagnose_unassigned_classes(db,s.id,run.id)
    assert len(diagnosis)==1
    assert diagnosis[0]['eligible_candidates_count']==0, diagnosis[0]

def test_manual_check_uses_current_solver_assignments(db):
    s,c,ls,gs=setup(db,2,1)
    run=solve(db,2,False,s.id)
    assigned={a.class_id for a in db.scalars(select(Assignment).where(Assignment.run_id==run.id))}
    unassigned=next(g for g in gs if g.id not in assigned)
    result=check_assignment_change(db,s.id,unassigned.id,ls[0].id)
    assert result['valid'] is False and 'TIMETABLE_CONFLICT' in result['blocking_reasons'],result

def test_manual_assignment_visible_without_lock_after_solver(db):
    s,c,ls,gs=setup(db,1,2)
    run=solve(db,2,False,s.id)
    old=db.scalar(select(Assignment).where(Assignment.run_id==run.id)).lecturer_id
    new=next(l.id for l in ls if l.id!=old)
    assert apply_manual_assignment(db,s.id,gs[0].id,new,lock=False)['valid']
    from app.api.routes import get_classes
    assert get_classes(semester_id=s.id,run_id=None,db=db)[0]['lecturer_id']==new

def test_hard_avoid_period_enforced_by_manual_checker(db):
    s,c,ls,gs=setup(db)
    db.add(Constraint(semester_id=s.id,name='QA hard avoid',constraint_type='AVOID_PERIOD',hardness='hard',weight=1,lecturer_id=ls[0].id,target={'weekday':2,'periods':[1,2,3]},confirmed=True,active=True));db.commit()
    assert check_assignment_change(db,s.id,gs[0].id,ls[0].id)['valid'] is False

def test_manual_checker_respects_shared_hard_seminar(db):
    s,c,ls,gs=setup(db)
    db.add(Seminar(semester_id=s.id,name='QA seminar',chair_name='QA',members=[ls[0].id],alternatives=[{'weekday':2,'periods':[1,2,3]}],hardness='hard',weight=1));db.commit()
    assert check_assignment_change(db,s.id,gs[0].id,ls[0].id)['valid'] is False

def test_known_soft_avoid_penalty_is_60(db):
    s,c,ls,gs=setup(db)
    db.add(Constraint(semester_id=s.id,name='QA soft',constraint_type='AVOID_PERIOD',hardness='soft',weight=1,lecturer_id=ls[0].id,target={'weekday':2,'periods':[1,2,3]},confirmed=True,active=True));db.commit()
    run=solve(db,2,False,s.id)
    assert len(list(db.scalars(select(Assignment).where(Assignment.run_id==run.id))))==1
    assert run.summary['experiment']['objective_components']['soft_preferences']==60

def test_approved_draft_forced_soft_violation(db):
    s,c,ls,gs=setup(db)
    app.dependency_overrides[get_db]=lambda: db
    try:
        client=TestClient(app)
        r=client.post(f'/api/v1/preference-drafts?semester_id={s.id}',json={'lecturer_id':ls[0].id,'context_type':'TEACHING','constraint_type':'AVOID_PERIOD','day_scope':'T2','periods':[1,2,3],'hardness':'soft','weight':1,'note':'QA force soft violation','status':'CONFIRMED'})
        assert r.status_code==200,r.text
        before=solve(db,2,False,s.id)
        assert before.summary['experiment']['objective_components']['soft_preferences']==0
        r=client.post(f'/api/v1/preference-drafts/apply?semester_id={s.id}',json={'draft_ids':r.json()['ids']})
        assert r.status_code==200,r.text
        after=solve(db,2,False,s.id)
        assert after.summary['experiment']['objective_components']['soft_preferences']==60
    finally: app.dependency_overrides.clear()

def test_no_solution_metrics_do_not_read_variables(db):
    s,c,ls,gs=setup(db)
    db.add(Constraint(semester_id=s.id,name='QA impossible',constraint_type='MIN_CLASSES',hardness='hard',lecturer_id=ls[0].id,target={'min':2},confirmed=True,active=True));db.commit()
    run=solve(db,2,False,s.id)
    assert run.status=='infeasible'
    assert run.score is None and run.summary['experiment']['objective_components'] is None

def test_unknown_zero_search_budget_safe(db):
    s,c,ls,gs=setup(db,8,7)
    run=solve(db,0,False,s.id)
    assert run.status=='unknown'
    assert run.summary['experiment']['objective_components'] is None

def test_empty_semester_safe(db):
    s,c,ls,gs=setup(db,0,1)
    run=solve(db,1,False,s.id)
    assert run.status in ['optimal','blocked']

def test_same_source_reactivation_preserves_explicit_lock(db,tmp_path):
    s,c,ls,gs=setup(db,0,1)
    path=workbook(tmp_path/'original.xlsx','[00000]Lecturer 0')
    src=stage_source(db,path,s.id,'CURRENT_SCHEDULE');db.commit()
    p=preview_activation(db,s.id,src.id,'CURRENT_SCHEDULE')
    activate_source(db,s.id,src.id,'CURRENT_SCHEDULE',confirmed=True,preview_token=p['preview_token'])
    g=db.scalar(select(ClassSection).where(ClassSection.semester_id==s.id));g.locked_assignment=True;db.commit()
    p=preview_activation(db,s.id,src.id,'CURRENT_SCHEDULE')
    activate_source(db,s.id,src.id,'CURRENT_SCHEDULE',confirmed=True,preview_token=p['preview_token'])
    db.refresh(g);assert g.locked_assignment

def test_source_change_preview_enumerates_lock_loss(db,tmp_path):
    s,c,ls,gs=setup(db,0,1)
    path=workbook(tmp_path/'original.xlsx','[00000]Lecturer 0')
    src=stage_source(db,path,s.id,'CURRENT_SCHEDULE');db.commit();p=preview_activation(db,s.id,src.id,'CURRENT_SCHEDULE')
    activate_source(db,s.id,src.id,'CURRENT_SCHEDULE',confirmed=True,preview_token=p['preview_token'])
    g=db.scalar(select(ClassSection).where(ClassSection.semester_id==s.id));g.locked_assignment=True;db.commit()
    changed=workbook(tmp_path/'changed.xlsx','[00000]Lecturer 0',room='NEW')
    src2=stage_source(db,changed,s.id,'CURRENT_SCHEDULE');db.commit();p=preview_activation(db,s.id,src2.id,'CURRENT_SCHEDULE')
    # A generic warning is present, but a reviewable impact list must identify the affected lock.
    assert any('lock' in str(k).lower() for k in p if k!='consequences'), p.keys()

from datetime import date
from pathlib import Path
from sqlalchemy import select
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format
from app.main import app
from app.db.session import get_db
from app.models.entities import Assignment,Constraint
from app.optimization.solver import solve

def test_nonoverlapping_actual_dates_allow_same_lecturer(db):
    s,c,ls,gs=setup(db,2,1)
    a,b=gs[0].sessions[0],gs[1].sessions[0]
    a.active_weeks=[1,2];a.start_date=date(2026,8,3);a.end_date=date(2026,8,9)
    b.active_weeks=[1,2];b.start_date=date(2026,8,4);b.end_date=date(2026,8,10)
    db.commit()
    # Monday occurrences are independently Aug 3 and Aug 10; no actual collision.
    run=solve(db,2,False,s.id)
    assert len(list(db.scalars(select(Assignment).where(Assignment.run_id==run.id))))==2

def test_alternating_weeks_do_not_exceed_one_session_per_day(db):
    s,c,ls,gs=setup(db,2,1)
    gs[1].sessions[0].active_weeks=[2]
    gs[1].sessions[0].raw_weeks=' 2'
    gs[1].sessions[0].start_date=gs[1].sessions[0].end_date=date(2026,8,10)
    db.add(Constraint(semester_id=s.id,name='QA once daily',constraint_type='MAX_SESSIONS_PER_DAY',hardness='hard',weight=1,lecturer_id=ls[0].id,target={'max':1},confirmed=True,active=True));db.commit()
    run=solve(db,2,False,s.id)
    assert len(list(db.scalars(select(Assignment).where(Assignment.run_id==run.id))))==2

def test_partial_week_merge_rejected_by_api(db):
    s,c,ls,gs=setup(db,2,1)
    gs[0].sessions[0].active_weeks=[1,2]
    gs[1].sessions[0].active_weeks=[2,3]
    for g in gs:g.sessions[0].end_date=date(2026,8,31)
    db.commit();app.dependency_overrides[get_db]=lambda:db
    try:
        r=TestClient(app).post(f'/api/v1/classes/merge?semester_id={s.id}',json={'class_ids':[g.id for g in gs]})
        assert r.status_code==400,r.text
    finally:app.dependency_overrides.clear()

def test_xls_template_export_preserves_date_display_format(tmp_path):
    from app.services.template_detector import detect_output_template
    from app.exporters.excel import _export_template
    from app.models.entities import OutputTemplateProfile
    source = Path(__file__).resolve().parents[3] / 'storage/uploads/schedule/phan-cong-giang-day-lich-hoc-to-bo-mon-03-09-2026.xls'
    assert source.exists(), "Read-only real XLS fixture is required"
    detected = detect_output_template(source)
    profile = OutputTemplateProfile(source_file=str(source), source_sheet='Sheet1', header_row=detected['header_row'], mappings=detected['mappings'])
    output = _export_template(profile, [], [], tmp_path, source=source)
    w = load_workbook(output)
    cell = w['Sheet1']['K10']
    assert is_date_format(cell.number_format), {'value':cell.value,'number_format':cell.number_format}


def test_assignment_api_missing_lecturer_returns_validation_error(db):
    s,c,ls,gs=setup(db)
    app.dependency_overrides[get_db]=lambda:db
    try:
        r=TestClient(app,raise_server_exceptions=False).post(f'/api/v1/classes/{gs[0].id}/assignment/check?semester_id={s.id}',json={})
        assert r.status_code==422,r.text
    finally:app.dependency_overrides.clear()
