from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import select
from app.db.session import Base, SessionLocal, engine
from app.exporters.excel import export_latest
from app.models.entities import ClassSession, OutputTemplateProfile
from app.optimization.solver import solve
from test_phase3_solver import setup, capability

@pytest.fixture(autouse=True)
def database():
    Base.metadata.drop_all(engine); Base.metadata.create_all(engine); yield; Base.metadata.drop_all(engine)

def test_template_export_updates_only_lecturer_source_rows(tmp_path):
    source=tmp_path/"source.xlsx"; book=Workbook(); sheet=book.active; sheet.title="S"; sheet.append(["Mã học phần","Mã lớp học","Giảng viên","Phòng"]); sheet.append(["C","L0","old","P"]); book.save(source)
    with SessionLocal() as db:
        s,c,a,b,items=setup(db); capability(db,a,c); run=solve(db,2,False,s.id)
        db.add(OutputTemplateProfile(semester_id=s.id,source_file=str(source),source_sheet="S",header_row=1,mappings={"course_code":{"column_index":1},"class_code":{"column_index":2},"lecturer":{"column_index":3}},missing_fields=[])); db.commit()
        result=export_latest(db,tmp_path/"out",s.id)
        output=load_workbook(result)["S"]
        assert output.max_row==2 and output.cell(2,1).value=="C" and output.cell(2,2).value=="L0" and output.cell(2,4).value=="P"
        assert output.cell(2,3).value=="A"


def test_prior_output_template_disambiguates_repeated_class_by_schedule_columns(tmp_path):
    source = tmp_path / "prior-output.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet.append(["Mã học phần", "Mã lớp học", "Giảng viên", "Thứ", "Tiết", "Phòng"])
    sheet.append(["C", "L0", "old", 2, "4-6", "P1"])
    sheet.append(["C", "L0", "old", 3, "7-9", "P2"])
    book.save(source)
    with SessionLocal() as db:
        semester, course, lecturer, _, items = setup(db)
        capability(db, lecturer, course)
        item = items[0]
        item.source_file = "current-import.xlsx"
        item.source_sheet = "S"
        item.sessions[0].room = "P1"
        db.add(ClassSession(
            class_id=item.id, weekday=3, start_period=7, end_period=9,
            room="P2", raw_weeks="34", active_weeks=[3, 4], source_row=2,
        ))
        db.commit()
        solve(db, 2, False, semester.id)
        db.add(OutputTemplateProfile(
            semester_id=semester.id,
            source_file=str(source),
            source_sheet="S",
            header_row=1,
            mappings={
                "course_code": {"column_index": 1},
                "class_code": {"column_index": 2},
                "lecturer": {"column_index": 3},
                "weekday": {"column_index": 4},
                "periods": {"column_index": 5},
                "room": {"column_index": 6},
            },
            missing_fields=[],
        ))
        db.commit()
        result = export_latest(db, tmp_path / "out", semester.id)
    output = load_workbook(result)["S"]
    assert output.cell(2, 3).value == "A"
    assert output.cell(3, 3).value == "A"


def test_prior_output_template_allows_repeated_meetings_of_one_teaching_group(tmp_path):
    source = tmp_path / "prior-output.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet.append(["Mã học phần", "Mã lớp học", "Giảng viên", "Thứ", "Tiết"])
    sheet.append(["C", "L0", "old", 2, "4-6"])
    sheet.append(["C", "L0", "old", 2, "4-6"])
    book.save(source)
    with SessionLocal() as db:
        semester, course, lecturer, _, items = setup(db)
        capability(db, lecturer, course)
        item = items[0]
        item.source_file = "current-import.xlsx"
        item.source_sheet = "S"
        db.add(ClassSession(
            class_id=item.id, weekday=2, start_period=4, end_period=6,
            room="P2", raw_weeks="34", active_weeks=[3, 4], source_row=2,
        ))
        db.commit()
        solve(db, 2, False, semester.id)
        db.add(OutputTemplateProfile(
            semester_id=semester.id,
            source_file=str(source),
            source_sheet="S",
            header_row=1,
            mappings={
                "course_code": {"column_index": 1},
                "class_code": {"column_index": 2},
                "lecturer": {"column_index": 3},
                "weekday": {"column_index": 4},
                "periods": {"column_index": 5},
            },
            missing_fields=[],
        ))
        db.commit()
        result = export_latest(db, tmp_path / "out", semester.id)
    output = load_workbook(result)["S"]
    assert output.cell(2, 3).value == "A"
    assert output.cell(3, 3).value == "A"


def test_prior_output_template_preserves_unmatched_historical_rows(tmp_path):
    source = tmp_path / "prior-output.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet.append(["Mã học phần", "Mã lớp học", "Giảng viên"])
    sheet.append(["C", "L0", "old-current"])
    sheet.append(["OLD", "HISTORICAL", "keep-history"])
    book.save(source)
    with SessionLocal() as db:
        semester, course, lecturer, _, _ = setup(db)
        capability(db, lecturer, course)
        solve(db, 2, False, semester.id)
        db.add(OutputTemplateProfile(
            semester_id=semester.id,
            source_file=str(source),
            source_sheet="S",
            header_row=1,
            mappings={
                "course_code": {"column_index": 1},
                "class_code": {"column_index": 2},
                "lecturer": {"column_index": 3},
            },
            missing_fields=[],
        ))
        db.commit()
        result = export_latest(db, tmp_path / "out", semester.id)
    output = load_workbook(result)["S"]
    assert output.cell(2, 3).value == "A"
    assert output.cell(3, 3).value == "keep-history"
