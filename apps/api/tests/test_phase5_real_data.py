from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.routes import get_problems
from app.db.session import Base, SessionLocal, engine
from app.exporters.excel import export_latest
from app.models.entities import (
    ClassSection,
    Constraint,
    Lecturer,
    LecturerCourseCapability,
    OptimizationRun,
    OutputTemplateProfile,
    Semester,
)
from app.optimization.solver import solve
from app.parsers.preferences import parse_preference_workbook
from app.parsers.schedule import parse_schedule
from app.services.importer import _import_files_impl as import_files
from app.services.manual_assignment import apply_manual_assignment, check_assignment_change
from app.services.readiness import capability_readiness, validate_schedule
from app.services.template_detector import detect_output_template


@pytest.fixture(autouse=True)
def database():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    yield
    Base.metadata.drop_all(engine)


def test_real_huce_schedule_and_preferences_normalize_without_losing_rows():
    root = Path(__file__).resolve().parents[4]
    schedule = root / "Phan_Cong_Giang_Day_HK1_2026_2027.xlsx"
    preferences = root / "Nguyện vọng TKB 2026-2027.xlsx"
    assert schedule.exists() and preferences.exists()
    result = parse_schedule(schedule)
    wishes = parse_preference_workbook(preferences)
    assert result.rows_accepted == 286
    assert result.rows_rejected == 0
    assert len(result.classes) == 158
    assert sum(len(item.sessions) for item in result.classes) == 286
    assert all(len({meeting.signature() for meeting in item.sessions}) == len(item.sessions) for item in result.classes)
    assert wishes.format == "LEGACY"
    assert len(wishes.drafts) == 39
    assert len(wishes.seminars) == 0
    assert wishes.raw_clauses == 38
    assert wishes.dropped_clauses == 0
    assert all(item.raw_text for item in wishes.drafts)
    lieu = [item for item in wishes.drafts if item.lecturer_alias == "Liễu"]
    assert len(lieu) >= 3 and {item.context_type for item in lieu} == {"TEACHING", "SEMINAR"}
    x_linh_mixed = [
        item for item in wishes.drafts
        if item.lecturer_alias == "X Linh" and item.source_cell.startswith('F6#')
    ]
    assert len(x_linh_mixed) == 2
    assert {item.context_type for item in x_linh_mixed} == {"TEACHING", "SEMINAR"}
    # Newlines delimit independent source clauses, preserving both intentions.
    assert len({item.raw_text for item in x_linh_mixed}) == 2
    assert any('12h30' in item.raw_text and item.context_type=='SEMINAR' for item in x_linh_mixed)
    assert any('tiết 10' in item.raw_text and item.context_type=='TEACHING' for item in x_linh_mixed)


def test_real_data_import_solve_problem_export_round_trip_and_persistence(tmp_path):
    """The real workbook is never mutated; its copied draft is round-trippable."""
    root = Path(__file__).resolve().parents[4]
    schedule = root / "Phan_Cong_Giang_Day_HK1_2026_2027.xlsx"
    preferences = root / "Nguyện vọng TKB 2026-2027.xlsx"

    with SessionLocal() as db:
        semester = Semester(
            name="HK1 2026-2027 real acceptance",
            department_name="Toán",
            start_date=date(2026, 8, 1),
            end_date=date(2027, 1, 31),
            head_name="Trưởng bộ môn",
        )
        db.add(semester)
        db.commit()
        semester_id = semester.id

        imported = import_files(
            db,
            [schedule, preferences],
            semester_id=semester_id,
            schedule_paths=[schedule],
            preference_paths=[preferences],
        )
        assert imported["summary"]["classes"] == 158
        assert imported["summary"]["sessions"] == 286
        readiness = capability_readiness(db, semester_id)
        assert readiness["total_teaching_groups"] == 158
        assert readiness["code"] == "CAPABILITY_DATA_INCOMPLETE"

        detected = detect_output_template(schedule)
        assert not detected["missing_fields"]
        db.add(OutputTemplateProfile(
            semester_id=semester_id,
            source_file=str(schedule.resolve()),
            source_sheet=detected["source_sheet"],
            header_row=detected["header_row"],
            mappings=detected["mappings"],
            missing_fields=detected["missing_fields"],
            preview=detected["preview"],
        ))
        db.commit()

        run = solve(db, 5, False, semester_id)
        # Imported teaching assignments are provenance only.  They remain
        # editable until the department head explicitly locks one.
        assert not db.scalars(select(ClassSection).where(
            ClassSection.semester_id == semester_id,
            ClassSection.assignment_source == "IMPORT",
            ClassSection.locked_assignment.is_(True),
        )).first()
        assert run.status in {"optimal", "feasible"}

        # Find a real eligible manual case instead of naming a lecturer in the
        # fixture.  It is then locked and re-solved; the explicit manual lock
        # survives while the rest of the imported schedule remains editable.
        manual_case = next(
            (
                (group, lecturer)
                for group in db.scalars(
                    select(ClassSection).where(
                        ClassSection.semester_id == semester_id,
                        ClassSection.locked_assignment.is_(False),
                    )
                )
                for lecturer in db.scalars(select(Lecturer))
                if check_assignment_change(db, semester_id, group.id, lecturer.id)["valid"]
            ),
            None,
        )
        if manual_case is None:
            pytest.skip("REAL_DATA_NO_ELIGIBLE_MANUAL_CASE")
        manual_group, manual_lecturer = manual_case
        assert apply_manual_assignment(db, semester_id, manual_group.id, manual_lecturer.id, lock=True)["valid"]
        rerun = solve(db, 5, False, semester_id)
        assert rerun.id != run.id
        db.refresh(manual_group)
        assert manual_group.assigned_lecturer_id == manual_lecturer.id
        assert manual_group.locked_assignment is True
        validation = validate_schedule(db, semester_id, rerun.id)
        assert validation["valid"]

        problems = get_problems(semester_id=semester_id, db=db)
        assert not any(item["code"] == "LOCKED_ASSIGNMENT_CONFLICT" for item in problems)
        assert len({(item["code"], item["entity_type"], item["entity_id"]) for item in problems}) == len(problems)

        source_book = load_workbook(schedule, data_only=False)
        source_sheet = source_book[detected["source_sheet"]]
        output_path = export_latest(db, tmp_path, semester_id, mode="draft")
        output_book = load_workbook(output_path, data_only=False)
        output_sheet = output_book[detected["source_sheet"]]
        assert output_sheet.max_row == source_sheet.max_row
        lecturer_col = detected["mappings"]["lecturer"]["column_index"]
        for row in range(1, source_sheet.max_row + 1):
            for column in range(1, source_sheet.max_column + 1):
                if column != lecturer_col:
                    assert output_sheet.cell(row, column).value == source_sheet.cell(row, column).value
        with pytest.raises(ValueError, match="FINAL_EXPORT_NOT_READY"):
            export_latest(db, tmp_path, semester_id, mode="final")

    # Fresh-session reads prove that import, profile, run and snapshots are
    # durable rather than accidental state held by the import session.
    with SessionLocal() as fresh:
        assert fresh.get(Semester, semester_id) is not None
        assert fresh.scalar(select(ClassSection).where(ClassSection.semester_id == semester_id)) is not None
        # Imported preferences remain review drafts and do not silently become
        # solver-visible constraints.
        assert fresh.scalar(select(Constraint).where(Constraint.semester_id == semester_id)) is None
        assert fresh.scalar(select(LecturerCourseCapability)) is not None
        assert fresh.scalar(select(OutputTemplateProfile).where(OutputTemplateProfile.semester_id == semester_id)) is not None
        assert fresh.scalars(select(OptimizationRun).where(OptimizationRun.semester_id == semester_id)).all()
        persisted_manual = fresh.get(ClassSection, manual_group.id)
        assert persisted_manual.assigned_lecturer_id == manual_lecturer.id
        assert persisted_manual.locked_assignment is True
