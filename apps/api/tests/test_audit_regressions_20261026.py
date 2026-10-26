"""Regressions for audit BUG-01..04 (output/qa-audit-20261026/BUGS.md). Expectations come from the audit probes."""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import Base, SessionLocal, engine
from app.main import app
from app.models.entities import Assignment, ClassSection, Constraint
from app.optimization.solver import solve
from app.schemas.api import OptimizationRequest
from test_phase3_solver import capability, setup


@pytest.fixture(autouse=True)
def database():
    Base.metadata.drop_all(engine); Base.metadata.create_all(engine); yield; Base.metadata.drop_all(engine)


def rule(db, s, lecturer, kind, target):
    db.add(Constraint(semester_id=s.id, name=kind, constraint_type=kind, hardness="hard", weight=1, lecturer_id=lecturer.id,
                      target=target, confirmed=True, active=True)); db.commit()


def plan(db, run):
    return {a.class_id: a.lecturer_id for a in db.scalars(select(Assignment).where(Assignment.run_id == run.id))}


def test_bug01_forbidden_course_holds_when_lecturer_is_not_candidate_for_one_class():
    with SessionLocal() as db:
        s, c, a, b, items = setup(db, 2); capability(db, a, c); capability(db, b, c)
        items[1].sessions[0].weekday = 3
        db.commit()
        rule(db, s, a, "UNAVAILABLE", {"weekday": 2, "periods": [4, 5, 6]})   # A is not a candidate for items[0]
        rule(db, s, b, "UNAVAILABLE", {"weekday": 3, "periods": [4, 5, 6]})   # B is not a candidate for items[1]
        rule(db, s, a, "FORBIDDEN_ASSIGNMENT", {"course_id": c.id})
        run = solve(db, 5, False, s.id)
        assert a.id not in plan(db, run).values()
        assert items[1].id in {u["class_id"] for u in run.summary["unassigned"]}
        assert run.summary["invalid_constraints"] == []


def test_bug02_pending_member_does_not_inherit_merge_exemption():
    with SessionLocal() as db:
        s, c, a, b, items = setup(db, 2); capability(db, a, c)
        items[0].merged_group_id = items[1].merged_group_id = "G"
        items[0].merged_confirmed = True; items[0].merge_status = "confirmed"
        items[1].merged_confirmed = False; items[1].merge_status = "candidate"
        db.commit()
        run = solve(db, 5, False, s.id)
        assert len(plan(db, run)) == 1


def test_bug02_merge_endpoint_rejects_group_with_unselected_members():
    with SessionLocal() as db:
        s, c, a, b, items = setup(db, 3)
        for item in items:
            item.merged_group_id = "G"
        db.commit(); sid, ids = s.id, [items[0].id, items[1].id]
    client = TestClient(app)
    assert client.post(f"/api/v1/classes/merge?semester_id={sid}", json={"class_ids": ids, "merged_group_id": "G"}).status_code == 400
    assert client.post(f"/api/v1/classes/merge?semester_id={sid}", json={"class_ids": ids}).status_code == 200


def test_bug03_draft_export_never_uses_source_file_lecturer_for_unassigned_class():
    with SessionLocal() as db:
        s, c, a, b, items = setup(db, 2); capability(db, a, c)
        for item in items:  # stale lecturer from the source file on every class, whichever one the solver leaves out
            item.assigned_lecturer_id = b.id; item.assignment_source = "IMPORTED"
        db.commit(); sid = s.id
        run = solve(db, 5, False, sid)
        assigned_class = next(iter(plan(db, run))); other = next(i.class_code for i in items if i.id != assigned_class)
    response = TestClient(app).get(f"/api/v1/exports/latest?semester_id={sid}&mode=draft")
    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.content))["Phân công"].iter_rows(min_row=3, values_only=True))
    lecturer_by_class = {r[3]: r[12] for r in rows if r[3]}
    assert lecturer_by_class[other] == "Chưa phân công"


def test_bug04_run_request_does_not_confirm_merge_suggestions_by_default():
    assert OptimizationRequest().confirm_merged_suggestions is False
    with SessionLocal() as db:
        s, c, a, b, items = setup(db, 2); capability(db, a, c)
        for item in items:
            item.merged_group_id = "G"
        db.commit(); sid = s.id
    assert TestClient(app).post(f"/api/v1/optimization/run?semester_id={sid}", json={}).status_code in (200, 422)
    with SessionLocal() as db:
        assert all(i.merge_status != "confirmed" for i in db.scalars(select(ClassSection).where(ClassSection.semester_id == sid)))
