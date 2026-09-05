from pathlib import Path

from openpyxl import Workbook

from app.parsers.preferences import describe_preference, _infer
from app.parsers.schedule import decode_weeks, parse_lecturer, parse_schedule


def _schedule_book(path, rows):
    book = Workbook(); sheet = book.active
    sheet.append(["STT", "Mã học phần", "Tên môn học", "Mã lớp học", "", "Thứ", "Tiết", "Phòng", "Số TC", "", "Bắt đầu", "Kết thúc", "Tuần học", "Giảng viên"])
    for row in rows: sheet.append(row)
    book.save(path)
from app.services.template_detector import detect_output_template


def test_week_decoder_preserves_positions():
    raw, active = decode_weeks("1234 678  12345678")
    assert raw == "1234 678  12345678"
    assert active == [1, 2, 3, 4, 6, 7, 8, 11, 12, 13, 14, 15, 16, 17, 18]


def test_parse_lecturer_uses_primary_teacher_for_co_teaching_cell():
    code, name = parse_lecturer("[00173]Nguyễn Bằng Giang,\n[00181]Nguyễn Xuân Linh")
    assert code == "00173"
    assert name == "Nguyễn Bằng Giang"


def test_parse_lecturer_treats_unassigned_placeholders_as_empty():
    for value in (
        "Chưa phân công",
        "CHUA PHAN CONG",
        "Chưa xếp giảng viên",
        "Chưa có GV",
        "Đang cập nhật",
        "N/A",
        "-",
    ):
        assert parse_lecturer(value) == (None, None)


def test_schedule_parser_keeps_first_data_row_after_header(tmp_path):
    path = tmp_path / "immediate.xlsx"
    _schedule_book(path, [[1, "M1", "Math", "L1", "", 2, "1-3", "P", 3, "", "01/01/2026", "01/06/2026", "123", "[T1]Teacher One"]])
    result = parse_schedule(path)
    assert len(result.classes) == 1 and result.classes[0].class_code == "L1"


def test_schedule_parser_supports_blank_subheader_before_data(tmp_path):
    path = tmp_path / "subheader.xlsx"
    _schedule_book(path, [[], [1, "M1", "Math", "L1", "", 2, "1-3", "P", 3, "", "01/01/2026", "01/06/2026", "123", "Teacher One"]])
    result = parse_schedule(path)
    assert len(result.classes) == 1


def test_schedule_parser_deduplicates_identical_meetings(tmp_path):
    row = [1, "M1", "Math", "L1", "", 2, "1-3", "P", 3, "", "01/01/2026", "01/06/2026", "123", "Teacher One"]
    path = tmp_path / "duplicate.xlsx"; _schedule_book(path, [row, row])
    assert len(parse_schedule(path).classes[0].sessions) == 1


def test_co_teaching_is_preserved_for_review(tmp_path):
    path = tmp_path / "cotaught.xlsx"
    _schedule_book(path, [[1, "M1", "Math", "L1", "", 2, "1-3", "P", 3, "", "01/01/2026", "01/06/2026", "123", "[T1]Teacher One, [T2]Teacher Two"]])
    result = parse_schedule(path)
    assert any(issue.code == "CO_TEACHING_REQUIRES_REVIEW" for issue in result.issues)
    assert result.classes[0].locked_assignment is False


def test_preference_normalization_extracts_unavailability_rule():
    constraint_type, target, confidence = _infer("Cô xin không dạy thứ 4 tiết 4 đến 6", None)

    assert constraint_type == "unavailable"
    assert target["weekday"] == 4
    assert target["periods"] == [4, 5, 6]
    assert confidence >= 0.8
    assert "Thứ 4" in describe_preference(constraint_type, target)


def test_preference_normalization_does_not_mistake_weekday_for_period():
    constraint_type, target, _ = _infer("Xin nghỉ sáng thứ 6", None)

    assert constraint_type == "unavailable"
    assert target["weekday"] == 6
    assert target["periods"] == []


def test_parse_real_schedule_when_present():
    path = Path(__file__).resolve().parents[3] / "data/local/source/schedule-2026-07-28.xls"
    if not path.exists():
        return
    result = parse_schedule(path)
    assert result.rows_accepted > 300
    assert len(result.classes) > 170
    assert sum(item.locked_assignment for item in result.classes) > 30
    assert result.merged_groups


def test_detect_output_template_columns(tmp_path):
    path = tmp_path / "mau-ky-truoc.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Mã lớp", "Tên môn học", "Giảng viên", "Thứ", "Tiết học"])
    sheet.append(["71CSQT", "Giải tích 1", "Nguyễn Bằng Giang", 3, "1-3"])
    book.save(path)

    result = detect_output_template(path)

    assert result["header_row"] == 1
    assert result["ready"] is True
    assert result["mappings"]["lecturer"]["column_letter"] == "C"
    assert result["mappings"]["class_code"]["column_letter"] == "A"


def test_detect_output_template_matrix(tmp_path):
    path = tmp_path / "tkb-bo-mon.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "TKB_Bo_Mon"
    sheet.append(["BẢNG PHÂN CÔNG GIẢNG DẠY"])
    sheet.append([])
    sheet.append(["Giảng viên / Thứ", "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5"])
    sheet.append(["Nguyễn Bằng Giang", "Giải tích 1 - 71CSQT\n(Tiết 1-3)", None, None, None])
    book.save(path)

    result = detect_output_template(path)

    assert result["layout"] == "matrix"
    assert result["header_row"] == 3
    assert result["ready"] is True
    assert result["missing_fields"] == []
    assert [item["weekday"] for item in result["weekday_columns"]] == [2, 3, 4, 5]
