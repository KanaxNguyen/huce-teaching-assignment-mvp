from datetime import date
import unicodedata

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select,func

from app.db.session import Base,engine,SessionLocal
from app.main import app
from app.models.entities import Lecturer,Semester,NormalizedPreferenceDraft,Constraint,LecturerIdentityAudit
from app.parsers.preferences import _legacy_rule,parse_preference_workbook,_periods_from_text,_weekday_from_text
from app.services.lecturer_master import resolve_identity,confirm_alias
from app.services.preference_validation import validate_preference_draft


@pytest.fixture(autouse=True)
def database():
    Base.metadata.create_all(engine)
    yield
    Base.metadata.drop_all(engine)


def setup(db):
    sem=Semester(name='S',department_name='D',head_name='H',start_date=date(2028,8,1),end_date=date(2029,1,31))
    a=Lecturer(code='GV_A',canonical_name='Canonical A',confirmed=True)
    b=Lecturer(code='GV_B',canonical_name='Canonical B',confirmed=True)
    db.add_all([sem,a,b]);db.flush()
    return sem,a,b


def make_draft(db,sid,lid,alias,row=1,status='DRAFT'):
    d=NormalizedPreferenceDraft(semester_id=sid,lecturer_id=lid,lecturer_alias=alias,source_file='source.xlsx',source_sheet='Data',source_row=row,source_cell=f'B{row}',raw_text=f'Original clause {row}',context_type='SEMINAR',constraint_type='SEMINAR_COMMITMENT',day_scope='T2',periods=[4,5,6],target={'weekday':2,'periods':[4,5,6]},status=status,confidence='HIGH',needs_review=False,rejected_reason='Keep reason' if status=='REJECTED' else None)
    db.add(d);db.flush();return d


@pytest.mark.parametrize('text,field',[
    ('Seminar chiều thứ 3, chưa chốt tiết','period'),
    ('buổi sáng nhưng chưa rõ tiết','period'),
    ('SEMINAR: Thứ 2, thời gian chưa xác định','period'),
    ('không dạy sáng thứ 2 nhưng chưa chốt giờ','period'),
    ('Xin dạy sáng T2, chưa xác định ngày','day'),
    ('Xin dạy tiết 4-6 T3, chưa rõ giảng viên','identity'),
])
def test_explicit_uncertainty_overrides_heuristics(tmp_path,text,field):
    wb=openpyxl.Workbook();ws=wb.active
    ws.append(['Họ và tên','Thứ 2']);ws.append(['Canonical A',text])
    path=tmp_path/'variation.xlsx';wb.save(path)
    parsed=parse_preference_workbook(path)
    assert parsed.drafts and parsed.dropped_clauses==0
    for d in parsed.drafts:
        assert d.confidence_label!='HIGH' and d.needs_review
        if field=='period': assert not d.target.get('periods') and d.start_period is None
        elif field=='day': assert d.day_scope is None
        else: assert d.target['identity_uncertain']


@pytest.mark.parametrize('change',[{'periods':[7,8,9]},{'weight':0.2},{'seminar_link':'Different seminar'}])
def test_each_semantic_edit_needs_reconfirmation(change):
    with SessionLocal() as db:
        sem,a,_=setup(db);d=make_draft(db,sem.id,a.id,'Source',status='CONFIRMED');db.commit();sid,did=sem.id,d.id
    client=TestClient(app);url=f'/api/v1/preference-drafts/{did}?semester_id={sid}'
    response=client.patch(url,json={**change,'status':'CONFIRMED'})
    assert response.status_code==200 and response.json()['status']=='NEEDS_REVIEW'
    assert client.post(f'/api/v1/preference-drafts/apply?semester_id={sid}',json={'draft_ids':[did]}).status_code==422
    assert client.patch(url,json={'status':'CONFIRMED'}).json()['status']=='CONFIRMED'


@pytest.mark.parametrize('change',[{'periods':16},{'target':'bad'},{'start_date':123},{'lecturer_id':[]},{'day_scope':{}},{'weight':None}])
def test_malformed_validation_api_is_structured_not_500(change):
    with SessionLocal() as db:
        sem,a,_=setup(db);db.commit();sid,aid=sem.id,a.id
    response=TestClient(app).post(f'/api/v1/preference-drafts/validate?semester_id={sid}',json={
        'lecturer_id':aid,'context_type':'SEMINAR','constraint_type':'SEMINAR_COMMITMENT','day_scope':'T2',
        'periods':[4,5,6],'hardness':'soft','weight':0,**change})
    assert response.status_code==200
    assert not response.json()['is_confirmable'] and response.json()['validation_errors']


def test_blank_source_identity_does_not_link_other_unknown_rows():
    with SessionLocal() as db:
        sem,a,_=setup(db);first=make_draft(db,sem.id,None,'',1);other=make_draft(db,sem.id,None,'',2)
        db.commit();sid,aid,did,oid=sem.id,a.id,first.id,other.id
    response=TestClient(app).patch(f'/api/v1/identity/draft/{did}?semester_id={sid}',json={'lecturer_id':aid})
    assert response.status_code==200 and response.json()['affected_draft_ids']==[did]
    with SessionLocal() as db:
        assert db.get(NormalizedPreferenceDraft,oid).lecturer_id is None
        assert resolve_identity(db,'',semester_id=sid)[0] is None


@pytest.mark.parametrize('text,day',[('T2',2),('Thứ 2',2),('thứ hai',2),('CN',8),('Chủ nhật',8)])
def test_day_spellings(text,day): assert _weekday_from_text(text)==day


@pytest.mark.parametrize('text,expected',[
    ('tiết 1-3',[1,2,3]),('TIẾT 4 - 6',[4,5,6]),('tiet 4-6',[4,5,6]),('Tiết 4–6',[4,5,6]),
    ('  tiết   4 - 6  ',[4,5,6]),('từ tiết 4 đến tiết 9',list(range(4,10))),
    ('tiết 10-12',[10,11,12]),('tiết 13-15',[13,14,15]),('tiết 1-5',list(range(1,6))),
    ('tiết 2-6',list(range(2,7))),('từ tiết 10 trở đi',list(range(10,16))),
    ('chưa rõ tiết',[]),('chưa chốt giờ',[]),('tiết 4-16',[]),('tiết 16',[]),
])
def test_period_families(text,expected):
    assert _periods_from_text(text)==expected
    assert _periods_from_text(unicodedata.normalize('NFD',text))==expected


@pytest.mark.parametrize('text',['đến 11/10','đến ngày 11 tháng 10','trước 12/10','đến 11/10/2028'])
def test_generic_dates(text):
    kind,target,conf,_=_legacy_rule('Không dạy '+text,None,semester_start=date(2028,8,1),semester_end=date(2029,1,31))
    assert kind=='UNAVAILABLE' and conf=='HIGH' and target['end_date']=='2028-10-11'


@pytest.mark.parametrize('text',['không dạy đến 31/02','không dạy đến 11/10'])
def test_ambiguous_dates_are_review_only(text):
    kind,target,conf,reason=_legacy_rule(text,None,semester_start=date(2028,1,1),semester_end=date(2030,1,1))
    assert kind=='RAW_NOTE' and conf!='HIGH' and reason


@pytest.mark.parametrize('alter,code',[
    ({'lecturer_id':None},'MISSING_LECTURER'),({'periods':[]},'MISSING_PERIOD'),({'day_scope':None},'MISSING_DAY'),
    ({'periods':[16]},'INVALID_PERIOD'),({'constraint_type':'RAW_NOTE'},'UNSUPPORTED_CONSTRAINT_TYPE'),
    ({'numeric_value':None,'constraint_type':'MIN_FREE_MORNING_PER_WEEK','context_type':'TEACHING'},'INVALID_NUMERIC_VALUE'),
])
def test_validator_contract(alter,code):
    item={'lecturer_id':1,'constraint_type':'SEMINAR_COMMITMENT','context_type':'SEMINAR','day_scope':'T2','periods':[4,5,6],'target':{},'weight':0,'hardness':'soft','status':'DRAFT',**alter}
    result=validate_preference_draft(item)
    assert not result['is_confirmable'] and code in {e['code'] for e in result['validation_errors']}


def test_unknown_seminar_cannot_confirm_via_api():
    with SessionLocal() as db:
        sem,a,_=setup(db);d=make_draft(db,sem.id,a.id,'Source');d.periods=[];d.target={};db.commit();sid,did=sem.id,d.id
    client=TestClient(app);base=f'/api/v1/preference-drafts/{did}'
    result=client.post(f'{base}/validate?semester_id={sid}',json={})
    assert result.status_code==200 and not result.json()['is_confirmable']
    assert client.patch(f'{base}?semester_id={sid}',json={'status':'CONFIRMED'}).status_code==422
    result=client.patch(f'{base}?semester_id={sid}',json={'status':'CONFIRMED','periods':[4,5,6]})
    assert result.status_code==200 and result.json()['is_confirmable']
    with SessionLocal() as db:
        assert db.scalar(select(func.count(NormalizedPreferenceDraft.id)))==1


def test_semantic_edit_invalidates_confirmation_but_note_does_not():
    with SessionLocal() as db:
        sem,a,_=setup(db);d=make_draft(db,sem.id,a.id,'Source',status='CONFIRMED');db.commit();sid,did=sem.id,d.id
    client=TestClient(app);url=f'/api/v1/preference-drafts/{did}?semester_id={sid}'
    assert client.patch(url,json={'review_reason':'Cosmetic'}).json()['status']=='CONFIRMED'
    result=client.patch(url,json={'weight':0,'status':'CONFIRMED'})
    assert result.status_code==200 and result.json()['status']=='NEEDS_REVIEW'
    assert result.json()['field_provenance']['weight']['origin']=='HUMAN_ENTERED'
    assert client.patch(url,json={'status':'CONFIRMED'}).json()['status']=='CONFIRMED'


@pytest.mark.parametrize('alias,target_name',[('Unfamiliar Source','Unrelated Canonical'),('Mai Hồng','Mai Thị Hồng')])
def test_identity_relink_correction_preserves_all_drafts_and_other_semester(alias,target_name):
    with SessionLocal() as db:
        sem,a,b=setup(db);b.canonical_name=target_name
        other=Semester(name='Other',department_name='D',head_name='H',start_date=sem.start_date,end_date=sem.end_date);db.add(other);db.flush()
        drafts=[make_draft(db,sem.id,None,alias,i,status) for i,status in enumerate(['DRAFT','CONFIRMED','REJECTED'],1)]
        foreign=make_draft(db,other.id,a.id,alias,8)
        db.commit();sid,oid,aid,bid=sem.id,other.id,a.id,b.id;ids=[d.id for d in drafts];fid=foreign.id
        snapshot={d.id:(d.raw_text,d.source_file,d.source_sheet,d.source_cell) for d in drafts}
    client=TestClient(app)
    url=f'/api/v1/identity/draft/{ids[0]}?semester_id={sid}'
    assert client.patch(url,json={'lecturer_id':aid}).status_code==200
    result=client.patch(url,json={'lecturer_id':bid})
    assert result.status_code==200 and set(result.json()['affected_draft_ids'])==set(ids)
    assert client.patch(f'/api/v1/identity/draft/{fid}?semester_id={sid}',json={'lecturer_id':bid}).status_code==404
    with SessionLocal() as db:
        assert db.scalar(select(func.count(NormalizedPreferenceDraft.id)))==4
        for did in ids:
            d=db.get(NormalizedPreferenceDraft,did)
            assert d.lecturer_id==bid and (d.raw_text,d.source_file,d.source_sheet,d.source_cell)==snapshot[did]
        assert db.get(NormalizedPreferenceDraft,ids[2]).status=='REJECTED'
        assert db.get(NormalizedPreferenceDraft,ids[2]).rejected_reason=='Keep reason'
        assert db.get(NormalizedPreferenceDraft,fid).lecturer_id==aid
        assert resolve_identity(db,alias,semester_id=sid)[0].id==bid
        assert resolve_identity(db,alias,semester_id=oid)[0] is None
        audits=db.scalars(select(LecturerIdentityAudit).order_by(LecturerIdentityAudit.id)).all()
        assert audits[-1].snapshot['before']==aid and audits[-1].snapshot['after']==bid
    review=client.get(f'/api/v1/lecturers/review?semester_id={sid}').json()
    assert next(i for i in review['items'] if i['id']==bid)['identity_status']=='STANDARDIZED'


def test_generic_identity_priority_and_similar_names():
    with SessionLocal() as db:
        sem,a,b=setup(db)
        a.canonical_name='Nguyễn Văn Tuyên';b.canonical_name='Nguyễn Đặng Tuyên'
        confirm_alias(db,a,'Confirmed short');db.commit()
        assert resolve_identity(db,b.canonical_name,code=a.code)[0].id==a.id
        assert resolve_identity(db,'Confirmed short')[0].id==a.id
        assert resolve_identity(db,b.canonical_name)[0].id==b.id
        assert resolve_identity(db,'Tuyên')[0] is None


@pytest.mark.parametrize('offset,reverse,merged',[(0,False,False),(4,True,False),(2,False,True)])
def test_file_structure_variations(tmp_path,offset,reverse,merged):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Arbitrary sheet'
    wb.create_sheet('Unrelated')
    for _ in range(offset): ws.append([])
    headers=['Họ và tên','Mã GV','Thứ hai','Nguyện vọng thêm'];data=['New person','NEW_CODE','SEMINAR: thời gian chưa xác định','Xin dạy tiết 4-6 Thứ 5']
    if reverse: headers.reverse();data.reverse()
    ws.append(headers);ws.append(data)
    if merged:
        ws.append([None,None,'Seminar tiết 13-15',None]);ws.merge_cells(start_row=offset+2,end_row=offset+3,start_column=1,end_column=1);ws.merge_cells(start_row=offset+2,end_row=offset+3,start_column=2,end_column=2)
    path=tmp_path/'arbitrary.xlsx';wb.save(path)
    parsed=parse_preference_workbook(path)
    assert parsed.dropped_clauses==0
    unknown=next(d for d in parsed.drafts if 'chưa xác định' in d.raw_text)
    assert unknown.lecturer_code=='NEW_CODE' and unknown.day_scope=='T2' and unknown.target['periods']==[]
    assert unknown.needs_review and unknown.confidence_label!='HIGH'
    assert unknown.target['_provenance']['periods']['origin']=='UNKNOWN'
    assert next(d for d in parsed.drafts if 'Thứ 5' in d.raw_text).target['weekday']==5


def test_unknown_structure_rejected_without_guessing(tmp_path):
    wb=openpyxl.Workbook();wb.active.append(['random','stuff']);path=tmp_path/'anything.xlsx';wb.save(path)
    with pytest.raises(ValueError,match='PREFERENCE_STRUCTURE_UNRECOGNIZED'):parse_preference_workbook(path)


def test_reimport_preserves_human_edits_ids_and_rejected_state(tmp_path):
    from pathlib import Path
    from app.services.importer import _import_files_impl as import_files
    schedule=Path('/Users/mac/AI/Huce_timetable/INPUT GỐC/Phan_Cong_Giang_Day_HK1_2026_2027.xlsx')
    if not schedule.exists(): pytest.skip('Exact private schedule unavailable')
    wb=openpyxl.Workbook();ws=wb.active;ws.append(['Họ và tên','Ghi chú']);ws.append(['Unknown source','Seminar Thứ 2 tiết 4-6']);ws.append(['Unknown source','Xin dạy tiết 7-9 Thứ 5'])
    path=tmp_path/'preference.xlsx';wb.save(path)
    with SessionLocal() as db:
        sem,a,_=setup(db);db.commit();sid,aid=sem.id,a.id
        import_files(db,[schedule,path],semester_id=sid,schedule_paths=[schedule],preference_paths=[path])
        drafts=db.scalars(select(NormalizedPreferenceDraft).where(NormalizedPreferenceDraft.semester_id==sid).order_by(NormalizedPreferenceDraft.id)).all()
        ids=[d.id for d in drafts]
        assert len(ids)==2 and all(d.confidence!='HIGH' for d in drafts)
    client=TestClient(app)
    assert client.patch(f'/api/v1/identity/draft/{ids[0]}?semester_id={sid}',json={'lecturer_id':aid}).status_code==200
    url=f'/api/v1/preference-drafts/{ids[0]}?semester_id={sid}'
    assert client.patch(url,json={'periods':[10,11,12],'status':'CONFIRMED'}).status_code==200
    assert client.patch(f'/api/v1/preference-drafts/{ids[1]}?semester_id={sid}',json={'status':'REJECTED','rejected_reason':'keep'}).status_code==200
    ws.cell(2,2,'Seminar Thứ 3 tiết 1-3');wb.save(path)
    with SessionLocal() as db:
        result=import_files(db,[schedule,path],semester_id=sid,schedule_paths=[schedule],preference_paths=[path])
        after=db.scalars(select(NormalizedPreferenceDraft).where(NormalizedPreferenceDraft.semester_id==sid).order_by(NormalizedPreferenceDraft.id)).all()
        assert [d.id for d in after]==ids
        assert after[0].periods==[10,11,12] and after[0].status=='CONFIRMED'
        assert after[0].raw_text=='Seminar Thứ 2 tiết 4-6' and after[0].lecturer_id==aid
        assert after[1].status=='REJECTED' and after[1].rejected_reason=='keep'
        assert any(w['code']=='SOURCE_CHANGED_REVIEW_REQUIRED' for w in result['summary']['preferences']['warnings'])


def test_metadata_edit_does_not_resolve_identity():
    with SessionLocal() as db:
        sem,a,_=setup(db);a.code=None;a.confirmed=False
        d=make_draft(db,sem.id,a.id,'Unconfirmed source');db.commit();sid,aid=sem.id,a.id
    client=TestClient(app)
    before=client.get(f'/api/v1/lecturers/review?semester_id={sid}').json()
    assert next(i for i in before['items'] if i['id']==aid)['identity_status']=='NEEDS_CONFIRMATION'
    assert client.patch(f'/api/v1/lecturers/{aid}?semester_id={sid}',json={'name':'Canonical A','note':'Only metadata'}).status_code==200
    after=client.get(f'/api/v1/lecturers/review?semester_id={sid}').json()
    assert next(i for i in after['items'] if i['id']==aid)['identity_status']=='NEEDS_CONFIRMATION'
