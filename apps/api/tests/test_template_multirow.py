from pathlib import Path
from openpyxl import Workbook

from app.services.template_detector import detect_output_template


def test_real_xls_multirow_header_golden():
    fixture_path = Path(__file__).resolve().parents[3] / "storage/uploads/schedule/phan-cong-giang-day-lich-hoc-to-bo-mon-03-09-2026.xls"
    assert fixture_path.exists(), f"Fixture missing: {fixture_path}"

    result = detect_output_template(fixture_path)

    # Header window
    assert result["header_start_row"] == 8
    assert result["header_end_row"] == 9
    assert result["header_row"] == 9
    assert result["ready"] is True
    assert result["missing_fields"] == []

    mappings = result["mappings"]
    # Golden assertions for this specific real fixture
    assert mappings["stt"]["column_letter"] == "A"
    assert mappings["course_code"]["column_letter"] == "B"
    assert mappings["course_name"]["column_letter"] == "C"
    assert mappings["class_code"]["column_letter"] == "D"
    assert mappings["merged_class"]["column_letter"] == "E"
    assert mappings["weekday"]["column_letter"] == "F"
    assert mappings["periods"]["column_letter"] == "G"
    assert mappings["room"]["column_letter"] == "H"
    assert mappings["credits"]["column_letter"] == "I"
    assert mappings["group"]["column_letter"] == "J"
    assert mappings["start_date"]["column_letter"] == "K"
    assert mappings["end_date"]["column_letter"] == "L"
    assert mappings["weeks"]["column_letter"] == "M"
    assert mappings["lecturer"]["column_letter"] == "N"

    # Verify composite headers
    assert mappings["weekday"]["header"] == "Lịch học > Thứ"
    assert mappings["periods"]["header"] == "Lịch học > Tiết"
    assert mappings["room"]["header"] == "Lịch học > Phòng học"
    assert mappings["start_date"]["header"] == "Thời gian học > Bắt đầu"
    assert mappings["end_date"]["header"] == "Thời gian học > Kết thúc"

    # Verify dynamic historical summary
    summary = result["historical_summary"]
    assert summary["columns_count"] >= 14
    assert summary["data_rows_count"] > 0
    assert summary["courses_count"] > 0
    assert summary["lecturers_count"] > 0


def test_synthetic_xlsx_multirow_header(tmp_path):
    path = tmp_path / "synthetic_multirow.xlsx"
    book = Workbook()
    sheet = book.active

    # Row 1: parent header
    sheet.append(["STT", "Môn học", "Mã lớp", "Lịch học", None, None, "Thời gian", None, "Giảng viên"])
    # Row 2: child header
    sheet.append([None, None, None, "Thứ", "Tiết", "Phòng", "Bắt đầu", "Kết thúc", None])
    # Row 3: sample data
    sheet.append([1, "Toán 1", "70IT1", 2, "1-3", "201.H1", "2026-09-01", "2026-12-31", "Nguyễn Văn A"])

    # Merge ranges
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:B2")
    sheet.merge_cells("C1:C2")
    sheet.merge_cells("D1:F1")
    sheet.merge_cells("G1:H1")
    sheet.merge_cells("I1:I2")

    book.save(path)

    result = detect_output_template(path)

    assert result["header_start_row"] == 1
    assert result["header_end_row"] == 2
    assert result["header_row"] == 2
    assert result["ready"] is True
    assert result["missing_fields"] == []

    mappings = result["mappings"]
    assert mappings["stt"]["column_letter"] == "A"
    assert mappings["course_name"]["column_letter"] == "B"
    assert mappings["class_code"]["column_letter"] == "C"
    assert mappings["weekday"]["column_letter"] == "D"
    assert mappings["periods"]["column_letter"] == "E"
    assert mappings["room"]["column_letter"] == "F"
    assert mappings["start_date"]["column_letter"] == "G"
    assert mappings["end_date"]["column_letter"] == "H"
    assert mappings["lecturer"]["column_letter"] == "I"

    assert mappings["weekday"]["header"] == "Lịch học > Thứ"
    assert mappings["periods"]["header"] == "Lịch học > Tiết"
    assert mappings["room"]["header"] == "Lịch học > Phòng"


def test_single_row_template_backward_compatibility(tmp_path):
    path = tmp_path / "single_row.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["STT", "Mã học phần", "Tên môn học", "Mã lớp", "Thứ", "Tiết học", "Phòng", "Giảng viên"])
    sheet.append([1, "391912", "Toán 6", "69CLC1", 3, "4-6", "406.H1", "Phạm Đức Thoan"])
    book.save(path)

    result = detect_output_template(path)

    assert result["header_start_row"] == 1
    assert result["header_end_row"] == 1
    assert result["header_row"] == 1
    assert result["ready"] is True
    assert result["missing_fields"] == []
    assert result["mappings"]["lecturer"]["column_letter"] == "H"
    assert result["mappings"]["weekday"]["column_letter"] == "E"
    assert result["mappings"]["periods"]["column_letter"] == "F"


def test_user_override_persists_over_redetection(tmp_path):
    """Test that human mapping overrides take precedence over redetection (Requirement 9)."""
    path = tmp_path / "test_override.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Mã lớp", "Tên môn", "Thứ", "Tiết", "Giảng viên", "Phòng khác"])
    sheet.append(["71CSQT", "Giải tích", 2, "1-3", "Nguyễn Văn A", "401.H1"])
    book.save(path)

    result = detect_output_template(path)
    # Simulate user changing room from unmapped to F
    user_saved_mappings = dict(result["mappings"])
    user_saved_mappings["room"] = {
        "column_index": 6,
        "column_letter": "F",
        "header": "Phòng khác",
        "confidence": 1.0,
    }

    # Simulate get_latest_template overlaying profile.mappings
    fresh_detect = detect_output_template(path)
    for field, mapping in user_saved_mappings.items():
        fresh_detect["mappings"][field] = mapping

    assert fresh_detect["mappings"]["room"]["column_letter"] == "F"
    assert fresh_detect["mappings"]["room"]["header"] == "Phòng khác"
