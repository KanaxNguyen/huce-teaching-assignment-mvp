import pytest
import os
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

# Use absolute path to the database
db_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../storage/database/huce_test.db"))
engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)

@pytest.fixture
def db():
    db = SessionLocal()
    yield db
    db.close()

def test_hard_unavailable_not_bypassed(db):
    from app.models.entities import ClassSection, Constraint, Semester, Lecturer
    from app.optimization.occurrences import meeting_occurrences
    from app.optimization.solver import _slot_match
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload
    
    semester = db.scalar(select(Semester).where(Semester.id == 2))
    assert semester is not None, "Semester 2 not found"
    
    khien = db.scalar(select(Lecturer).where(Lecturer.canonical_name == "Trần Văn Khiên"))
    assert khien
    # Find Khiên's T5 4-6 constraint
    c_khien = None
    for c in db.scalars(select(Constraint).where(Constraint.lecturer_id == khien.id, Constraint.constraint_type == 'unavailable')).all():
        if c.target.get("weekday") == 5 and 4 in c.target.get("periods", []):
            c_khien = c
            break
    
    g_71kte = db.scalar(select(ClassSection).options(selectinload(ClassSection.sessions)).where(ClassSection.class_code == "71KTE", ClassSection.course_id == 4))
    
    s_71kte_t5 = [s for s in g_71kte.sessions if s.weekday == 5][0]
    
    occ = meeting_occurrences(s_71kte_t5, semester, c_khien.target)
    assert len(occ) > 0, "Khiên's constraint must overlap 71KTE T5 session"
    assert _slot_match(s_71kte_t5, c_khien.target, semester) is True

    g_71kde2 = db.scalar(select(ClassSection).options(selectinload(ClassSection.sessions)).where(ClassSection.class_code == "71KDE2"))
    s_71kde2_t5 = [s for s in g_71kde2.sessions if s.weekday == 5][0]
    assert _slot_match(s_71kde2_t5, c_khien.target, semester) is True

    thuy = db.scalar(select(Lecturer).where(Lecturer.canonical_name == "Vũ Thị Thủy"))
    c_thuy = None
    for c in db.scalars(select(Constraint).where(Constraint.lecturer_id == thuy.id, Constraint.constraint_type == 'unavailable')).all():
        if c.target.get("weekday") == 5 and 4 in c.target.get("periods", []):
            c_thuy = c
            break

    g_71kte_thuy = db.scalar(select(ClassSection).options(selectinload(ClassSection.sessions)).where(ClassSection.class_code == "71KTE", ClassSection.id == 368))
    s_71kte_thuy_t5 = [s for s in g_71kte_thuy.sessions if s.weekday == 5][0]
    assert _slot_match(s_71kte_thuy_t5, c_thuy.target, semester) is True

def test_export_unassigned_fallback(db):
    import openpyxl
    from pathlib import Path
    excel_path = Path("FINAL_TKB_HUCE_HK1_2026_2027.xlsx")
    if not excel_path.exists():
        excel_path = Path(__file__).resolve().parents[3] / "FINAL_TKB_HUCE_HK1_2026_2027.xlsx"
    assert excel_path.exists(), f"FINAL_TKB file not found at {excel_path}"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Phân công chi tiết"]
    unassigned_checked = 0
    for r in range(5, ws.max_row + 1):
        status = ws.cell(r, 8).value
        lec = ws.cell(r, 7).value
        if status == "CHƯA PHÂN CÔNG":
            assert lec == "Chưa phân công"
            unassigned_checked += 1
    assert unassigned_checked > 0, "Expected at least one unassigned row in export"

def test_run_persistence_unassigned(db):
    with open("apps/api/app/optimization/solver.py") as f:
        content = f.read()
    assert "flag_modified(run, \"summary\")" in content

def test_dashboard_metrics(db):
    from app.services.importer import dashboard
    metrics = dashboard(db, 2)
    assert "unassigned_classes" in metrics
    assert "optimization_status" in metrics
    assert metrics["classes"] == 158
    assert metrics["unassigned_classes"] == 6

def test_solver_maximum_valid_solution(db):
    from app.models.entities import OptimizationRun, Assignment
    from sqlalchemy import select, func
    run = db.scalars(select(OptimizationRun).where(OptimizationRun.semester_id == 2).order_by(OptimizationRun.id.desc())).first()
    assert run is not None, "OptimizationRun for semester 2 must exist"
    assert run.status == "optimal"
    assert len(run.summary.get("unassigned", [])) == 6
    assigned_count = db.scalar(select(func.count(Assignment.id)).where(Assignment.run_id == run.id))
    assert assigned_count == 152
