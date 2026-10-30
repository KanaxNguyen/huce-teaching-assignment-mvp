"""Locking assignments that come with an imported schedule (report from 05-10-2026)."""
from openpyxl import load_workbook
from sqlalchemy import select
from app.models.entities import Assignment, ClassSection
from app.optimization.solver import solve
from app.parsers.schedule import parse_periods
from app.services.manual_assignment import lock_assigned_classes
from app.services.source_authority import stage_source, preview_activation, activate_source
from test_qa_full_audit_regressions import db, setup, workbook


def _source(db, tmp_path, rows, teacher='[00000]Lecturer 0'):
    s, *_ = setup(db, 0, 1)
    path = workbook(tmp_path / 'source.xlsx', teacher)
    book = load_workbook(path)
    base = [cell.value for cell in book.active[2]]
    for index, changes in enumerate(rows, start=2):
        row = list(base); row[0] = index
        for column, value in changes.items(): row[column] = value
        book.active.append(row)
    book.save(path)
    src = stage_source(db, path, s.id, 'CURRENT_SCHEDULE'); db.commit()
    return s, src


def _activate(db, s, src, lock=True):
    preview = preview_activation(db, s.id, src.id, 'CURRENT_SCHEDULE')
    return activate_source(db, s.id, src.id, 'CURRENT_SCHEDULE', confirmed=True,
                           preview_token=preview['preview_token'], lock_imported=lock)


def test_same_lecturer_identical_schedule_is_confirmed_merge_and_both_lock(db, tmp_path):
    s, src = _source(db, tmp_path, [{3: 'QA1'}])
    result = _activate(db, s, src)
    groups = list(db.scalars(select(ClassSection)))
    assert result['lock']['locked'] == 2 and result['lock']['skipped'] == []
    assert len(result['lock']['confirmed_merges']) == 1
    assert all(g.locked_assignment and g.merge_status == 'confirmed' and g.merged_confirmed for g in groups)
    run = solve(db, 2, False, s.id)
    assert len(list(db.scalars(select(Assignment).where(Assignment.run_id == run.id)))) == 2


def test_blank_room_identical_schedule_with_same_lecturer_is_merged(db, tmp_path):
    # Blank rooms: the parser makes no suggestion, the source lecturer still merges them.
    s, *_ = setup(db, 0, 1)
    path = workbook(tmp_path / 'blank.xlsx', '[00000]Lecturer 0', room='')
    book = load_workbook(path); row = [cell.value for cell in book.active[2]]; row[0] = 2; row[3] = 'QA1'
    book.active.append(row); book.save(path)
    src = stage_source(db, path, s.id, 'CURRENT_SCHEDULE'); db.commit()
    assert all(g.merged_group_id is None for g in db.scalars(select(ClassSection)))
    result = _activate(db, s, src)
    assert result['lock']['locked'] == 2 and result['lock']['skipped'] == []
    assert all(g.merge_status == 'confirmed' for g in db.scalars(select(ClassSection)))


def test_different_lecturers_never_auto_merge(db, tmp_path):
    from app.models.entities import Lecturer, LecturerCourseCapability, Course
    s, src = _source(db, tmp_path, [{3: 'QA1', 13: '[00001]Lecturer 1'}])
    other = Lecturer(code='00001', canonical_name='Lecturer 1', confirmed=True); db.add(other); db.flush()
    db.add(LecturerCourseCapability(lecturer_id=other.id, course_id=db.scalar(select(Course.id)), allowed=True, confirmed=True, source='MANUAL')); db.commit()
    result = _activate(db, s, src)
    assert result['lock']['confirmed_merges'] == []
    assert all(g.merge_status == 'candidate' for g in db.scalars(select(ClassSection)))


def test_lock_all_button_confirms_source_merges(db, tmp_path):
    s, src = _source(db, tmp_path, [{3: 'QA1'}])
    _activate(db, s, src, lock=False)
    result = lock_assigned_classes(db, s.id)
    assert result['locked'] == 2 and result['skipped'] == [] and len(result['confirmed_merges']) == 1


def test_period_16_online_session_is_accepted():
    assert parse_periods('14-16') == (14, 16)
