from datetime import date
from pathlib import Path

from fastapi.testclient import TestClient
import openpyxl
import pytest
from sqlalchemy import select

from app.db.session import Base, SessionLocal, engine
from app.exporters.excel import export_latest, export_matrix
from app.main import app
from app.models.entities import (
    Assignment,
    ClassSection,
    ClassSession,
    Course,
    Lecturer,
    LecturerCourseCapability,
    LecturerSemesterProfile,
    OptimizationRun,
    OutputTemplateProfile,
    Semester,
    Seminar,
)


@pytest.fixture(autouse=True)
def database():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    yield
    Base.metadata.drop_all(engine)


def test_matrix_export_dynamic_lecturer_count_and_cell_multiline(tmp_path):
    """User Correction 1: Lecturer count is dynamic (NOT hardcoded to 20)."""
    with SessionLocal() as db:
        semester = Semester(
            name="HK1 2026-2027",
            department_name="Toán học",
            start_date=date(2026, 9, 7),
            end_date=date(2027, 1, 24),
            head_name="Trưởng bộ môn",
        )
        c1 = Course(code="MATH1", name="Giải tích 1")
        c2 = Course(code="MATH2", name="Đại số tuyến tính")

        # Exactly 3 lecturers
        l1 = Lecturer(code="001", canonical_name="Giảng viên Một", confirmed=True)
        l2 = Lecturer(code="002", canonical_name="Giảng viên Hai", confirmed=True)
        l3 = Lecturer(code="003", canonical_name="Giảng viên Ba", confirmed=True)
        db.add_all([semester, c1, c2, l1, l2, l3])
        db.flush()

        # Set participation profiles for 3 lecturers
        for lec in (l1, l2, l3):
            db.add(LecturerSemesterProfile(
                semester_id=semester.id,
                lecturer_id=lec.id,
                participation_status="ACTIVE",
            ))

        # Classes and sessions (including evening period 13-15)
        s1 = ClassSection(semester_id=semester.id, course_id=c1.id, class_code="L01", source_file="s", source_sheet="s", source_row=1)
        s2 = ClassSection(semester_id=semester.id, course_id=c2.id, class_code="L02", source_file="s", source_sheet="s", source_row=2)
        s3 = ClassSection(semester_id=semester.id, course_id=c1.id, class_code="L03_EVENING", source_file="s", source_sheet="s", source_row=3)
        db.add_all([s1, s2, s3])
        db.flush()

        # s1: Thứ 2, tiết 1-3
        db.add(ClassSession(class_id=s1.id, weekday=2, start_period=1, end_period=3, room="101.H1", raw_weeks="12", active_weeks=[1, 2], source_row=1))
        # s2: Thứ 2, tiết 4-6 (same day as s1 for l1 to verify multiline separator)
        db.add(ClassSession(class_id=s2.id, weekday=2, start_period=4, end_period=6, room="102.H1", raw_weeks="12", active_weeks=[1, 2], source_row=2))
        # s3: Thứ 3, tiết 13-15 (evening)
        db.add(ClassSession(class_id=s3.id, weekday=3, start_period=13, end_period=15, room="208.H1", raw_weeks="12", active_weeks=[1, 2], source_row=3))

        # Shared Seminar on Thứ 4, tiết 4-6 for l1 and l2
        seminar = Seminar(
            semester_id=semester.id,
            name="Giáo trình Toán",
            chair_name="Nguyễn Bằng Giang",
            members=[l1.id, l2.id],
            alternatives=[{"weekday": 4, "start_period": 4, "end_period": 6}],
        )
        db.add(seminar)

        run = OptimizationRun(semester_id=semester.id, status="feasible", summary={})
        db.add(run)
        db.flush()

        # Assignments: l1 gets s1 and s2 (both on Thứ 2)
        db.add(Assignment(semester_id=semester.id, run_id=run.id, class_id=s1.id, lecturer_id=l1.id))
        db.add(Assignment(semester_id=semester.id, run_id=run.id, class_id=s2.id, lecturer_id=l1.id))
        # l2 gets s3 (Thứ 3 evening 13-15)
        db.add(Assignment(semester_id=semester.id, run_id=run.id, class_id=s3.id, lecturer_id=l2.id))
        db.commit()

        sid = semester.id

    out_file = export_matrix(db, tmp_path / "matrix", sid, mode="draft")
    assert out_file.exists()
    assert out_file.suffix == ".xlsx"

    wb = openpyxl.load_workbook(out_file)
    assert "TKB_Bo_Mon" in wb.sheetnames
    ws = wb["TKB_Bo_Mon"]

    # Exactly 3 lecturer rows: Row 4, Row 5, Row 6 (Header is rows 1-3)
    lecturer_rows = [ws.cell(r, 1).value for r in range(4, ws.max_row + 1)]
    assert len(lecturer_rows) == 3
    assert set(lecturer_rows) == {"Giảng viên Một", "Giảng viên Hai", "Giảng viên Ba"}

    # Check multiline separator in l1's Thứ 2 cell (col 2)
    l1_row = next(r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value == "Giảng viên Một")
    thu_2_content = ws.cell(l1_row, 2).value
    assert "Giải tích 1 - L01" in thu_2_content
    assert "Đại số tuyến tính - L02" in thu_2_content
    assert "\n----------------\n" in thu_2_content

    # Check seminar in l1's Thứ 4 cell (col 4)
    thu_4_content = ws.cell(l1_row, 4).value
    assert "SEMINAR GIÁO TRÌNH TOÁN" in thu_4_content
    assert "Nguyễn Bằng Giang chủ trì" in thu_4_content

    # Check evening period 13-15 in l2's Thứ 3 cell (col 3)
    l2_row = next(r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value == "Giảng viên Hai")
    thu_3_content = ws.cell(l2_row, 3).value
    assert "Tiết 13-15" in thu_3_content
    assert "208.H1" in thu_3_content


def test_matrix_export_endpoint(tmp_path):
    with SessionLocal() as db:
        sem = Semester(name="Test Sem", department_name="Toán", start_date=date(2026, 9, 1), end_date=date(2027, 1, 1), head_name="H")
        lec = Lecturer(code="G1", canonical_name="GV Test", confirmed=True)
        db.add_all([sem, lec]); db.flush()
        c = Course(code="C1", name="Môn Test")
        db.add(c); db.flush()
        sec = ClassSection(semester_id=sem.id, course_id=c.id, class_code="T1", source_file="f", source_sheet="s", source_row=1)
        db.add(sec); db.flush()
        db.add(ClassSession(class_id=sec.id, weekday=2, start_period=1, end_period=3, room="101", raw_weeks="1", active_weeks=[1], source_row=1))
        run = OptimizationRun(semester_id=sem.id, status="feasible", summary={})
        db.add(run); db.flush()
        db.add(Assignment(semester_id=sem.id, run_id=run.id, class_id=sec.id, lecturer_id=lec.id))
        db.commit()
        sid = sem.id

    client = TestClient(app)
    resp = client.get(f"/api/v1/exports/matrix?semester_id={sid}")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.content[:2] == b"PK"  # Valid ZIP/XLSX header
