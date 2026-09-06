from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.api.routes import apply_preference_drafts
from app.db.session import Base, SessionLocal, engine
from app.main import app
from app.models.entities import Constraint, ImportBatch, Lecturer, NormalizedPreferenceDraft, Semester
from app.parsers.preferences import detect_preference_format, parse_preference_workbook
from app.schemas.api import PreferenceDraftApply
from app.services.importer import _clear_imported_data


@pytest.fixture(autouse=True)
def database():
    Base.metadata.drop_all(engine); Base.metadata.create_all(engine)
    yield
    Base.metadata.drop_all(engine)


def _semester(db, name="Preference V2"):
    item = Semester(name=name, department_name="Toán", start_date=date(2026, 8, 1), end_date=date(2027, 1, 31), head_name="H")
    db.add(item); db.commit(); return item


def test_canonical_template_has_exact_structure_validations_and_examples():
    path = Path(__file__).resolve().parents[3] / "data/fixtures/Template_Nguyen_vong_Giang_v2.xlsx"
    assert detect_preference_format(path) == "STRUCTURED_V2"
    book = load_workbook(path)
    assert book.sheetnames == ["Huong_dan", "Nguyen_vong_GV", "Seminar_Shared", "Danh_muc"]
    assert [cell.value for cell in book["Nguyen_vong_GV"][1]] == [
        "Mã GV*", "Họ tên GV", "Ngữ cảnh*", "Loại ràng buộc*", "Thứ/Phạm vi*", "Tiết bắt đầu",
        "Tiết kết thúc", "Ngày bắt đầu", "Ngày kết thúc", "Độ cứng*", "Trọng số",
        "Giá trị số", "Ghi chú / Nguyên văn", "Trạng thái",
    ]
    assert len(book["Nguyen_vong_GV"].data_validations.dataValidation) >= 6
    parsed = parse_preference_workbook(path)
    assert parsed.invalid_rows == 0 and len(parsed.drafts) == 8 and len(parsed.seminars) == 1
    soft_one = next(item for item in parsed.drafts if item.weight == 1.0 and item.hardness == "soft")
    assert soft_one.constraint_type == "PREFERRED_PERIOD"


def test_structured_period_range_is_inclusive_and_invalid_rows_require_review(tmp_path):
    path = tmp_path / "anything.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Nguyen_vong_GV"
    sheet.append(["Mã GV*", "Họ tên GV", "Ngữ cảnh*", "Loại ràng buộc*", "Thứ/Phạm vi*", "Tiết bắt đầu", "Tiết kết thúc", "Ngày bắt đầu", "Ngày kết thúc", "Độ cứng*", "Trọng số", "Giá trị số", "Ghi chú / Nguyên văn", "Trạng thái"])
    sheet.append(["GV1", "A", "TEACHING", "UNAVAILABLE", "T4", 4, 6, None, None, "SOFT", 1, None, "4-6", "DRAFT"])
    sheet.append([None, "Unknown", "TEACHING", "UNKNOWN", "Monday", 9, 7, None, None, "maybe", 3, None, "bad", "CONFIRMED"])
    book.save(path)
    result = parse_preference_workbook(path)
    assert result.format == "STRUCTURED_V2" and result.invalid_rows == 1
    assert result.drafts[0].target["periods"] == [4, 5, 6]
    assert result.drafts[0].hardness == "soft" and result.drafts[0].weight == 1
    assert result.drafts[1].constraint_type == "RAW_NOTE" and result.drafts[1].needs_review


def test_structured_context_and_rule_mismatch_requires_review(tmp_path):
    path = tmp_path / "context-mismatch.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Nguyen_vong_GV"
    sheet.append(["Mã GV*", "Họ tên GV", "Ngữ cảnh*", "Loại ràng buộc*", "Thứ/Phạm vi*", "Tiết bắt đầu", "Tiết kết thúc", "Ngày bắt đầu", "Ngày kết thúc", "Độ cứng*", "Trọng số", "Giá trị số", "Ghi chú / Nguyên văn", "Trạng thái"])
    sheet.append(["GV1", "A", "SEMINAR", "UNAVAILABLE", "T4", 4, 6, None, None, "SOFT", 1, None, "Sai ngữ cảnh", "CONFIRMED"])
    book.save(path)

    item = parse_preference_workbook(path).drafts[0]

    assert item.status == "NEEDS_REVIEW" and item.needs_review
    assert item.review_reason == "Loại rule chưa phù hợp ngữ cảnh đã chọn."


def test_legacy_golden_semantics_and_zero_clause_loss(tmp_path):
    path = tmp_path / "legacy.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append([]); sheet.append([]); sheet.append(["STT", "Tên", "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "CN", "Ghi chú"])
    sheet.append([1, "A", "Xin nghỉ sáng thứ 6. Và xin dạy liền 6 tiết nếu được", None, "không dạy tiết 10-12 thứ 4", None, None, None, None, "Xin dạy ít"])
    book.save(path)
    result = parse_preference_workbook(path)
    unavailable = next(item for item in result.drafts if item.constraint_type == "UNAVAILABLE" and item.day_scope == "T6")
    assert unavailable.target["periods"] == [1, 2, 3, 4, 5, 6]
    by_type = {item.constraint_type: item for item in result.drafts}
    assert by_type["PREFER_CONSECUTIVE_PERIODS"].numeric_value == 6
    assert any(item.day_scope == "T4" and item.target["periods"] == [10, 11, 12] for item in result.drafts)
    assert any(item.constraint_type == "RAW_NOTE" and item.needs_review for item in result.drafts)
    assert result.dropped_clauses == 0


def test_legacy_golden_day_and_period_phrases_are_deterministic(tmp_path):
    cases = [
        "Seminar 4-6", "xin không dạy tiết 10-12 thứ 4", "dạy sáng", "dạy chiều",
        "T2 đến T6", "3 tiết cuối sáng", "3 tiết đầu chiều", "Xin tránh tiết 123",
        "Xin dạy ít", "Xin nghỉ ít nhất 01 buổi sáng", "xin dạy liền 6 tiết",
    ]
    path = tmp_path / "legacy-goldens.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append([]); sheet.append([]); sheet.append(["STT", "Tên", "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "CN", "Ghi chú"])
    for index, source in enumerate(cases, 1):
        sheet.append([index, f"GV {index}", source])
    book.save(path)

    result = parse_preference_workbook(path)
    by_source = {item.raw_text: item for item in result.drafts}

    assert by_source["Seminar 4-6"].target["periods"] == [4, 5, 6]
    assert by_source["xin không dạy tiết 10-12 thứ 4"].target["periods"] == [10, 11, 12]
    assert by_source["xin không dạy tiết 10-12 thứ 4"].day_scope == "T4"
    assert by_source["dạy sáng"].constraint_type == "PREFERRED_PERIOD"
    assert by_source["dạy sáng"].target["periods"] == list(range(1, 7))
    assert by_source["dạy chiều"].target["periods"] == list(range(7, 13))
    assert by_source["T2 đến T6"].day_scope == "ALL_WEEKDAYS"
    assert by_source["3 tiết cuối sáng"].target["periods"] == [4, 5, 6]
    assert by_source["3 tiết đầu chiều"].target["periods"] == [7, 8, 9]
    assert by_source["Xin tránh tiết 123"].target["periods"] == [1, 2, 3]
    assert by_source["Xin dạy ít"].needs_review
    assert by_source["Xin nghỉ ít nhất 01 buổi sáng"].numeric_value == 1
    assert by_source["xin dạy liền 6 tiết"].numeric_value == 6
    assert result.raw_clauses == len(cases) and result.dropped_clauses == 0


def test_apply_route_is_semester_scoped_and_unresolved_identity_never_becomes_global():
    with SessionLocal() as db:
        semester = _semester(db); other = _semester(db, "Other")
        lecturer = Lecturer(code="GV1", canonical_name="Teacher One", confirmed=True); db.add(lecturer); db.flush()
        batch = ImportBatch(semester_id=semester.id, source_files=["v2.xlsx"], summary={}); db.add(batch); db.flush()
        valid = NormalizedPreferenceDraft(semester_id=semester.id, import_batch_id=batch.id, lecturer_id=lecturer.id, lecturer_code="GV1", lecturer_alias="A", constraint_type="UNAVAILABLE", day_scope="T4", periods=[4, 5, 6], target={"weekday": 4, "periods": [4, 5, 6]}, source_file="v2.xlsx", source_sheet="Nguyen_vong_GV", source_row=2, source_cell="A2:M2", confidence="HIGH", needs_review=False, status="CONFIRMED")
        unresolved = NormalizedPreferenceDraft(semester_id=semester.id, import_batch_id=batch.id, lecturer_alias="Unknown", constraint_type="UNAVAILABLE", periods=[1, 2, 3], target={"periods": [1, 2, 3]}, source_file="v2.xlsx", source_sheet="Nguyen_vong_GV", source_row=3, source_cell="A3:M3", confidence="LOW", needs_review=True, status="NEEDS_REVIEW")
        db.add_all([valid, unresolved]); db.commit(); sid, oid, valid_id, unresolved_id = semester.id, other.id, valid.id, unresolved.id
    with TestClient(app) as client:
        assert client.get(f"/api/v1/preference-drafts?semester_id={oid}").json() == []
        rejected = client.post(f"/api/v1/preference-drafts/apply?semester_id={sid}", json={"draft_ids": [valid_id, unresolved_id]})
        assert rejected.status_code == 422
        applied = client.post(f"/api/v1/preference-drafts/apply?semester_id={sid}", json={"draft_ids": [valid_id]})
        assert applied.status_code == 200 and applied.json()["constraints"] == 1
    with SessionLocal() as db:
        rules = db.scalars(select(Constraint).where(Constraint.semester_id == sid)).all()
        assert len(rules) == 1 and rules[0].lecturer_id is not None
        assert not db.scalar(select(Constraint).where(Constraint.lecturer_id.is_(None)))


def test_apply_rolls_back_all_mutations_on_failure(monkeypatch):
    with SessionLocal() as db:
        semester = _semester(db); lecturer = Lecturer(code="GV", canonical_name="Teacher", confirmed=True); db.add(lecturer); db.flush()
        drafts = [NormalizedPreferenceDraft(semester_id=semester.id, lecturer_id=lecturer.id, constraint_type="UNAVAILABLE", periods=[period], target={"weekday": 2, "periods": [period]}, source_file="x", source_sheet="s", source_row=period, source_cell=f"A{period}", confidence="HIGH", needs_review=False, status="CONFIRMED") for period in (1, 2)]
        db.add_all(drafts); db.commit(); sid = semester.id; ids = [item.id for item in drafts]
        original_flush = db.flush; calls = 0
        def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2: raise RuntimeError("forced")
            return original_flush(*args, **kwargs)
        monkeypatch.setattr(db, "flush", fail_second)
        with pytest.raises(RuntimeError):
            apply_preference_drafts(PreferenceDraftApply(draft_ids=ids), semester_id=sid, db=db)
    with SessionLocal() as fresh:
        assert not fresh.scalars(select(Constraint).where(Constraint.semester_id == sid)).all()


def test_manual_teaching_and_seminar_context_are_owned_and_persisted():
    with SessionLocal() as db:
        semester = _semester(db); lecturer = Lecturer(code="GV1", canonical_name="Teacher", confirmed=True)
        db.add(lecturer); db.commit(); sid, lid = semester.id, lecturer.id
    teaching = {"lecturer_id": lid, "context_type": "TEACHING", "constraint_type": "UNAVAILABLE", "day_scope": "T4", "periods": [10, 11, 12], "hardness": "soft", "weight": 1.0, "note": "Không dạy T4", "status": "DRAFT"}
    seminar = {"lecturer_id": lid, "context_type": "SEMINAR", "constraint_type": "SEMINAR_COMMITMENT", "day_scope": "T5", "periods": [4, 5, 6], "hardness": "soft", "weight": 0.8, "seminar_link": "SEM-1", "note": "Seminar", "status": "DRAFT"}
    with TestClient(app) as client:
        assert client.post(f"/api/v1/preference-drafts?semester_id={sid}", json=teaching).status_code == 200
        assert client.post(f"/api/v1/preference-drafts?semester_id={sid}", json=seminar).status_code == 200
        rows = client.get(f"/api/v1/preference-drafts?semester_id={sid}").json()
    assert {(row["context_type"], row["lecturer_id"]) for row in rows} == {("TEACHING", lid), ("SEMINAR", lid)}
    assert next(row for row in rows if row["context_type"] == "SEMINAR")["seminar_link"] == "SEM-1"


def test_manual_mixed_creates_atomic_drafts_and_cannot_be_applied_as_one_rule():
    with SessionLocal() as db:
        semester = _semester(db); lecturer = Lecturer(code="GV1", canonical_name="Teacher", confirmed=True)
        db.add(lecturer); db.commit(); sid, lid = semester.id, lecturer.id
    payload = {"lecturer_id": lid, "context_type": "MIXED", "constraint_type": "RAW_NOTE", "periods": [], "hardness": "soft", "weight": 0.8, "note": "T2,4 seminar nên dạy T3,5", "status": "DRAFT", "parts": [
        {"context_type": "SEMINAR", "constraint_type": "SEMINAR_NOTE", "day_scope": "T2", "periods": [], "hardness": "soft", "weight": 0.8, "note": "seminar", "status": "DRAFT"},
        {"context_type": "TEACHING", "constraint_type": "PREFERRED_PERIOD", "day_scope": "T3", "periods": [1, 2, 3], "hardness": "soft", "weight": 0.8, "note": "teaching", "status": "DRAFT"},
    ]}
    with TestClient(app) as client:
        response = client.post(f"/api/v1/preference-drafts?semester_id={sid}", json=payload)
        assert response.status_code == 200 and response.json() == {"ids": response.json()["ids"], "created": 2, "atomic_split": True}
        rows = client.get(f"/api/v1/preference-drafts?semester_id={sid}").json()
        assert {row["context_type"] for row in rows} == {"TEACHING", "SEMINAR"}
        assert all(row["context_type"] != "MIXED" for row in rows)
        blocked = client.post(f"/api/v1/preference-drafts/apply?semester_id={sid}", json={"draft_ids": [row["id"] for row in rows]})
        assert blocked.status_code == 422


def test_manual_context_correction_persists_and_clears_seminar_only_metadata():
    with SessionLocal() as db:
        semester = _semester(db); lecturer = Lecturer(code="GV1", canonical_name="Teacher", confirmed=True); db.add(lecturer); db.flush()
        draft = NormalizedPreferenceDraft(semester_id=semester.id, lecturer_id=lecturer.id, context_type="SEMINAR", context_confidence="HIGH", constraint_type="SEMINAR_COMMITMENT", seminar_link="SEM-OLD", periods=[4, 5, 6], target={"seminar_link": "SEM-OLD", "periods": [4, 5, 6]}, source_file="x", source_sheet="s", source_row=1, source_cell="A1", confidence="HIGH", needs_review=False, status="DRAFT")
        db.add(draft); db.commit(); sid, did = semester.id, draft.id
    with TestClient(app) as client:
        changed = client.patch(f"/api/v1/preference-drafts/{did}?semester_id={sid}", json={"context_type": "TEACHING"})
        assert changed.status_code == 200
        row = client.get(f"/api/v1/preference-drafts?semester_id={sid}").json()[0]
        assert row["context_type"] == "TEACHING" and row["context_confirmed"] is True
        assert row["seminar_link"] is None and "seminar_link" not in row["target"]
        changed_back = client.patch(f"/api/v1/preference-drafts/{did}?semester_id={sid}", json={"context_type": "SEMINAR", "constraint_type": "SEMINAR_NOTE"})
        assert changed_back.status_code == 200
        persisted = client.get(f"/api/v1/preference-drafts?semester_id={sid}").json()[0]
    assert persisted["context_type"] == "SEMINAR" and persisted["status"] == "NEEDS_REVIEW"


def test_edit_rebuilds_scope_dates_and_numeric_target_without_stale_values():
    with SessionLocal() as db:
        semester = _semester(db); lecturer = Lecturer(code="GV1", canonical_name="Teacher", confirmed=True); db.add(lecturer); db.flush()
        draft = NormalizedPreferenceDraft(semester_id=semester.id, lecturer_id=lecturer.id, context_type="TEACHING", constraint_type="MAX_CLASSES", day_scope="T2", periods=[1], numeric_value=4, start_date=date(2026, 8, 1), end_date=date(2026, 8, 31), target={"weekday": 2, "periods": [1], "start_date": "2026-08-01", "end_date": "2026-08-31", "max": 4}, source_file="x", source_sheet="s", source_row=1, source_cell="A1", confidence="HIGH", needs_review=False, status="DRAFT")
        db.add(draft); db.commit(); sid, did = semester.id, draft.id
    with TestClient(app) as client:
        response = client.patch(
            f"/api/v1/preference-drafts/{did}?semester_id={sid}",
            json={"day_scope": "ALL_WEEKDAYS", "periods": [4, 5, 6], "start_date": None, "end_date": None, "constraint_type": "MIN_CLASSES", "numeric_value": 2},
        )
    assert response.status_code == 200
    target = response.json()["target"]
    assert "weekday" not in target and "start_date" not in target and "end_date" not in target
    assert target["min"] == 2 and "max" not in target
    assert target["slots"] == [{"weekday": day, "periods": [4, 5, 6]} for day in range(2, 7)]


def test_reimport_cleanup_preserves_human_confirmed_context():
    with SessionLocal() as db:
        semester = _semester(db); lecturer = Lecturer(code="GV1", canonical_name="Teacher", confirmed=True); db.add(lecturer); db.flush()
        draft = NormalizedPreferenceDraft(semester_id=semester.id, lecturer_id=lecturer.id, context_type="SEMINAR", context_confidence="HIGH", context_confirmed=True, constraint_type="SEMINAR_NOTE", periods=[], target={"context_type": "SEMINAR"}, source_file="preferences.xlsx", source_sheet="Nguyen_vong_GV", source_row=2, source_cell="A2:N2", raw_text="original source", confidence="HIGH", needs_review=True, status="CONFIRMED")
        db.add(draft); db.commit(); sid, did = semester.id, draft.id

        _clear_imported_data(db, sid)
        db.commit()

    with SessionLocal() as db:
        persisted = db.get(NormalizedPreferenceDraft, did)
        assert persisted is not None
        assert persisted.context_type == "SEMINAR" and persisted.context_confirmed is True
        assert persisted.raw_text == "original source"


def test_legacy_mixed_statement_splits_and_seminar_note_never_reaches_solver(tmp_path):
    path = tmp_path / "legacy-mixed.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append([]); sheet.append([]); sheet.append(["STT", "Tên", "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "CN", "Ghi chú"])
    source = "T2,4 có lịch seminar nên xin ghép dạy gọn T3,5,6,7"
    sheet.append([1, "A", source])
    book.save(path)
    parsed = parse_preference_workbook(path)
    assert len(parsed.drafts) == 2 and {item.context_type for item in parsed.drafts} == {"TEACHING", "SEMINAR"}
    assert all(item.raw_text == source for item in parsed.drafts)
    assert next(item for item in parsed.drafts if item.context_type == "SEMINAR").constraint_type == "SEMINAR_NOTE"
    assert all(item.needs_review for item in parsed.drafts)


def test_old_draft_without_context_is_read_as_teaching():
    with SessionLocal() as db:
        semester = _semester(db)
        draft = NormalizedPreferenceDraft(semester_id=semester.id, context_type=None, constraint_type="RAW_NOTE", periods=[], target={}, source_file="legacy", source_sheet="s", source_row=1, source_cell="A1", confidence="LOW", needs_review=True, status="NEEDS_REVIEW")
        db.add(draft); db.commit(); sid = semester.id
    with TestClient(app) as client:
        row = client.get(f"/api/v1/preference-drafts?semester_id={sid}").json()[0]
    assert row["context_type"] == "TEACHING"
