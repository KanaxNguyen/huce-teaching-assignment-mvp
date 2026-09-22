from __future__ import annotations

import os
import sqlite3
import subprocess
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.api.routes import _temporary_typed_upload
from app.core.config import Settings, settings
from app.db.session import SessionLocal
from app.exporters.excel import export_latest
from app.main import app
from app.models.entities import (
    Assignment,
    ClassSection,
    ClassSession,
    Course,
    Lecturer,
    LecturerCourseCapability,
    OptimizationRun,
    OutputTemplateProfile,
    Semester,
)
from app.storage import LocalStorage, S3CompatibleStorage

ROOT = Path(__file__).resolve().parents[3]


class FakeS3Client:
    def __init__(self):
        self.objects: dict[tuple[str, str], bytes] = {}

    def upload_file(self, filename: str, bucket: str, key: str) -> None:
        self.objects[(bucket, key)] = Path(filename).read_bytes()

    def download_file(self, bucket: str, key: str, filename: str) -> None:
        Path(filename).write_bytes(self.objects[(bucket, key)])


class FakeMissingObjectError(Exception):
    response = {"Error": {"Code": "NoSuchKey"}}


class FakeMissingS3Client(FakeS3Client):
    def download_file(self, bucket: str, key: str, filename: str) -> None:
        raise FakeMissingObjectError


def _workbook_bytes() -> bytes:
    stream = BytesIO()
    book = Workbook()
    book.active.append(["Mã học phần", "Mã lớp học", "Giảng viên", "Thứ", "Tiết học"])
    book.active.append(["C", "L1", "A", "Thứ 2", "1-3"])
    book.save(stream)
    return stream.getvalue()


def test_local_storage_template_round_trip_without_overwrite(tmp_path):
    source = tmp_path / "template.xlsx"
    source.write_bytes(_workbook_bytes())
    storage = LocalStorage(tmp_path / "objects")
    reference = storage.put(source, "templates/1/unique/template.xlsx")
    assert reference == "local://templates/1/unique/template.xlsx"
    with storage.materialize(reference) as materialized:
        assert materialized.read_bytes() == source.read_bytes()
    with pytest.raises(ValueError, match="STORAGE_OBJECT_ALREADY_EXISTS"):
        storage.put(source, "templates/1/unique/template.xlsx")
    with pytest.raises(ValueError, match="STORAGE_KEY_INVALID"):
        storage.put(source, "../outside.xlsx")


def test_temporary_upload_strips_path_and_avoids_collisions():
    first = UploadFile(filename="../../same.xlsx", file=BytesIO(b"first"))
    second = UploadFile(filename="same.xlsx", file=BytesIO(b"second"))
    with _temporary_typed_upload(first) as first_path:
        with _temporary_typed_upload(second) as second_path:
            assert first_path.name == second_path.name == "same.xlsx"
            assert first_path.parent != second_path.parent
            assert first_path.read_bytes() == b"first"
            assert second_path.read_bytes() == b"second"
        assert not second_path.exists()
    assert not first_path.exists()


def test_s3_compatible_storage_adapter_round_trip(tmp_path):
    source = tmp_path / "template.xlsx"
    source.write_bytes(_workbook_bytes())
    fake = FakeS3Client()
    storage = S3CompatibleStorage(bucket="private", region="test", client=fake)
    reference = storage.put(source, "templates/2/unique/template.xlsx")
    assert reference == "s3://private/templates/2/unique/template.xlsx"
    with storage.materialize(reference) as materialized:
        assert materialized.name == "template.xlsx"
        assert materialized.read_bytes() == source.read_bytes()


def test_s3_missing_object_becomes_controlled_missing_file():
    storage = S3CompatibleStorage(
        bucket="private", region="test", client=FakeMissingS3Client()
    )
    with pytest.raises(FileNotFoundError, match="s3://private/templates/missing.xlsx"):
        with storage.materialize("s3://private/templates/missing.xlsx"):
            pass


def test_staging_configuration_fails_fast_when_persistence_or_security_is_missing():
    with pytest.raises(ValueError, match="Thiếu cấu hình deployment"):
        Settings(environment="staging", _env_file=None)


def test_staging_configuration_rejects_replica_startup_migrations():
    with pytest.raises(ValueError, match="RUN_MIGRATIONS_ON_STARTUP=false"):
        Settings(
            environment="staging",
            database_url="postgresql+psycopg://example.invalid/app",
            allowed_origins="https://staging.example.invalid",
            storage_backend="s3",
            s3_bucket="private",
            s3_region="test",
            s3_access_key_id="key",
            s3_secret_access_key="secret",
            internal_api_token="token",
            run_migrations_on_startup=True,
            _env_file=None,
        )


def test_staging_configuration_rejects_wildcard_cors():
    with pytest.raises(ValueError, match="không dùng wildcard"):
        Settings(
            environment="staging",
            database_url="postgresql+psycopg://example.invalid/app",
            allowed_origins="*",
            storage_backend="s3",
            s3_bucket="private",
            s3_region="test",
            s3_access_key_id="key",
            s3_secret_access_key="secret",
            internal_api_token="token",
            _env_file=None,
        )


@pytest.mark.parametrize("scheme", ["postgres://", "postgresql://"])
def test_managed_postgresql_urls_select_psycopg_v3(scheme):
    configured = Settings(database_url=f"{scheme}user:pass@db.example/app", _env_file=None)
    assert configured.database_url == "postgresql+psycopg://user:pass@db.example/app"


def test_release_migration_command_uses_database_url_environment(tmp_path):
    database = tmp_path / "release-command.db"
    environment = {**os.environ, "DATABASE_URL": f"sqlite:///{database}"}
    result = subprocess.run(
        [sys.executable, "-m", "alembic", "-c", "alembic.ini", "upgrade", "head"],
        cwd=ROOT,
        env=environment,
        capture_output=True,
        text=True,
        timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    with sqlite3.connect(database) as connection:
        assert connection.execute("SELECT version_num FROM alembic_version").fetchone() == (
            "0009_department_agnostic_capabilities",
        )


def test_development_startup_creates_sqlite_parent_directory(tmp_path):
    database = tmp_path / "nested" / "database" / "development.db"
    environment = {
        **os.environ,
        "DATABASE_URL": f"sqlite:///{database}",
        "ENVIRONMENT": "development",
        "PYTHONPATH": str(ROOT / "apps/api"),
        "UPLOAD_DIR": str(tmp_path / "uploads"),
        "EXPORT_DIR": str(tmp_path / "exports"),
    }
    script = """
from fastapi.testclient import TestClient
from app.main import app
with TestClient(app) as client:
    assert client.get('/api/v1/health').status_code == 200
"""
    result = subprocess.run(
        [sys.executable, "-c", script],
        cwd=ROOT,
        env=environment,
        capture_output=True,
        text=True,
        timeout=45,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert database.is_file()


def test_staging_api_protection_uses_server_side_header(monkeypatch):
    monkeypatch.setattr(settings, "environment", "staging")
    monkeypatch.setattr(settings, "internal_api_token", "deployment-test-token")
    with TestClient(app) as client:
        assert client.get("/api/v1/health").status_code == 200
        assert client.get("/").status_code == 401
        assert client.get(
            "/", headers={"x-internal-api-key": "wrong"}
        ).status_code == 401
        assert client.get(
            "/", headers={"x-internal-api-key": "deployment-test-token"}
        ).status_code == 200


def test_cors_allows_configured_origin_and_rejects_unlisted_origin():
    request_headers = {"Access-Control-Request-Method": "GET"}
    with TestClient(app) as client:
        allowed = client.options(
            "/api/v1/health",
            headers={"Origin": "http://localhost:3000", **request_headers},
        )
        rejected = client.options(
            "/api/v1/health",
            headers={"Origin": "https://unlisted.example", **request_headers},
        )
    assert allowed.status_code == 200
    assert allowed.headers["access-control-allow-origin"] == "http://localhost:3000"
    assert "access-control-allow-origin" not in rejected.headers


def test_cors_allows_configured_origin_and_rejects_other_origin():
    headers = {
        "origin": "http://127.0.0.1:3000",
        "access-control-request-method": "GET",
    }
    with TestClient(app) as client:
        allowed = client.options("/api/v1/semesters", headers=headers)
        rejected = client.options(
            "/api/v1/semesters",
            headers={**headers, "origin": "https://not-configured.example"},
        )
    assert allowed.headers["access-control-allow-origin"] == "http://127.0.0.1:3000"
    assert "access-control-allow-origin" not in rejected.headers


def test_upload_size_extension_and_malformed_workbook_are_controlled(monkeypatch):
    monkeypatch.setattr(settings, "max_upload_bytes", 32)
    with TestClient(app) as client:
        semester = client.post("/api/v1/semesters", json={
            "name": "Upload security", "department_name": "Department",
            "start_date": "2026-01-01", "end_date": "2026-06-01", "head_name": "Head",
        })
        assert semester.status_code == 200
        unsupported = client.post(
            "/api/v1/templates/detect", files={"template_file": ("bad.txt", b"x", "text/plain")}
        )
        oversized = client.post(
            "/api/v1/templates/detect",
            files={"template_file": ("large.xlsx", b"x" * 33, "application/octet-stream")},
        )
    assert unsupported.status_code == 415
    assert oversized.status_code == 413

    monkeypatch.setattr(settings, "max_upload_bytes", 1024)
    with TestClient(app) as client:
        malformed = client.post(
            "/api/v1/templates/detect",
            files={"template_file": ("bad.xlsx", b"not a workbook", "application/octet-stream")},
        )
    assert malformed.status_code == 422
    assert "traceback" not in malformed.text.casefold()


def test_template_upload_persists_storage_reference_and_reloads(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    workbook_data = _workbook_bytes()
    with TestClient(app) as client:
        semester = client.post(
            "/api/v1/semesters",
            json={
                "name": "Persistent template",
                "department_name": "Department",
                "start_date": "2026-01-01",
                "end_date": "2026-06-01",
                "head_name": "Head",
            },
        )
        semester_id = semester.json()["id"]
        detected = client.post(
            f"/api/v1/templates/detect?semester_id={semester_id}",
            files={
                "template_file": (
                    "template.xlsx",
                    workbook_data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        reloaded = client.get(f"/api/v1/templates/latest?semester_id={semester_id}")

    assert detected.status_code == 200
    assert reloaded.status_code == 200
    assert reloaded.json()["profile_id"] == detected.json()["profile_id"]
    with SessionLocal() as db:
        profile = db.scalar(
            select(OutputTemplateProfile).where(
                OutputTemplateProfile.id == detected.json()["profile_id"]
            )
        )
        assert profile is not None
        assert profile.source_file.startswith(f"local://templates/{semester_id}/")
        with LocalStorage(tmp_path / "uploads" / "objects").materialize(
            profile.source_file
        ) as stored:
            assert stored.read_bytes() == workbook_data


def test_missing_persisted_template_is_controlled_export_error(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    with SessionLocal() as db:
        semester = Semester(
            name="Missing template",
            department_name="Department",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 6, 1),
            head_name="Head",
        )
        db.add(semester)
        db.flush()
        run = OptimizationRun(semester_id=semester.id, status="feasible", summary={})
        db.add(run)
        db.add(
            OutputTemplateProfile(
                semester_id=semester.id,
                source_file="local://templates/missing.xlsx",
                source_sheet="Sheet",
                header_row=1,
                mappings={},
                missing_fields=[],
                preview=[],
            )
        )
        db.commit()
        with pytest.raises(ValueError, match="EXPORT_TEMPLATE_SOURCE_NOT_FOUND"):
            export_latest(db, tmp_path / "exports", semester.id)


def test_persisted_template_export_round_trip_opens_workbook(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    source = tmp_path / "template.xlsx"
    source.write_bytes(_workbook_bytes())
    reference = LocalStorage(tmp_path / "uploads" / "objects").put(
        source, "templates/semester/unique/template.xlsx"
    )
    with SessionLocal() as db:
        semester = Semester(
            name="Export round trip",
            department_name="Department",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 6, 1),
            head_name="Head",
        )
        course = Course(code="C", name="Course")
        lecturer = Lecturer(code="L", canonical_name="Lecturer")
        db.add_all([semester, course, lecturer])
        db.flush()
        section = ClassSection(
            semester_id=semester.id,
            course_id=course.id,
            class_code="L1",
            credits=3,
            assigned_lecturer_id=lecturer.id,
            source_file="template.xlsx",
            source_sheet="Sheet",
            source_row=2,
        )
        db.add(section)
        db.flush()
        db.add_all(
            [
                ClassSession(
                    class_id=section.id,
                    weekday=2,
                    start_period=1,
                    end_period=3,
                    room="P",
                    raw_weeks="1",
                    active_weeks=[1],
                    source_row=2,
                ),
                LecturerCourseCapability(
                    lecturer_id=lecturer.id,
                    course_id=course.id,
                    allowed=True,
                    confirmed=True,
                ),
            ]
        )
        run = OptimizationRun(semester_id=semester.id, status="feasible", summary={})
        db.add(run)
        db.flush()
        db.add(
            Assignment(
                semester_id=semester.id,
                run_id=run.id,
                class_id=section.id,
                lecturer_id=lecturer.id,
                locked=False,
                penalty=0,
            )
        )
        db.add(
            OutputTemplateProfile(
                semester_id=semester.id,
                source_file=reference,
                source_sheet="Sheet",
                header_row=1,
                mappings={
                    "course_code": {"column_index": 1},
                    "class_code": {"column_index": 2},
                    "lecturer": {"column_index": 3},
                },
                missing_fields=[],
                preview=[],
            )
        )
        db.commit()
        result = export_latest(db, tmp_path / "exports", semester.id)

    workbook = load_workbook(result)
    assert workbook["Sheet"].cell(2, 3).value == "Lecturer"
    workbook.close()
