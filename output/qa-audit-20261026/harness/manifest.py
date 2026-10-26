"""Dataset manifest with an independent reader (openpyxl / xlrd only, no app parser).

For sheets in the 14-column HUCE schedule layout (STT, Mã học phần, ..., Tuần, Giảng viên)
it counts valid meeting rows, unique classes (course::class code) and lecturer identities.
Usage: python manifest.py <out_json>
"""
import hashlib
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import xlrd
from openpyxl import load_workbook

ROOT = Path("/Users/mac/AI/Huce_timetable")
FILES = sorted(
    [p for p in (ROOT / "test_fixtures").glob("TC*/*.xlsx")]
    + [p for p in (ROOT / "INPUT Gốc").iterdir() if p.suffix in (".xls", ".xlsx")]
    + [p for p in (ROOT / "INPUT TEST").iterdir() if p.suffix in (".xls", ".xlsx")]
    + [Path("/Users/mac/Downloads") / n for n in (
        "phan-cong-giang-day-lich-hoc-to-bo-mon-03-08-2026.xls",
        "phan-cong-giang-day-lich-hoc-to-bo-mon-19-07-2026.xls",
        "Copy of phan-cong-giang-day-lich-hoc-to-bo-mon-26-07-2026_TKB.xls",
        "phan-cong-giang-day-lich-hoc-to-bo-mon-ky1-2025-2026.xls")]
)


def role(path: Path) -> str:
    import unicodedata
    n = unicodedata.normalize("NFC", path.name).lower()
    if path.parent.name.startswith("TC"):
        return {"timetable_input.xlsx": "FIXTURE_SCHEDULE", "preference_input.xlsx": "FIXTURE_PREFERENCE",
                "expected_assignments.xlsx": "FIXTURE_EXPECTED_OUTPUT (never input)"}.get(path.name, "FIXTURE_OTHER")
    if "nguyen" in n or "nguyện" in n or "phieu" in n:
        return "PREFERENCE"
    if "output" in n or "template" in n:
        return "OUTPUT_TEMPLATE/REFERENCE"
    if "tkb-ncm" in n:
        return "REFERENCE_MATRIX"
    return "SCHEDULE (CURRENT_SCHEDULE or HISTORICAL)"


def rows_of(path):
    if path.suffix == ".xls":
        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            yield sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]


def schedule_stats(rows):
    header_idx = None
    for i, r in enumerate(rows[:15]):
        cells = [str(v or "").strip().lower() for v in r]
        if any("mã lớp" in c for c in cells) and any(c == "thứ" or c.startswith("thứ") for c in cells) and any("tuần" in c for c in cells):
            header_idx = i; break
    if header_idx is None:
        # University portal export: multi-row header, data rows start with a numeric STT, fixed 14 columns.
        numeric = [i for i, r in enumerate(rows) if r and str(r[0]).strip().replace(".0", "").isdigit() and len(r) >= 14]
        if not numeric:
            return None
        header_idx = numeric[0] - 1
        c_course, c_class, c_day, c_period, c_weeks, c_lect = 1, 3, 5, 6, 12, 13
    else:
        head = [str(v or "").strip().lower() for v in rows[header_idx]]
        col = lambda key: next((j for j, h in enumerate(head) if h.startswith(key)), None)  # noqa: E731
        c_course, c_class, c_day, c_period, c_weeks = (col("mã học phần") if col("mã học phần") is not None else col("mã hp")), col("mã lớp"), col("thứ"), col("tiết"), col("tuần")
        c_lect = next((j for j, h in enumerate(head) if "giảng viên" in h), None)
        if None in (c_course, c_class, c_day, c_period):
            return {"header_row": header_idx + 1, "note": f"partial header {head[:16]}"}
    classes, meetings, lecturers, invalid, multi = defaultdict(int), 0, set(), 0, 0
    for r in rows[header_idx + 1:]:
        if not r or c_class is None or c_class >= len(r):
            continue
        key = f"{str(r[c_course]).strip()}::{str(r[c_class]).strip()}"
        if not str(r[c_class] or "").strip():
            continue
        try:
            day = int(float(r[c_day])); periods = [int(x) for x in re.findall(r"\d+", str(r[c_period]))]
            assert 2 <= day <= 8 and periods and 1 <= periods[0] <= periods[-1] <= 15
        except Exception:  # noqa: BLE001
            invalid += 1; continue
        meetings += 1; classes[key] += 1
        raw = str(r[c_lect] or "").strip() if c_lect is not None and c_lect < len(r) else ""
        ids = re.findall(r"\[([^\]]+)\]", raw)
        if len(ids) > 1:
            multi += 1
        lecturers.update(ids)
    return {"header_row": header_idx + 1, "meeting_rows": meetings, "classes": len(classes), "invalid_rows": invalid,
            "rows_with_multiple_lecturers": multi, "distinct_lecturer_codes": len(lecturers)}


def main(out):
    records = []
    for path in FILES:
        if not path.exists():
            records.append({"file": str(path), "status": "MISSING"}); continue
        rec = {"file": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size, "role": role(path), "sheets": []}
        try:
            for name, rows in rows_of(path):
                non_empty = sum(1 for r in rows if any(v not in (None, "") for v in r))
                rec["sheets"].append({"name": name, "rows": len(rows), "non_empty_rows": non_empty, "schedule": schedule_stats(rows)})
        except Exception as exc:  # noqa: BLE001
            rec["error"] = f"{type(exc).__name__}: {exc}"
        records.append(rec)
    Path(out).write_text(json.dumps(records, ensure_ascii=False, indent=1))
    for r in records:
        sched = [s["schedule"] for s in r.get("sheets", []) if s.get("schedule")]
        print(Path(r["file"]).name[:60], r.get("role"), r.get("sha256", "")[:12], [(s.get("classes"), s.get("meeting_rows"), s.get("rows_with_multiple_lecturers"), s.get("note","")[:60]) for s in sched], r.get("error", ""))


if __name__ == "__main__":
    main(sys.argv[1])
