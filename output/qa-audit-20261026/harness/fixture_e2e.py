"""TC01–TC17 through the real HTTP API (TestClient), one fresh SQLite file per case.

Flow: semester -> upload/preview/activate schedule -> upload/preview/activate preferences
-> capabilities from fixture_meta (PUT, global catalog) -> human review simulated by
confirming lecturer-resolved drafts -> apply -> lock fixture_meta.locked to the expected
lecturer -> POST /optimization/run (confirm_merged=False) -> independent checker ->
compare with expected_result.json -> draft export read back with openpyxl.

expected_assignments.xlsx is never used as input.
Usage: PYTHONPATH=apps/api .venv/bin/python fixture_e2e.py <fixtures_dir> <evidence_dir> <work_dir>
"""
import importlib
import json
import os
import sys
from pathlib import Path

FIX, OUT, WORK = (Path(p) for p in sys.argv[1:4])
OUT.mkdir(parents=True, exist_ok=True); WORK.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(Path(__file__).parent))
os.environ.update(ENVIRONMENT="development", UPLOAD_DIR=str(WORK / "up"), EXPORT_DIR=str(WORK / "ex"),
                  DATABASE_URL=f"sqlite:///{WORK / 'fixtures.db'}")
import independent_checker as ic  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from app.db.session import Base, engine  # noqa: E402
from app.main import app  # noqa: E402

P = "/api/v1"
DB = str(WORK / "fixtures.db")


def call(c, method, path, log, **kw):
    r = c.request(method, P + path, **kw)
    try:
        body = r.json()
    except Exception:  # noqa: BLE001
        body = f"<{len(r.content)} bytes {r.headers.get('content-type')}>"
    log.append({"method": method, "path": path, "status": r.status_code, "body": body if len(json.dumps(body, default=str)) < 4000 else str(body)[:4000]})
    return r


def stage(c, sid, f, role, log):
    s = call(c, "POST", f"/sources/upload?semester_id={sid}&source_type={role}", log, files={"file": (f.name, f.read_bytes())}).json()
    p = call(c, "POST", f"/sources/{s['id']}/preview?semester_id={sid}", log, json={"source_type": role}).json()
    a = call(c, "POST", f"/sources/{s['id']}/activate?semester_id={sid}", log, json={"source_type": role, "confirmed": True, "preview_token": p.get("preview_token")})
    return a.status_code


def run_case(case: Path):
    engine.dispose()
    if os.path.exists(DB):
        os.remove(DB)
    Base.metadata.create_all(engine)
    c = TestClient(app); log = []
    meta = json.loads((case / "fixture_meta.json").read_text())
    expected = json.loads((case / "expected_result.json").read_text())
    sid = call(c, "POST", "/semesters", log, json=dict(name=case.name[:100], department_name="Toán học", head_name="QA",
                                                         start_date="2026-09-07", end_date="2027-01-24")).json()["id"]
    act_schedule = stage(c, sid, case / "timetable_input.xlsx", "CURRENT_SCHEDULE", log)
    act_pref = stage(c, sid, case / "preference_input.xlsx", "PREFERENCE", log) if (case / "preference_input.xlsx").exists() else None
    lecturers = call(c, "GET", "/lecturers", log).json()
    by_code = {l["code"]: l["id"] for l in lecturers}
    classes = call(c, "GET", f"/classes?semester_id={sid}", log).json()
    course_ids = {k["course_code"]: k["course_id"] for k in classes}
    missing = []
    for code, courses in meta.get("capabilities", {}).items():
        if code not in by_code:
            r = call(c, "POST", "/lecturers", log, json={"code": code, "canonical_name": f"GV {code}"})
            by_code[code] = r.json().get("id")
        for course_code in courses:
            if course_code not in course_ids:
                missing.append(course_code); continue
            call(c, "PUT", f"/lecturers/{by_code[code]}/capabilities", log, json={"course_id": course_ids[course_code], "allowed": True, "confirmed": True})
    drafts = call(c, "GET", f"/preference-drafts?semester_id={sid}", log).json()
    to_apply = []
    for d in drafts:
        if d.get("lecturer_id") and d.get("draft_kind") == "CONSTRAINT" and d.get("status") not in ("REJECTED",):
            r = call(c, "PATCH", f"/preference-drafts/{d['id']}?semester_id={sid}", log, json={"status": "CONFIRMED"})
            if r.status_code == 200:
                to_apply.append(d["id"])
    applied = call(c, "POST", f"/preference-drafts/apply?semester_id={sid}", log, json={"draft_ids": to_apply}).status_code if to_apply else None
    constraints = call(c, "GET", f"/constraints?semester_id={sid}", log).json()
    lecturer_names = {l["id"]: (l.get("canonical_name") or l.get("name")) for l in call(c, "GET", "/lecturers", log).json()}
    name_to_id = {}
    for d in drafts:
        if d.get("lecturer_id"):
            name_to_id[d.get("lecturer")] = d["lecturer_id"]
    exact = {e["class_code"]: e["lecturer"] for e in expected.get("exact_assignments", [])}
    class_by_code = {k["class_code"]: k for k in classes}
    locks = []
    for code in meta.get("locked", []):
        lid = name_to_id.get(exact.get(code))
        r = call(c, "PATCH", f"/classes/{class_by_code[code]['id']}/assignment?semester_id={sid}", log, json={"lecturer_id": lid, "lock": True}) if lid and code in class_by_code else None
        locks.append((code, lid, r.status_code if r else "NO_LECTURER"))
    run = call(c, "POST", f"/optimization/run?semester_id={sid}", log, json={"time_limit_seconds": 10, "confirm_merged_suggestions": False})
    body = run.json()
    issues, assigned = ([], {})
    if run.status_code == 200 and body.get("status") in {"optimal", "feasible"}:
        issues, assigned = ic.check(DB, sid, body["run_id"])
    code_of = {k["id"]: k["class_code"] for k in classes}
    # Lecturer names come from the preference workbook (drafts), not from the solver.
    id_to_name = {v: k for k, v in name_to_id.items()}
    actual_named = {code_of[cid]: id_to_name.get(lid, lecturer_names.get(lid)) for cid, lid in assigned.items()}
    mismatches = []
    if len(assigned) != expected["expected_assigned"]:
        mismatches.append(f"assigned {len(assigned)} != {expected['expected_assigned']}")
    for code, name in exact.items():
        if actual_named.get(code) != name:
            mismatches.append(f"{code}: expected {name}, got {actual_named.get(code)}")
    for u in expected.get("unassigned", []):
        ucode = u.get("class_code") if isinstance(u, dict) else u
        if ucode in actual_named:
            mismatches.append(f"{ucode} expected unassigned, got {actual_named[ucode]}")
    if issues:
        mismatches.append(f"checker: {issues}")
    export = call(c, "GET", f"/exports/latest?semester_id={sid}&mode=draft", log)
    export_check = None
    if export.status_code == 200 and export.content[:2] == b"PK":
        path = OUT / f"export-{case.name}.xlsx"; path.write_bytes(export.content)
        wb = load_workbook(path, read_only=True)
        cells = [str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None]
        export_check = {code: any(code in cell for cell in cells) for code in class_by_code}
        for code, name in actual_named.items():
            if name and not any(name in cell for cell in cells):
                mismatches.append(f"export missing lecturer {name} for {code}")
        missing_classes = [code for code, ok in export_check.items() if not ok]
        if missing_classes:
            mismatches.append(f"export missing classes {missing_classes}")
    else:
        mismatches.append(f"draft export HTTP {export.status_code}")
    result = {"case": case.name, "activate_schedule": act_schedule, "activate_preference": act_pref, "missing_courses": missing,
              "drafts": len(drafts), "applied_status": applied, "constraints_active": len(constraints), "locks": locks,
              "run_http": run.status_code, "run_status": body.get("status"), "score": body.get("score"),
              "assigned": actual_named, "checker": issues, "mismatches": mismatches,
              "verdict": "PASS" if not mismatches else "FAIL"}
    (OUT / f"api-trace-{case.name}.json").write_text(json.dumps(log, ensure_ascii=False, indent=1, default=str))
    return result


if __name__ == "__main__":
    rows = [run_case(case) for case in sorted(FIX.glob("TC*")) if case.is_dir()]
    (OUT / "fixture-e2e.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str))
    for r in rows:
        print(r["case"], r["verdict"], r["run_status"], r["mismatches"][:3])
