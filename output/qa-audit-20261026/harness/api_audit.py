"""API-level audit on a fresh SQLite file: import edges, source lifecycle, semester isolation,
manual/lock chain, merge confirmation, preference review, export round trip.

Synthetic workbooks are generated into the work directory from the 14-column HUCE layout.
Usage: PYTHONPATH=apps/api .venv/bin/python api_audit.py <evidence_dir> <work_dir>
"""
import json
import os
import sys
from pathlib import Path

OUT, WORK = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True); WORK.mkdir(parents=True, exist_ok=True)
DB = WORK / "api.db"
if DB.exists():
    DB.unlink()
os.environ.update(ENVIRONMENT="development", UPLOAD_DIR=str(WORK / "up"), EXPORT_DIR=str(WORK / "ex"), DATABASE_URL=f"sqlite:///{DB}")
sys.path.insert(0, str(Path(__file__).parent))
import independent_checker as ic  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from app.db.session import Base, engine  # noqa: E402
from app.main import app  # noqa: E402

Base.metadata.create_all(engine)
c = TestClient(app)
P = "/api/v1"
RESULTS, TRACE = [], []
HEADER = ("STT", "Mã học phần", "Tên môn học", "Mã lớp học", "Ghi chú", "Thứ", "Tiết", "Phòng học", "Số TC", "Số SV", "Bắt đầu", "Kết thúc", "Tuần", "Giảng viên")


def call(method, path, **kw):
    r = c.request(method, P + path, **kw)
    try:
        body = r.json()
    except Exception:  # noqa: BLE001
        body = f"<{len(r.content)} bytes>"
    TRACE.append({"method": method, "path": path, "status": r.status_code, "body": str(body)[:1500]})
    return r


def record(case_id, area, scenario, expected, actual, ok, **detail):
    RESULTS.append({"id": case_id, "area": area, "scenario": scenario, "expected": expected, "actual": actual,
                    "verdict": "PASS" if ok else "FAIL", **detail})


def workbook(name, rows):
    wb = Workbook(); ws = wb.active; ws.title = "Schedule"; ws.append(HEADER)
    for i, row in enumerate(rows, 1):
        ws.append((str(i), *row))
    path = WORK / name; wb.save(path); return path


def row(course, cname, code, day, periods, weeks, lecturer=None, room="P1", credits="3"):
    return (course, cname, code, None, str(day), periods, room, credits, "60", "07/09/2026", "24/01/2027", weeks, lecturer)


def semester(name):
    return call("POST", "/semesters", json=dict(name=name, department_name="Toán học", head_name="QA", start_date="2026-09-07", end_date="2027-01-24")).json()["id"]


def stage(sid, path, role="CURRENT_SCHEDULE"):
    up = call("POST", f"/sources/upload?semester_id={sid}&source_type={role}", files={"file": (path.name, path.read_bytes())})
    if up.status_code != 200:
        return up, None, None
    pv = call("POST", f"/sources/{up.json()['id']}/preview?semester_id={sid}", json={"source_type": role})
    return up, pv, (pv.json().get("preview_token") if pv.status_code == 200 else None)


def activate(sid, source_id, token, role="CURRENT_SCHEDULE"):
    return call("POST", f"/sources/{source_id}/activate?semester_id={sid}", json={"source_type": role, "confirmed": True, "preview_token": token})


def lecturer(code, name):
    r = call("POST", "/lecturers", json={"code": code, "name": name})
    return r.json().get("id")


def classes(sid):
    return call("GET", f"/classes?semester_id={sid}").json()


def solve(sid, **extra):
    return call("POST", f"/optimization/run?semester_id={sid}", json={"time_limit_seconds": 10, **extra})


# ------------------------------------------------------------------ import edges
sid_a = semester("QA-A")
empty = WORK / "empty.xlsx"; empty.write_bytes(b"")
bad_ext = WORK / "notes.txt"; bad_ext.write_text("not a workbook")
corrupt = WORK / "corrupt.xlsx"; corrupt.write_bytes(b"PK\x03\x04garbage")
for case_id, path in (("IMP-EMPTY", empty), ("IMP-EXT", bad_ext), ("IMP-CORRUPT", corrupt)):
    up, _, _ = stage(sid_a, path)
    sem = call("GET", f"/semesters").json()
    active = next(s for s in sem if s["id"] == sid_a).get("active_schedule_source_id")
    record(case_id, "Import", f"upload {path.name}", "4xx, active source unchanged (None)", f"{up.status_code}, active={active}",
           400 <= up.status_code < 500 and active is None)
os.environ["MAX_UPLOAD_BYTES"] = "1"  # settings are cached at import, so this only documents the knob
big = WORK / "big.xlsx"; big.write_bytes(b"PK" + b"0" * (21 * 1024 * 1024))
up, _, _ = stage(sid_a, big)
record("IMP-SIZE", "Import", "21 MiB upload", "413", up.status_code, up.status_code == 413)

base_rows = [
    row("391001", "Giải tích 1", "K1", 2, "1-3", "12345", "[00187] Phạm Đức Thoan"),
    row("391001", "Giải tích 1", "K2", 2, "1-3", "12345"),
    row("391001", "Giải tích 1", "K3", 2, "1-3", "12345"),
    row("391002", "Đại số", "K4", 3, "4-6", "  345"),
    row("391002", "Đại số", "K4", 5, "13-15", "12345"),  # second meeting of the same class, evening
    row("391003", "Giải tích 2", "0K5", 4, "7-9", "1 3 5"),
]
wb_a = workbook("schedule-a.xlsx", base_rows)
up, pv, token = stage(sid_a, wb_a)
before = call("GET", "/semesters").json()
staged_not_active = next(s for s in before if s["id"] == sid_a).get("active_schedule_source_id") is None
record("SRC-CANDIDATE", "Source", "upload + preview without activate", "active source stays None", staged_not_active, staged_not_active)
act = activate(sid_a, up.json()["id"], token)
record("SRC-ACTIVATE", "Source", "activate with fresh token", 200, act.status_code, act.status_code == 200)
cls = classes(sid_a)
by_code = {k["class_code"]: k for k in cls}
record("IMP-GROUP", "Import", "6 rows -> 5 classes (K4 has 2 meetings, evening 13-15 kept)", "5 classes; K4 2 sessions",
       f"{len(cls)} classes; K4 {len(by_code.get('K4', {}).get('sessions', []))} sessions", len(cls) == 5 and len(by_code.get("K4", {}).get("sessions", [])) == 2)
record("IMP-LEADING-ZERO", "Import", "class code '0K5' keeps leading zero", "0K5 present", sorted(by_code), "0K5" in by_code)
k4_weeks = sorted({w for s in by_code.get("K4", {}).get("sessions", []) for w in s["active_weeks"]})
record("IMP-WEEKS-GAP", "Import", "week mask '  345' -> weeks 3,4,5", "[1..5] union with [3,4,5]", k4_weeks, k4_weeks == [1, 2, 3, 4, 5])
k5_weeks = by_code.get("0K5", {}).get("sessions", [{}])[0].get("active_weeks")
record("IMP-WEEKS-SPACED", "Import", "week mask '1 3 5' (space = week off)", "[1,3,5]", k5_weeks, k5_weeks == [1, 3, 5])

# Re-upload identical file: no duplicate classes.
up2, pv2, token2 = stage(sid_a, wb_a)
act2 = activate(sid_a, up2.json()["id"], token2) if up2.status_code == 200 else up2
record("SRC-REIMPORT", "Source", "re-import identical schedule", "still 5 classes, no duplicates", f"{act2.status_code}/{len(classes(sid_a))}",
       len(classes(sid_a)) == 5)

# Stale preview token: preview, then change state (lock), then activate.
wb_b = workbook("schedule-b.xlsx", base_rows[:-1])
up3, pv3, token3 = stage(sid_a, wb_b)
k2 = next(k for k in classes(sid_a) if k["class_code"] == "K2")
k4 = next(k for k in classes(sid_a) if k["class_code"] == "K4")
th = next(l for l in call("GET", "/lecturers").json() if l["code"] == "00187")["id"]
call("PUT", f"/lecturers/{th}/capabilities", json={"course_id": k4["course_id"], "allowed": True, "confirmed": True})
lock_r = call("PATCH", f"/classes/{k4['id']}/assignment?semester_id={sid_a}", json={"lecturer_id": th, "lock": True})
stale = activate(sid_a, up3.json()["id"], token3)
record("SRC-STALE-TOKEN", "Source", "activate with token issued before a lock change", "409/422 rejected",
       f"lock={lock_r.status_code}, activate={stale.status_code}", lock_r.status_code == 200 and stale.status_code in (409, 422))
record("SRC-CANCEL", "Source", "abandon preview keeps data", "5 classes", len(classes(sid_a)), len(classes(sid_a)) == 5)

# ------------------------------------------------------------------ semester isolation
sid_b = semester("QA-B")
foreign = by_code["K3"]["id"]
for case_id, method, path, kwargs in (
    ("ISO-LOCK", "POST", f"/classes/{foreign}/lock?semester_id={sid_b}", {}),
    ("ISO-UNLOCK", "POST", f"/classes/{foreign}/unlock?semester_id={sid_b}", {}),
    ("ISO-ASSIGN", "PATCH", f"/classes/{foreign}/assignment?semester_id={sid_b}", {"json": {"lecturer_id": th}}),
    ("ISO-CHECK", "POST", f"/classes/{foreign}/assignment/check?semester_id={sid_b}", {"json": {"lecturer_id": th}}),
    ("ISO-CANDIDATES", "GET", f"/classes/{foreign}/candidates?semester_id={sid_b}", {}),
    ("ISO-ANALYSIS", "GET", f"/classes/{foreign}/candidate-analysis?semester_id={sid_b}", {}),
    ("ISO-SCENARIOS", "GET", f"/classes/{foreign}/scenarios?semester_id={sid_b}", {}),
    ("ISO-SCEN-APPLY", "POST", f"/classes/{foreign}/scenarios/apply?semester_id={sid_b}", {"json": {"scenario": "direct", "lecturer_id": th}}),
    ("ISO-RESOLUTION", "PATCH", f"/classes/{foreign}/resolution-status?semester_id={sid_b}", {"json": {"status": "ESCALATED"}}),
    ("ISO-OVERRIDE", "POST", f"/classes/{foreign}/override-lock?semester_id={sid_b}", {"json": {"lecturer_id": th, "reason": "cross semester"}}),
    ("ISO-MERGE", "POST", f"/classes/merge?semester_id={sid_b}", {"json": {"class_ids": [by_code["K2"]["id"], foreign]}}),
    ("ISO-UNMERGE", "POST", f"/classes/unmerge?semester_id={sid_b}", {"json": {"class_ids": [foreign]}}),
    ("ISO-SOURCE-ACT", "POST", f"/sources/{up.json()['id']}/activate?semester_id={sid_b}", {"json": {"source_type": "CURRENT_SCHEDULE", "confirmed": True, "preview_token": "x"}}),
):
    r = call(method, path, **kwargs)
    body = r.json() if r.headers.get("content-type", "").startswith("application/json") else None
    harmless_empty = r.status_code == 200 and (body == [] or (isinstance(body, dict) and body.get("unmerged") == 0))
    record(case_id, "Isolation", f"{method} {path.split('?')[0]} with other semester", "rejected (4xx) or empty no-op; no foreign data",
           f"{r.status_code} {str(body)[:80]}", r.status_code in (400, 403, 404, 409, 422) or harmless_empty,
           note="200 with empty body: no leak, but inconsistent with the 404 used by sibling endpoints" if harmless_empty else None)
foreign_constraint = call("POST", f"/constraints?semester_id={sid_b}", json={"name": "x", "constraint_type": "REQUIRED_ASSIGNMENT", "hardness": "hard",
                                                                         "lecturer_id": th, "target": {"class_id": foreign}})
record("ISO-CONSTRAINT-TARGET", "Isolation", "create hard REQUIRED in semester B targeting a class of semester A", "422 rejected",
       foreign_constraint.status_code, foreign_constraint.status_code in (400, 404, 422))

# ------------------------------------------------------------------ manual / lock chain
cs = classes(sid_a)
course_ids = {k["course_id"] for k in cs}
cu = lecturer("00172", "Lê Viết Cường")
hv = lecturer("00173", "Hoàng Văn")
for lid in (th, cu, hv):
    for cid in course_ids:
        call("PUT", f"/lecturers/{lid}/capabilities", json={"course_id": cid, "allowed": True, "confirmed": True})
call("POST", f"/classes/{k4['id']}/unlock?semester_id={sid_a}")
default_body = call("POST", f"/optimization/run?semester_id={sid_a}", json={}).json()
record("API-DEFAULT-MERGE", "Solver API", "POST /optimization/run with {}", "confirm_merged False (merge suggestions need review)",
       (default_body.get("summary") or {}).get("experiment", {}).get("configuration", {}).get("confirm_merged"),
       (default_body.get("summary") or {}).get("experiment", {}).get("configuration", {}).get("confirm_merged") is False)
merge_state_after_default = {k["class_code"]: k["merge_status"] for k in classes(sid_a) if k["merged_group_id"]}
record("API-DEFAULT-MERGE-SIDE-EFFECT", "Solver API", "merge status after a default run", "no class silently confirmed",
       merge_state_after_default, all(v != "confirmed" for v in merge_state_after_default.values()))

# Fresh semester for the manual chain (default run above confirmed merges).
sid_m = semester("QA-MANUAL")
up_m, _, tok_m = stage(sid_m, wb_a)
activate(sid_m, up_m.json()["id"], tok_m)
cm = {k["class_code"]: k for k in classes(sid_m)}
r1 = solve(sid_m, confirm_merged_suggestions=False).json()
issues1, plan1 = ic.check(str(DB), sid_m, r1["run_id"])
record("SOLVE-INDEP", "Solver", "solve real-format synthetic semester", "0 independent violations", issues1, not issues1, run=r1["run_id"])
target = cm["0K5"]["id"]
current = plan1.get(target)
other = next(l for l in (th, cu, hv) if l != current)
manual = call("PATCH", f"/classes/{target}/assignment?semester_id={sid_m}", json={"lecturer_id": other})
names_by_id = {l["id"]: l.get("name") for l in call("GET", "/lecturers").json()}
assignments = {a["class_id"]: a["lecturer"] for a in call("GET", f"/assignments?semester_id={sid_m}").json()}
record("MAN-EDIT", "Manual", "manual change visible in /assignments", names_by_id.get(other), assignments.get(target),
       manual.status_code == 200 and assignments.get(target) == names_by_id.get(other))
runs = call("GET", f"/optimization/runs?semester_id={sid_m}").json()
solver_run = next((r for r in runs if r["id"] == r1["run_id"]), None)
record("MAN-VERSION", "Manual", "manual edit stored as separate version; solver run metrics untouched",
       "solver run score unchanged", (solver_run or {}).get("score"), solver_run is not None and solver_run.get("score") == r1["score"])
lock = call("POST", f"/classes/{target}/lock?semester_id={sid_m}")
r2 = solve(sid_m, confirm_merged_suggestions=False).json()
issues2, plan2 = ic.check(str(DB), sid_m, r2["run_id"])
record("LOCK-RESOLVE", "Manual", "lock manual choice then re-solve", other, plan2.get(target), plan2.get(target) == other and not issues2)
call("POST", f"/classes/{target}/unlock?semester_id={sid_m}")
r3 = solve(sid_m, confirm_merged_suggestions=False).json()
issues3, plan3 = ic.check(str(DB), sid_m, r3["run_id"])
record("UNLOCK-RESOLVE", "Manual", "after unlock the solver may reconsider (no hard violation)", "0 violations", issues3, not issues3)
bad = call("PATCH", f"/classes/{cm['K2']['id']}/assignment?semester_id={sid_m}", json={"lecturer_id": plan3.get(cm["K1"]["id"]) or th})
clash = plan3.get(cm["K1"]["id"]) is not None
record("MAN-REJECT-CLASH", "Manual", "manual assign K2 to K1's lecturer (same slot)", "422 rejected", bad.status_code, (bad.status_code == 422) if clash else True)

# ------------------------------------------------------------------ merge half-confirmed via API
sid_g = semester("QA-MERGE")
up_g, _, tok_g = stage(sid_g, wb_a)
activate(sid_g, up_g.json()["id"], tok_g)
cg = {k["class_code"]: k for k in classes(sid_g)}
group = cg["K1"].get("merged_group_id")
merge = call("POST", f"/classes/merge?semester_id={sid_g}", json={"class_ids": [cg["K1"]["id"], cg["K2"]["id"]], "merged_group_id": group})
cg = {k["class_code"]: k for k in classes(sid_g)}
statuses = {code: (cg[code]["merged_group_id"], cg[code]["merge_status"]) for code in ("K1", "K2", "K3")}
rg = solve(sid_g, confirm_merged_suggestions=False).json()
issues_g, plan_g = ic.check(str(DB), sid_g, rg["run_id"])
record("MERGE-PARTIAL-API", "Merge", "merge 2 of 3 suggested classes into the existing group id, then solve",
       "no double booking for the still-unconfirmed member", {"statuses": statuses, "issues": issues_g},
       not [i for i in issues_g if i[0] == "DOUBLE_BOOKING"])

# Capacity-constrained variant: only one lecturer can teach the course.
sid_g2 = semester("QA-MERGE-TIGHT")
solo = lecturer("00555", "Solo Lecturer")
tight = workbook("tight.xlsx", [row("391555", "Môn hiếm", f"T{i}", 6, "7-9", "12345") for i in range(3)])
up_t, _, tok_t = stage(sid_g2, tight); activate(sid_g2, up_t.json()["id"], tok_t)
ct = {k["class_code"]: k for k in classes(sid_g2)}
call("PUT", f"/lecturers/{solo}/capabilities", json={"course_id": ct["T0"]["course_id"], "allowed": True, "confirmed": True})
call("POST", f"/classes/merge?semester_id={sid_g2}", json={"class_ids": [ct["T0"]["id"], ct["T1"]["id"]], "merged_group_id": ct["T0"].get("merged_group_id")})
ct = {k["class_code"]: k for k in classes(sid_g2)}
rt = solve(sid_g2, confirm_merged_suggestions=False).json()
issues_t, plan_t = ic.check(str(DB), sid_g2, rt["run_id"])
record("MERGE-PARTIAL-API-TIGHT", "Merge", "merge 2 of 3 identical suggested classes into the existing group id; one capable lecturer",
       "T2 (still candidate) not double-booked with T0/T1", {"statuses": {k: (v["merged_group_id"], v["merge_status"]) for k, v in ct.items()}, "plan": plan_t, "issues": issues_t},
       not [i for i in issues_t if i[0] == "DOUBLE_BOOKING"])

# ------------------------------------------------------------------ export round trip
for export_type in ("detailed", "matrix"):
    path = "/exports/latest" if export_type == "detailed" else "/exports/matrix"
    r = call("GET", f"{path}?semester_id={sid_m}&mode=draft" + ("&export_type=detailed" if export_type == "detailed" else ""))
    ok = r.status_code == 200 and r.content[:2] == b"PK"
    detail = {}
    if ok:
        f = OUT / f"export-{export_type}.xlsx"; f.write_bytes(r.content)
        wb = load_workbook(f, read_only=True)
        text = [str(v) for ws in wb.worksheets for rr in ws.iter_rows(values_only=True) for v in rr if v is not None]
        current_plan = {a["class_id"]: a for a in call("GET", f"/assignments?semester_id={sid_m}").json()}
        missing = [k["class_code"] for k in classes(sid_m) if export_type == "detailed" and not any(k["class_code"] in t for t in text)]
        names = {a.get("lecturer") or a.get("lecturer_name") for a in current_plan.values()}
        absent = [n for n in names if n and not any(n in t for t in text)]
        detail = {"missing_classes": missing, "absent_lecturers": absent}
        ok = not missing and not absent
    record(f"EXPORT-{export_type.upper()}", "Export", f"draft {export_type} export contains current plan", "all classes and lecturers present",
           detail or r.status_code, ok)
final = call("GET", f"/exports/latest?semester_id={sid_m}&mode=final")
ready = call("GET", f"/readiness?semester_id={sid_m}").json()
is_ready = bool(ready.get("ready") or ready.get("final_ready") or ready.get("status") == "READY")
record("EXPORT-FINAL-GATE", "Export", "final export allowed iff readiness is ready", f"{'200' if is_ready else '422'} (ready={is_ready})",
       final.status_code, final.status_code == (200 if is_ready else 422), readiness=str(ready)[:400])
# Make the semester not ready (unassign by removing all capability of one course is global; instead add a class-less blocker):
sid_nr = semester("QA-NOT-READY")
up_nr, _, tok_nr = stage(sid_nr, workbook("not-ready.xlsx", [row("391777", "Môn không ai dạy", "X1", 6, "1-3", "12345")]))
activate(sid_nr, up_nr.json()["id"], tok_nr)
solve(sid_nr, confirm_merged_suggestions=False)
final_nr = call("GET", f"/exports/latest?semester_id={sid_nr}&mode=final")
record("EXPORT-FINAL-GATE-UNREADY", "Export", "final export with an unassigned class and no override", "422", final_nr.status_code, final_nr.status_code == 422)

# ------------------------------------------------------------------ preference review dead-ends
pref = Workbook(); ws = pref.active; ws.title = "Preferences"
ws.append(("Mã GV*", "Họ tên GV", "Ngữ cảnh*", "Loại ràng buộc*", "Thứ/Phạm vi*", "Tiết bắt đầu", "Tiết kết thúc", "Ngày bắt đầu", "Ngày kết thúc", "Độ cứng*", "Trọng số", "Giá trị số", "Ghi chú / Nguyên văn", "Trạng thái"))
ws.append(("00187", "Phạm Đức Thoan", "TEACHING", "PREFERRED_DAYS", "T3", None, None, None, None, "SOFT", 0.8, None, "Ưu tiên T3", "CONFIRMED"))
ws.append(("00187", "Phạm Đức Thoan", "TEACHING", "REQUIRED_ASSIGNMENT", "ALL_WEEKDAYS", None, None, None, None, "HARD", 1, None, "Bắt buộc dạy K1", "CONFIRMED"))
ws.append(("99999", "Người Lạ", "TEACHING", "UNAVAILABLE", "T2", 1, 3, None, None, "HARD", 1, None, "GV chưa có trong danh mục", "CONFIRMED"))
pref_path = WORK / "pref.xlsx"; pref.save(pref_path)
upp, pvp, tokp = stage(sid_m, pref_path, "PREFERENCE")
activate(sid_m, upp.json()["id"], tokp, "PREFERENCE")
drafts = call("GET", f"/preference-drafts?semester_id={sid_m}").json()
for d in drafts:
    kind = d["constraint_type"]
    if kind == "PREFERRED_DAYS":
        call("PATCH", f"/preference-drafts/{d['id']}?semester_id={sid_m}", json={"day_scope": "T3"})
        r = call("PATCH", f"/preference-drafts/{d['id']}?semester_id={sid_m}", json={"status": "CONFIRMED"})
        record("PREF-DAYS-CONFIRM", "Preference", "confirm V2 PREFERRED_DAYS 'T3' after re-selecting the day", "200 confirmable", r.status_code, r.status_code == 200,
               errors=r.json().get("detail") if r.status_code != 200 else None)
    elif kind == "REQUIRED_ASSIGNMENT":
        r = call("PATCH", f"/preference-drafts/{d['id']}?semester_id={sid_m}", json={"status": "CONFIRMED"})
        record("PREF-REQUIRED-CONFIRM", "Preference", "V2 REQUIRED_ASSIGNMENT row: reviewer can pick the class and confirm",
               "a field to choose the TeachingGroup exists", "PreferenceDraftUpdate has no class/target field; status " + str(r.status_code), False)
    elif d.get("lecturer_code") == "99999" or d.get("lecturer_alias") == "Người Lạ":
        cons = [x for x in call("GET", f"/constraints?semester_id={sid_m}").json() if x.get("lecturer_id") is None]
        record("PREF-UNRESOLVED", "Preference", "unknown lecturer code stays a draft, never a global constraint", "lecturer_id None, needs review, no active constraint",
               {"lecturer_id": d.get("lecturer_id"), "status": d.get("status"), "needs_review": d.get("needs_review"), "global_constraints": len(cons)},
               d.get("lecturer_id") is None and not cons)
lecs = {l["code"]: l for l in call("GET", "/lecturers").json()}
record("PREF-AUTO-LECTURER", "Lecturer", "preference import with unknown code 99999 does not create a confirmed master lecturer",
       "no ACTIVE confirmed lecturer 99999", lecs.get("99999"), "99999" not in lecs or not lecs["99999"].get("confirmed"))

(OUT / "api-audit.json").write_text(json.dumps(RESULTS, ensure_ascii=False, indent=1, default=str))
(OUT / "api-audit-trace.json").write_text(json.dumps(TRACE, ensure_ascii=False, indent=1, default=str))
for r in RESULTS:
    print(r["verdict"], r["id"], "|", str(r["actual"])[:160])
